"""One-off/rerunnable import: pull vendor classifications out of a folder of
historical "Creditors Ageing" report exports and add any vendor the
Creditors Ageing app's own mapping table doesn't already know about.

These monthly reports were built by hand over a run of different export
shapes (an older desktop app's pivot-style export, then this suite's own
report format) and never got round-tripped back into the app's central
mapping table, so a chunk of vendor classifications have only ever existed
inside these standalone files. This script recovers them.

Handles both export shapes seen in the historical files:
  - This suite's own report layout: sheets "Only Creditors"/"Advances"/
    "Intercompany", header row with a literal "Vendor Name" column.
  - The older desktop app's Excel PivotTable exports: sheets named
    "Only Crs"/"Advances" (no "Intercompany" sheet), header row uses
    "Row Labels" instead of "Vendor Name" (Excel's own default pivot
    field header) but the same Location/Vendor Type/Vendor Sub Type
    columns alongside it.
Sheet roles are matched by substring (e.g. "Only Crs" and "Only
Creditors" both count as the non-intercompany creditors sheet), so sheet
naming drift across exports doesn't need special-casing here.

Purely additive, like every other seed path in this app: an existing
active OR archived vendor_key is left completely alone (see
app.soft_delete.seed_missing_keyed_rows) - re-running this after an admin
has since edited a mapping in the UI can never clobber that edit. Safe to
run more than once and against more than one folder of files.

Usage (from backend/, with the venv active):
    python -m scripts.import_creditors_ageing_history_mappings <folder-or-file> [more...]
    python -m scripts.import_creditors_ageing_history_mappings <folder> --dry-run

Same command works unchanged on the production server: copy the same
folder of .xlsx exports there and point this script at it. Run with
--dry-run first to preview what would be inserted before touching the
live database.
"""

import argparse
import re
import sys
from pathlib import Path

import openpyxl

from app.database import SessionLocal
from app.services.creditors_ageing import mapping_store
from app.services.creditors_ageing.models import VendorMapping
from app.soft_delete import seed_missing_keyed_rows

MONTHS = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "oct": 10,
    "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}
AS_AT_RE = re.compile(r"as\s+(?:at|on)\s+(\d{1,2})\D+([A-Za-z]+)\D+(\d{2,4})", re.IGNORECASE)
NAME_HEADER_TOKENS = ("vendor name", "row labels")
SKIP_NAMES = {"grand total", "subtotal", "(blank)", "total", "-", "--", "n/a", ""}


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and value == 0:
        return ""
    return str(value).strip()


def _sheet_role(sheet_name: str) -> str | None:
    normalized = re.sub(r"[^a-z]", "", sheet_name.lower())
    if "intercompany" in normalized:
        return "intercompany"
    if "advance" in normalized:
        return "advances"
    if "onlycr" in normalized:
        return "only_creditors"
    return None


def _report_date(workbook) -> tuple[int, int] | None:
    """(year, month) from the "As at <date>" text embedded near the top of
    any of this file's report sheets. Returns None if no sheet has it (the
    two oldest pivot-style exports don't)."""
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        for r in range(1, min(5, ws.max_row or 0) + 1):
            for c in range(1, min(6, ws.max_column or 0) + 1):
                value = ws.cell(r, c).value
                if not isinstance(value, str):
                    continue
                m = AS_AT_RE.search(value)
                if not m:
                    continue
                _day, month_name, year = m.groups()
                month = MONTHS.get(month_name.strip().lower())
                if month is None:
                    continue
                year_i = int(year)
                if year_i < 100:
                    year_i += 2000
                return (year_i, month)
    return None


def _date_from_filename(stem: str) -> tuple[int, int] | None:
    lowered = stem.lower()
    month = None
    for name, num in MONTHS.items():
        if re.search(rf"\b{name}\b", lowered):
            month = num
            break
    if month is None:
        return None
    year_m = re.search(r"\b(20\d{2})\b", lowered) or re.search(r"\b(\d{2})\b", lowered)
    if not year_m:
        return None
    year = int(year_m.group(1))
    if year < 100:
        year += 2000
    return (year, month)


def _find_header_row(ws) -> int | None:
    for row_number in range(1, min(10, ws.max_row or 0) + 1):
        values = [_s(ws.cell(row_number, c).value).lower() for c in range(1, (ws.max_column or 0) + 1)]
        if any(any(token in v for token in NAME_HEADER_TOKENS) for v in values if v):
            return row_number
    return None


def _extract_sheet_rows(ws, intercompany: bool) -> list[dict]:
    header_row = _find_header_row(ws)
    if header_row is None:
        return []
    columns: dict[str, int] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        header = _s(ws.cell(header_row, c).value).lower()
        if header and header not in columns:
            columns[header] = c

    name_col = next((c for h, c in columns.items() if any(t in h for t in NAME_HEADER_TOKENS)), None)
    if name_col is None:
        return []
    location_col = next((c for h, c in columns.items() if "location" in h), None)
    type_col = next((c for h, c in columns.items() if "vendor type" in h), None)
    sub_type_col = next((c for h, c in columns.items() if "vendor sub type" in h), None)

    rows = []
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        name = _s(ws.cell(r, name_col).value)
        if not name or name.lower() in SKIP_NAMES:
            continue
        rows.append({
            "vendor_name": name,
            "location": _s(ws.cell(r, location_col).value) if location_col else "",
            "vendor_type": _s(ws.cell(r, type_col).value) if type_col else "",
            "vendor_sub_type": _s(ws.cell(r, sub_type_col).value) if sub_type_col else "",
            "intercompany": intercompany,
        })
    return rows


# Sheet roles are applied in this order so that, if the same vendor genuinely
# appears under more than one role within a single monthly file, the more
# specific/deliberate classification wins - same precedence the app's own
# packaged seed template uses (see creditors_ageing.mapping_store.REPORT_SHEETS).
ROLE_ORDER = ("only_creditors", "advances", "intercompany")


def parse_file(path: Path) -> tuple[dict[str, dict], tuple[int, int] | None]:
    """Returns ({vendor_key: mapping_fields}, (year, month) | None)."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        by_role: dict[str, list[dict]] = {}
        for sheet_name in workbook.sheetnames:
            role = _sheet_role(sheet_name)
            if role is None:
                continue
            by_role[role] = _extract_sheet_rows(workbook[sheet_name], intercompany=(role == "intercompany"))

        merged: dict[str, dict] = {}
        for role in ROLE_ORDER:
            for row in by_role.get(role, []):
                merged[mapping_store.vendor_key(row["vendor_name"])] = row

        report_date = _report_date(workbook) or _date_from_filename(path.stem)
        return merged, report_date
    finally:
        workbook.close()


def _iter_xlsx_paths(inputs: list[str]) -> list[Path]:
    paths: list[Path] = []
    for raw in inputs:
        p = Path(raw)
        if p.is_dir():
            paths.extend(sorted(f for f in p.glob("*.xlsx") if not f.name.startswith("~$")))
        elif p.is_file():
            paths.append(p)
        else:
            print(f"WARNING: {p} does not exist, skipping", file=sys.stderr)
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("paths", nargs="+", help="Folder(s) of .xlsx report exports, or individual .xlsx files")
    parser.add_argument("--dry-run", action="store_true", help="Parse and report only - don't write to the database")
    args = parser.parse_args()

    files = _iter_xlsx_paths(args.paths)
    if not files:
        print("No .xlsx files found.", file=sys.stderr)
        sys.exit(1)

    dated: list[tuple[tuple[int, int], Path, dict[str, dict]]] = []
    undated: list[tuple[Path, dict[str, dict]]] = []
    for path in files:
        print(f"Reading {path.name} ...")
        merged, report_date = parse_file(path)
        print(f"  -> {len(merged)} vendor rows" + (f", dated {report_date[1]:02d}/{report_date[0]}" if report_date else ", no date found"))
        if report_date:
            dated.append((report_date, path, merged))
        else:
            undated.append((path, merged))

    if undated:
        print(
            "\nWARNING: could not determine a report date for: "
            + ", ".join(p.name for p, _ in undated)
            + " - applying these FIRST (lowest priority), so any dated file's "
            "values for the same vendor will override them.",
        )

    # Oldest -> newest, so a later file's classification for the same vendor
    # overwrites an earlier one. Undated files are treated as oldest/lowest
    # priority (see warning above) rather than guessed into the timeline.
    dated.sort(key=lambda item: item[0])
    ordered = [(path, merged) for path, merged in undated] + [(path, merged) for _date, path, merged in dated]

    combined: dict[str, dict] = {}
    sources: dict[str, str] = {}
    conflicts: dict[str, list[tuple[str, dict]]] = {}
    for path, merged in ordered:
        for key, row in merged.items():
            previous = combined.get(key)
            if previous is not None and previous != row:
                history = conflicts.setdefault(key, [(sources[key], previous)])
                history.append((path.name, row))
            combined[key] = row
            sources[key] = path.name

    if conflicts:
        print(f"\n{len(conflicts)} vendor(s) had a differing classification across files (latest file wins):")
        for key, versions in list(conflicts.items())[:20]:
            print(f"  {key}:")
            for filename, row in versions:
                print(f"    [{filename}] location={row['location']!r} type={row['vendor_type']!r} sub_type={row['vendor_sub_type']!r}")
        if len(conflicts) > 20:
            print(f"  ... and {len(conflicts) - 20} more")

    print(f"\n{len(combined)} distinct vendors found across {len(files)} file(s).")

    db = SessionLocal()
    try:
        existing_keys = {
            row.vendor_key
            for row in db.query(VendorMapping.vendor_key).all()
        }
        new_keys = [key for key in combined if key not in existing_keys]
        print(f"{len(new_keys)} are new (not already in the mapping table, active or archived).")
        print(f"{len(combined) - len(new_keys)} already exist in the mapping table and will be left untouched.")

        if args.dry_run:
            print("\n--dry-run: no changes written. New vendors that WOULD be added:")
            for key in sorted(new_keys)[:50]:
                row = combined[key]
                print(f"  {row['vendor_name']} | {row['location']} | {row['vendor_type']} | {row['vendor_sub_type']} | intercompany={row['intercompany']}")
            if len(new_keys) > 50:
                print(f"  ... and {len(new_keys) - 50} more")
            return

        inserted = seed_missing_keyed_rows(db, VendorMapping, ("vendor_key",), combined)
        db.commit()
        print(f"\nInserted {inserted} new vendor mapping(s).")
    finally:
        db.close()


if __name__ == "__main__":
    main()
