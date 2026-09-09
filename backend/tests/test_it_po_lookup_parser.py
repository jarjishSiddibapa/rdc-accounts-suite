import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from app.services.it_po_lookup.bank_statement_parser import parse_bank_statement
from app.services.it_po_lookup.errors import StatementParseError


def _build_statement_xlsx(path: Path, *, account_number="00600200000050", from_date="28/08/2026", to_date="28/08/2026", rows=None) -> None:
    """Builds a minimal file matching the real export's proven shape: a
    metadata block located by label (not fixed position), then a header
    row starting with "Transaction Date"."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Account Number", account_number])
    ws.append([])
    ws.append(["Customer Name", "RDC CONCRETE INDIA  LIMITED"])
    ws.append(["Account Currency", "INR"])
    ws.append(["Opening Balance", -100.0])
    ws.append(["Closing Balance", -200.0])
    ws.append(["From Date", from_date])
    ws.append(["To Date", to_date])
    ws.append([])
    ws.append([
        "Transaction Date", "Transaction Description", "Transaction Amount",
        "Debit / Credit", "Reference No.", "Value Date", "Transaction Branch", "Running Balance",
    ])
    for row in rows or []:
        ws.append(row)
    wb.save(path)


class BankStatementParserTests(unittest.TestCase):
    def test_parses_metadata_and_rows_by_label_not_fixed_position(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "statement.xlsx"
            _build_statement_xlsx(
                path,
                rows=[
                    [
                        "28/08/2026 22:28:51",
                        "NEFT - UTIB0002904 - CX00000000VNGXT9 - 20270001444 - 917020079381137 -  HARRIER SECURITY",
                        35396.71, "D", "HDFCH01251322240", "28/08/2026", "CBX Internet System Mumbai", -74441149.89,
                    ],
                ],
            )
            parsed = parse_bank_statement(str(path))

        self.assertEqual(parsed.account_number, "00600200000050")
        self.assertEqual(parsed.from_date, date(2026, 8, 28))
        self.assertEqual(parsed.to_date, date(2026, 8, 28))
        self.assertEqual(len(parsed.rows), 1)
        row = parsed.rows[0]
        self.assertEqual(row.transaction_date, datetime(2026, 8, 28, 22, 28, 51))
        self.assertIn("20270001444", row.transaction_description)
        self.assertEqual(row.transaction_amount, Decimal("35396.71"))
        self.assertEqual(row.debit_credit, "D")
        self.assertEqual(row.reference_no, "HDFCH01251322240")
        self.assertEqual(row.value_date, date(2026, 8, 28))
        self.assertEqual(row.running_balance, Decimal("-74441149.89"))

    def test_column_reordering_is_tolerated_since_headers_are_matched_by_name(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Account Number", "111"])
        ws.append(["From Date", "01/09/2026"])
        ws.append(["To Date", "01/09/2026"])
        # Amount and Description swapped relative to the real export.
        ws.append(["Transaction Date", "Transaction Amount", "Transaction Description", "Reference No."])
        ws.append(["01/09/2026 10:00:00", 500.0, "Some narration", "REF123"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "statement.xlsx"
            wb.save(path)
            parsed = parse_bank_statement(str(path))

        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.rows[0].transaction_amount, Decimal("500"))
        self.assertEqual(parsed.rows[0].transaction_description, "Some narration")
        self.assertEqual(parsed.rows[0].reference_no, "REF123")

    def test_rejects_a_file_with_no_recognizable_header_row(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["This is not a bank statement"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "not_a_statement.xlsx"
            wb.save(path)
            with self.assertRaises(StatementParseError):
                parse_bank_statement(str(path))

    def test_rejects_a_header_row_missing_required_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Transaction Date", "Something Else"])
        ws.append(["01/09/2026", "value"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "incomplete.xlsx"
            wb.save(path)
            with self.assertRaises(StatementParseError):
                parse_bank_statement(str(path))

    def test_a_row_missing_amount_is_skipped_not_fatal(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "statement.xlsx"
            _build_statement_xlsx(
                path,
                rows=[
                    ["28/08/2026 10:00:00", "Good row", 100.0, "C", "REF1", "28/08/2026", "BR", -1.0],
                    ["28/08/2026 11:00:00", "Bad row - no amount", None, "C", "REF2", "28/08/2026", "BR", -1.0],
                ],
            )
            parsed = parse_bank_statement(str(path))
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.rows[0].reference_no, "REF1")


if __name__ == "__main__":
    unittest.main()
