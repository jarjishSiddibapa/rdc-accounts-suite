"""Coverage for the Untagged Invoices Report Generator's processing
pipeline: header/row cleanup, Location mapping, ageing-bucket recompute
(from an as-on date, not the source file's own frozen buckets, and not
duplicated when the source file's own bucket columns don't exactly match
this module's bucket names), the two Ageing/Summary pairs (Below 1k,
Untagged - both restricted to Type = Transactions, with the Type column
kept in their output), and the 5-sheet workbook's shape."""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from app.services.untagged_invoices.processor import (
    BUCKET_COLS,
    COUNT_COL,
    LOCATION_COL,
    _bucket_index,
    add_location_column,
    build_below_1k_ageing,
    build_below_1k_summary,
    build_untagged_ageing,
    build_untagged_summary,
    missing_incharge_mappings,
    missing_location_mappings,
    read_ageing_file,
    recompute_ageing_buckets,
    write_report,
)

_TYPE_COL = "Type (Transaction/Receipt)"

_HEADERS = [
    "Customer Name", "Customer Account", "Customer Group Name", "Location Number",
    "Location Name", "Sales Person Name", "Credit Limit", "Payment Terms",
    "Invoice No/ Receipt No", "Invoice/ Receipt Date", "Type (Transaction/Receipt)",
    "GL Date", "Invoice Type", "Accounted Invoice", "Accounted Outstanding",
    "Sum of On Account/Prepayment Amount", "Sum of Unapplied Amount", "Days O/S",
    *BUCKET_COLS, "Receivable Account",
]


class _LogQueue:
    def __init__(self):
        self.messages = []

    def put(self, item):
        self.messages.append(item)


def _make_ageing_workbook(path, data_rows, include_grand_total=True):
    """Build a minimal Oracle-shaped Aging export: 13 metadata rows, the
    real header at row 14, then data rows, an optional Grand Total row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aging"
    for r in range(1, 14):
        ws.cell(row=r, column=1, value=f"meta {r}")
    for c, h in enumerate(_HEADERS, 1):
        ws.cell(row=14, column=c, value=h)
    row_num = 15
    for row in data_rows:
        for c, h in enumerate(_HEADERS, 1):
            ws.cell(row=row_num, column=c, value=row.get(h))
        row_num += 1
    ws.cell(row=row_num, column=1, value=None)  # blank separator row
    row_num += 1
    if include_grand_total:
        ws.cell(row=row_num, column=1, value="Grand Total:")
        ws.cell(row=row_num, column=15, value=999999.99)
    wb.save(path)


def _base_row(**overrides):
    row = {
        "Customer Name": "Acme Co",
        "Customer Account": 1001,
        "Location Name": "Coimbatore - Saravanampatti",
        "Invoice/ Receipt Date": dt.datetime(2026, 1, 1),
        "Type (Transaction/Receipt)": "Transactions",
        "Accounted Outstanding": 500.0,
    }
    row.update(overrides)
    return row


class BucketIndexTests(unittest.TestCase):
    def test_boundaries(self):
        self.assertEqual(_bucket_index(0), 0)
        self.assertEqual(_bucket_index(15), 0)
        self.assertEqual(_bucket_index(16), 1)
        self.assertEqual(_bucket_index(30), 1)
        self.assertEqual(_bucket_index(31), 2)
        self.assertEqual(_bucket_index(60), 2)
        self.assertEqual(_bucket_index(61), 3)
        self.assertEqual(_bucket_index(360), 7)
        self.assertEqual(_bucket_index(361), 8)
        self.assertEqual(_bucket_index(10_000), 8)


class ReadAgeingFileTests(unittest.TestCase):
    def test_drops_grand_total_and_blank_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ageing.xlsx"
            _make_ageing_workbook(path, [_base_row(), _base_row(**{"Customer Name": "Beta Ltd"})])
            df = read_ageing_file(str(path), _LogQueue())
            self.assertEqual(len(df), 2)
            self.assertNotIn("Grand Total:", df["Customer Name"].tolist())

    def test_empty_file_raises(self):
        from app.jobs import JobUserError
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ageing.xlsx"
            _make_ageing_workbook(path, [])
            with self.assertRaises(JobUserError):
                read_ageing_file(str(path), _LogQueue())


class LocationMappingTests(unittest.TestCase):
    def test_add_location_column_after_location_name(self):
        df = pd.DataFrame([_base_row(), _base_row(**{"Location Name": "unmapped place"})])
        location_map = {"COIMBATORE - SARAVANAMPATTI": "COIMBATORE+TRICHY"}
        out = add_location_column(df, location_map)
        cols = list(out.columns)
        self.assertEqual(cols.index(LOCATION_COL), cols.index("Location Name") + 1)
        self.assertEqual(out[LOCATION_COL].tolist(), ["COIMBATORE+TRICHY", ""])

    def test_missing_location_and_incharge_mappings(self):
        df = pd.DataFrame([_base_row(), _base_row(**{"Location Name": "unmapped place"})])
        location_map = {"COIMBATORE - SARAVANAMPATTI": "COIMBATORE+TRICHY"}
        missing = missing_location_mappings(df, location_map)
        self.assertEqual(missing, ["UNMAPPED PLACE"])

        out = add_location_column(df, location_map)
        missing_incharge = missing_incharge_mappings(out[LOCATION_COL].tolist(), {})
        self.assertIn("COIMBATORE+TRICHY", missing_incharge)
        # blank/unmapped Location never shows up as an "incharge" gap of its own
        self.assertNotIn("", missing_incharge)


class RecomputeAgeingBucketsTests(unittest.TestCase):
    def test_places_amount_in_the_correct_single_bucket(self):
        df = pd.DataFrame([_base_row(
            **{"Invoice/ Receipt Date": dt.datetime(2026, 1, 1), "Accounted Outstanding": 16812.05},
        )])
        as_on = dt.date(2026, 1, 1) + dt.timedelta(days=655)  # -> bucket "360+ Days"
        out = recompute_ageing_buckets(df, as_on, _LogQueue())
        self.assertEqual(out["Days O/S"].iloc[0], 655)
        for col in BUCKET_COLS[:-1]:
            self.assertEqual(out[col].iloc[0], 0.0)
        self.assertEqual(out["360+ Days"].iloc[0], 16812.05)

    def test_unparseable_date_leaves_buckets_blank(self):
        df = pd.DataFrame([_base_row(**{"Invoice/ Receipt Date": "not a date"})])
        out = recompute_ageing_buckets(df, dt.date(2026, 1, 1), _LogQueue())
        self.assertTrue(pd.isna(out["Days O/S"].iloc[0]))
        self.assertTrue(pd.isna(out[BUCKET_COLS[0]].iloc[0]))

    def test_stale_source_bucket_columns_are_replaced_not_duplicated(self):
        """Regression test for a real production bug: the ERP export's own
        frozen bucket columns must be dropped before the new ones are
        computed, not left behind alongside them under a near-identical
        name. This is exactly what happened when the source file's
        "361-999999 Days " (trailing space) didn't exact-match this
        module's own bucket name after header-stripping - the sheet showed
        that bucket twice."""
        row = _base_row(**{"Invoice/ Receipt Date": dt.datetime(2026, 1, 1)})
        row["361-999999 Days"] = 999.0  # stale raw ERP column, old naming
        df = pd.DataFrame([row])
        out = recompute_ageing_buckets(df, dt.date(2026, 1, 31), _LogQueue())
        self.assertNotIn("361-999999 Days", out.columns)
        self.assertEqual(out.columns.nunique(), len(out.columns))
        for col in BUCKET_COLS:
            self.assertIn(col, out.columns)


class PivotTests(unittest.TestCase):
    def _prepared(self, rows, as_on):
        df = pd.DataFrame(rows)
        df = add_location_column(df, {"LOC A": "Location A", "LOC B": "Location B"})
        return recompute_ageing_buckets(df, as_on, _LogQueue())

    def test_below_1k_ageing_filters_by_type_and_amount_and_groups_by_location(self):
        as_on = dt.date(2026, 1, 31)  # 30 days after invoice date -> bucket 16-30
        rows = [
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 500.0}),
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 250.0}),
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 5000.0}),  # excluded: >= 1000
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Receipts",
                         "Accounted Outstanding": 300.0}),  # excluded: not a Transaction
            _base_row(**{"Location Name": "LOC B", "Accounted Outstanding": -10.0}),   # excluded: not > 0
        ]
        df = self._prepared(rows, as_on)
        ageing = build_below_1k_ageing(df)
        self.assertEqual(len(ageing), 2)  # the >=1000, Receipts, and negative rows are excluded
        self.assertIn(_TYPE_COL, ageing.columns)
        pivot = build_below_1k_summary(ageing, {"Location A": "Alice"})

        self.assertEqual(list(pivot[LOCATION_COL]), ["Location A"])
        self.assertAlmostEqual(pivot["Total O/S"].iloc[0], 750.0)
        self.assertEqual(pivot[COUNT_COL].iloc[0], 2)
        self.assertEqual(
            list(pivot.columns[:4]), [LOCATION_COL, "Account Incharges", COUNT_COL, "Total O/S"],
        )
        self.assertEqual(pivot["Account Incharges"].iloc[0], "Alice")
        self.assertAlmostEqual(
            pivot["Above 30 days"].iloc[0],
            pivot["Total O/S"].iloc[0] - pivot[BUCKET_COLS[0]].iloc[0] - pivot[BUCKET_COLS[1]].iloc[0],
        )

    def test_below_1k_summary_sorted_by_total_os_descending(self):
        as_on = dt.date(2026, 1, 31)
        rows = [
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 100.0}),
            _base_row(**{"Location Name": "LOC B", "Accounted Outstanding": 900.0}),
        ]
        df = self._prepared(rows, as_on)
        ageing = build_below_1k_ageing(df)
        pivot = build_below_1k_summary(ageing, {})
        self.assertEqual(list(pivot[LOCATION_COL]), ["Location B", "Location A"])

    def test_untagged_ageing_and_summary(self):
        as_on = dt.date(2026, 1, 31)
        rows = [
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Transactions",
                         "Accounted Outstanding": -1000.0}),
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Receipts",
                         "Accounted Outstanding": -500.0}),   # excluded: not a Transaction
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Transactions",
                         "Accounted Outstanding": 50.0}),      # excluded: outstanding not < 0
        ]
        df = self._prepared(rows, as_on)
        ageing = build_untagged_ageing(df)
        self.assertEqual(len(ageing), 1)
        self.assertEqual(ageing["Accounted Outstanding"].iloc[0], -1000.0)
        self.assertIn(_TYPE_COL, ageing.columns)

        summary = build_untagged_summary(ageing, {"Location A": "Alice"})
        self.assertEqual(list(summary[LOCATION_COL]), ["Location A"])
        # Untagged Summary is in Lakhs (divide by 1,00,000)
        self.assertAlmostEqual(summary["Accounted Outstanding"].iloc[0], -1000.0 / 100_000)
        self.assertEqual(summary[COUNT_COL].iloc[0], 1)  # a count, not scaled to Lakhs
        self.assertEqual(
            list(summary.columns[:4]),
            [LOCATION_COL, "Account Incharges", COUNT_COL, "Accounted Outstanding"],
        )
        self.assertAlmostEqual(
            summary["Above 30 Days"].iloc[0],
            summary["Accounted Outstanding"].iloc[0] - summary[BUCKET_COLS[0]].iloc[0] - summary[BUCKET_COLS[1]].iloc[0],
        )

    def test_untagged_summary_sorted_by_accounted_outstanding_ascending(self):
        as_on = dt.date(2026, 1, 31)
        rows = [
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Transactions",
                         "Accounted Outstanding": -100.0}),
            _base_row(**{"Location Name": "LOC B", _TYPE_COL: "Transactions",
                         "Accounted Outstanding": -900.0}),
        ]
        df = self._prepared(rows, as_on)
        ageing = build_untagged_ageing(df)
        summary = build_untagged_summary(ageing, {})
        self.assertEqual(list(summary[LOCATION_COL]), ["Location B", "Location A"])


class ErpFrameTests(unittest.TestCase):
    _COLUMNS = [
        "CUSTOMER_NAME", "CUSTOMER_ACCOUNT", "LOCATION", "TRX_NUMBER", "INVOICE_RCPT_CLASS",
        "CREDITLIMIT", "PAYMENT_TERM", "OUTSTANDING_AMOUNT", "UNAPPLIED_AMOUNT", "DUE_DATE",
        "OVER_DUE_DAYS", "SALESPERSON",
    ]

    def _rows(self):
        d = dt.datetime(2026, 1, 1)
        return [
            ("Acme", "1001", " Pune ", "INV1", "Invoice", None, 30, 500, 0, d, 10, "Ravi"),
            ("Acme", "1001", "Pune", "CM1", "Credit Memo", None, 30, -250.5, 0, d, 10, None),
            ("Acme", "1001", None, "RCT1", "Receipts", None, None, -90, 0, d, 10, None),
        ]

    def test_columns_match_what_the_pipeline_reads(self):
        from app.services.erp_ageing import erp_frame as _erp_frame
        df = _erp_frame(self._COLUMNS, self._rows())
        for needed in ("Customer Name", "Location Name", "Invoice/ Receipt Date",
                       "Accounted Outstanding", _TYPE_COL):
            self.assertIn(needed, df.columns)
        self.assertEqual(df["Location Name"].tolist(), ["Pune", "Pune", ""])

    def test_receipts_are_receipts_and_everything_else_is_a_transaction(self):
        from app.services.erp_ageing import erp_frame as _erp_frame
        df = _erp_frame(self._COLUMNS, self._rows())
        self.assertEqual(df[_TYPE_COL].tolist(), ["Transactions", "Transactions", "Receipts"])

    def test_erp_frame_runs_through_the_rest_of_the_pipeline(self):
        from app.services.erp_ageing import erp_frame as _erp_frame
        df = _erp_frame(self._COLUMNS, self._rows())
        df = add_location_column(df, {"PUNE": "Pune"})
        df = recompute_ageing_buckets(df, dt.date(2026, 1, 31), _LogQueue())
        self.assertEqual(df["Days O/S"].tolist(), [30, 30, 30])
        untagged = build_untagged_ageing(df)
        self.assertEqual(untagged["Invoice No/ Receipt No"].tolist(), ["CM1"])  # the receipt is not untagged
        self.assertEqual(len(build_below_1k_ageing(df)), 1)


class TinyAmountTests(unittest.TestCase):
    def test_an_amount_too_small_to_show_is_written_as_zero_not_minus_zero(self):
        as_on = dt.date(2026, 1, 31)
        # -2 rupees = -0.00002 Lakhs: displays as -0.00 unless treated as zero
        rows = [_base_row(**{"Location Name": "LOC A", "Accounted Outstanding": -2.0,
                             "Invoice/ Receipt Date": dt.datetime(2026, 1, 30)}),     # 0-15 Days
                _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": -500000.0})]  # 16-30 Days
        df = add_location_column(pd.DataFrame(rows), {"LOC A": "Location A"})
        df = recompute_ageing_buckets(df, as_on, _LogQueue())
        untagged_ageing = build_untagged_ageing(df)
        untagged_summary = build_untagged_summary(untagged_ageing, {})
        below_ageing = build_below_1k_ageing(df)
        below_summary = build_below_1k_summary(below_ageing, {})
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "o.xlsx"
            write_report(df, untagged_summary, untagged_ageing, below_summary, below_ageing, out, as_on,
                         log_q=_LogQueue())
            ws = openpyxl.load_workbook(out)["Untagged Summary"]
            header = [c.value for c in ws[1]]
            first = {h: c.value for h, c in zip(header, ws[2])}
            self.assertEqual(first["0-15 Days"], 0.0)            # the tiny 0-15 amount stays out
            self.assertEqual(first["16-30 Days"], -5.0)           # real amounts are untouched


class SheetLayoutTests(unittest.TestCase):
    """Widths and the frozen header have to be in the file, or Excel shows
    ####### for numbers and the header scrolls away."""

    def test_every_sheet_has_column_widths_and_a_frozen_header(self):
        as_on = dt.date(2026, 1, 31)
        rows = [_base_row(**{"Location Name": "LOC A", "Accounted Outstanding": -776767.06}),
                _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 500.0})]
        df = add_location_column(pd.DataFrame(rows), {"LOC A": "Location A"})
        df = recompute_ageing_buckets(df, as_on, _LogQueue())
        ua = build_untagged_ageing(df)
        us = build_untagged_summary(ua, {})
        ba = build_below_1k_ageing(df)
        bs = build_below_1k_summary(ba, {})
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "o.xlsx"
            write_report(df, us, ua, bs, ba, out, as_on, log_q=_LogQueue())
            wb = openpyxl.load_workbook(out)
            for ws in wb.worksheets:
                self.assertEqual(ws.freeze_panes, "A2", ws.title)
                self.assertGreaterEqual(ws.column_dimensions["D"].width, 10, ws.title)
            ageing = wb["Ageing"]
            header = [c.value for c in ageing[1]]
            amount = ageing.column_dimensions[openpyxl.utils.get_column_letter(header.index("Accounted Outstanding") + 1)]
            self.assertGreaterEqual(amount.width, len("Accounted Outstanding"))   # wide enough for "-776,767.06"


class NumericColumnTests(unittest.TestCase):
    def test_identifier_columns_are_not_amounts(self):
        from app.services.untagged_invoices.processor import _is_numeric_col
        self.assertFalse(_is_numeric_col("Invoice No/ Receipt No"))
        self.assertFalse(_is_numeric_col("Invoice Type"))
        self.assertTrue(_is_numeric_col("Accounted Outstanding"))
        self.assertTrue(_is_numeric_col("Credit Limit"))


class DashForEmptyAndZeroTests(unittest.TestCase):
    def test_zero_format_and_blank_cells_show_a_dash(self):
        as_on = dt.date(2026, 1, 31)
        rows = [
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 500.0}),
            _base_row(**{"Location Name": "UNKNOWN", "Accounted Outstanding": 700.0}),  # no Location mapped
        ]
        df = add_location_column(pd.DataFrame(rows), {"LOC A": "Location A"})
        df = recompute_ageing_buckets(df, as_on, _LogQueue())
        below_1k_ageing = build_below_1k_ageing(df)
        below_1k_summary = build_below_1k_summary(below_1k_ageing, {})
        untagged_ageing = build_untagged_ageing(df)
        untagged_summary = build_untagged_summary(untagged_ageing, {})

        with tempfile.TemporaryDirectory() as tmp:
            out_path = Path(tmp) / "out.xlsx"
            write_report(df, untagged_summary, untagged_ageing, below_1k_summary, below_1k_ageing,
                         out_path, as_on, log_q=_LogQueue())
            wb = openpyxl.load_workbook(out_path)

            ws = wb["Ageing"]
            header = [c.value for c in ws[1]]
            rows_out = [{h: c for h, c in zip(header, row)} for row in ws.iter_rows(min_row=2)]
            unmapped = next(r for r in rows_out if r["Location Name"].value == "UNKNOWN")
            self.assertEqual(unmapped[LOCATION_COL].value, "-")     # empty text
            # a zero bucket keeps its number but its format turns it into a dash
            zero_bucket = unmapped["0-15 Days"]
            self.assertEqual(zero_bucket.value, 0.0)
            self.assertIn('"-"', zero_bucket.number_format)

            pivot = wb["Below 1k Summary"]
            first = [c.value for c in pivot[2]]
            self.assertEqual(first[1], "-")   # Account Incharges not set
            self.assertIn('"-"', pivot.cell(row=2, column=5).number_format)
            self.assertTrue(str(pivot.cell(row=2, column=pivot.max_column).value).startswith("=ROUND("))


class WriteReportTests(unittest.TestCase):
    def test_sheet_order_and_grand_total_formula(self):
        as_on = dt.date(2026, 1, 31)
        rows = [
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": -1000.0}),  # untagged
            _base_row(**{"Location Name": "LOC A", "Accounted Outstanding": 500.0}),    # below 1k
            _base_row(**{"Location Name": "LOC A", _TYPE_COL: "Receipts",
                         "Accounted Outstanding": -2000.0}),  # Receipts: in Ageing only
        ]
        df = pd.DataFrame(rows)
        df = add_location_column(df, {"LOC A": "Location A"})
        df = recompute_ageing_buckets(df, as_on, _LogQueue())
        below_1k_ageing = build_below_1k_ageing(df)
        below_1k_summary = build_below_1k_summary(below_1k_ageing, {"Location A": "Alice"})
        untagged_ageing = build_untagged_ageing(df)
        untagged_summary = build_untagged_summary(untagged_ageing, {"Location A": "Alice"})

        with tempfile.TemporaryDirectory() as tmp:
            out_path = Path(tmp) / "out.xlsx"
            write_report(df, untagged_summary, untagged_ageing, below_1k_summary, below_1k_ageing,
                         out_path, as_on, log_q=_LogQueue())

            wb = openpyxl.load_workbook(out_path)
            self.assertEqual(
                wb.sheetnames,
                ["Untagged Summary", "Untagged Ageing", "Below 1k Summary", "Below 1k Ageing",
                 "Ageing"],
            )

            below_1k_summary_ws = wb["Below 1k Summary"]
            grand_total_row = below_1k_summary_ws.max_row
            self.assertEqual(below_1k_summary_ws.cell(row=grand_total_row, column=1).value, "Grand Total")
            self.assertEqual(
                [c.value for c in below_1k_summary_ws[1]][:4],
                ["Location Name", "Account Incharges", "Count of Invoices", "Total O/S"],
            )
            count_cell = below_1k_summary_ws.cell(row=grand_total_row, column=3).value
            self.assertTrue(str(count_cell).startswith("=SUBTOTAL(9,C"))
            total_cell = below_1k_summary_ws.cell(row=grand_total_row, column=4).value
            self.assertTrue(str(total_cell).startswith("=SUBTOTAL(9,D"))
            self.assertEqual(below_1k_summary_ws.cell(row=2, column=3).value, 1)

            summary_ws = wb["Untagged Summary"]
            # Untagged Summary carries a "values are in Lakhs" footnote two
            # rows below the Grand Total row, so max_row is no longer the
            # Grand Total row itself.
            summary_grand_row = summary_ws.max_row - 2
            self.assertEqual(summary_ws.cell(row=summary_grand_row, column=1).value, "Grand Total")
            self.assertEqual(
                [c.value for c in summary_ws[1]][:4],
                ["Location Name", "Account Incharges", "Count of Invoices", "Accounted Outstanding"],
            )
            summary_total_cell = summary_ws.cell(row=summary_grand_row, column=4).value
            self.assertTrue(str(summary_total_cell).startswith("=SUBTOTAL(9,D"))
            self.assertEqual(summary_ws.cell(row=2, column=3).value, 1)
            footnote_cell = summary_ws.cell(row=summary_ws.max_row, column=1).value
            self.assertIn("Lakhs", str(footnote_cell))

            below_1k_ws = wb["Below 1k Ageing"]
            self.assertEqual(below_1k_ws.max_row - 1, len(below_1k_ageing))

            # Type column kept on every Ageing sub-sheet, including the raw
            # "Ageing" sheet, which also keeps the Receipts row.
            untagged_header = [c.value for c in wb["Untagged Ageing"][1]]
            self.assertIn(_TYPE_COL, untagged_header)
            below_1k_header = [c.value for c in wb["Below 1k Ageing"][1]]
            self.assertIn(_TYPE_COL, below_1k_header)

            ageing_ws = wb["Ageing"]
            ageing_header = [c.value for c in ageing_ws[1]]
            self.assertIn(_TYPE_COL, ageing_header)
            self.assertEqual(ageing_ws.max_row - 1, len(df))


if __name__ == "__main__":
    unittest.main()

