"""State-wise Summary sheet of the combined GSTR-2B workbook."""
import unittest

from openpyxl import Workbook

from app.services.gstr2b.summary import write_summary_sheet

B2B = ["State Code", "State Name", "Month & Year", "GSTIN of supplier", "Trade/Legal name",
       "Invoice Details - Invoice number", "Invoice Details - Invoice type", "Invoice Details - Invoice Date",
       "Invoice Details - Invoice Value(₹)", "Place of supply", "Supply Attract Reverse Charge",
       "Taxable Value (₹)", "Tax Amount - Integrated Tax(₹)", "Tax Amount - Central Tax(₹)",
       "Tax Amount - State/UT Tax(₹)", "Tax Amount - Cess(₹)", "GSTR-1/1A/IFF/GSTR-5 Period",
       "GSTR-1/1A/IFF/GSTR-5 Filing Date", "ITC Availability"]
CDNR = ["State Code", "State Name", "Month & Year", "GSTIN of supplier", "Trade/Legal name",
        "Credit note/Debit note details - Note number", "Credit note/Debit note details - Note type",
        "Credit note/Debit note details - Note Supply type", "Credit note/Debit note details - Note date",
        "Credit note/Debit note details - Note Value (₹)", "Place of supply", "Supply Attract Reverse Charge",
        "Taxable Value (₹)", "Tax Amount - Integrated Tax(₹)", "Tax Amount - Central Tax(₹)",
        "Tax Amount - State/UT Tax(₹)"]
IMPG = ["State Code", "State Name", "Month & Year", "Icegate Reference Date", "Port Code",
        "Bill of Entry Details - Number", "Bill of Entry Details - Date", "Bill of Entry Details - Taxable Value",
        "Amount of tax (₹) - Integrated Tax(₹)", "Amount of tax (₹) - Cess(₹)"]


def _summary(states):
    wb = Workbook()
    write_summary_sheet(wb, states, {"B2B": B2B, "B2B-CDNR": CDNR, "IMPG": IMPG})
    return wb


class SummarySheetTests(unittest.TestCase):
    def test_summary_is_first_sheet_with_states_sorted_including_data_less_ones(self):
        wb = _summary(["Gujarat", "Kerala", "Andhra Pradesh (New)", "Gujarat"])
        ws = wb["Summary"]
        self.assertEqual(wb.sheetnames[0], "Summary")
        self.assertEqual([ws.cell(row=r, column=2).value for r in range(4, 8)],
                         ["Andhra Pradesh (New)", "Gujarat", "Kerala", "Total"])

    def test_headers_match_the_team_layout(self):
        ws = _summary(["Goa"])["Summary"]
        self.assertEqual(ws["B3"].value, "Particulars")
        self.assertEqual([ws[c].value for c in ("C2", "F2", "I2", "L2", "O2", "P2", "S2", "V2", "Y2")],
                         ["B2B (Non-RCM)", "B2B (RCM)", "B2B (Blocked)", "CDNR (Credit Notes - Debit Notes)",
                          "Imports", "Manual RCM", "Net", "Net (Consultants)", "Difference"])
        self.assertEqual([ws.cell(row=3, column=c).value for c in range(3, 28)],
                         ["IGST", "CGST", "SGST"] * 4 + ["IGST"] + ["IGST", "CGST", "SGST"] * 4)

    def test_row_formulas_follow_the_manual_summary(self):
        ws = _summary(["Goa"])["Summary"]
        self.assertEqual(ws["C4"].value, '=SUMIFS(B2B!$M:$M,B2B!$B:$B,Summary!$B4,B2B!$K:$K,"no")')
        self.assertEqual(ws["H4"].value, '=SUMIFS(B2B!$O:$O,B2B!$B:$B,Summary!$B4,B2B!$K:$K,"Yes")')
        self.assertEqual(ws["I4"].value, '=SUMIFS(B2B!$M:$M,B2B!$B:$B,Summary!$B4,B2B!$S:$S,"no")')
        self.assertEqual(ws["O4"].value, "=SUMIF(IMPG!$B:$B,Summary!$B4,IMPG!$I:$I)")
        self.assertEqual(ws["S4"].value, "=C4+F4-I4-L4+O4+P4")
        self.assertEqual(ws["T4"].value, "=D4+G4-J4-M4+Q4")
        self.assertEqual(ws["U4"].value, "=E4+H4-K4-N4+R4")
        # credit notes less debit notes, reverse-charge "no" only
        self.assertIn("'B2B-CDNR'!$N:$N", ws["L4"].value)
        self.assertIn('"Debit Note"', ws["L4"].value)

    def test_manual_rcm_and_consultants_are_empty_and_difference_waits_for_them(self):
        ws = _summary(["Goa", "Gujarat"])["Summary"]
        for r in (4, 5):
            for col in ("P", "Q", "R", "V", "W", "X"):
                self.assertIsNone(ws[f"{col}{r}"].value)
        self.assertEqual(ws["Y4"].value, '=IF(V4="","",S4-V4)')
        self.assertEqual(ws["AA5"].value, '=IF(X5="","",U5-X5)')

    def test_total_row_sums_every_state_row(self):
        ws = _summary(["A", "B", "C", "D"])["Summary"]
        self.assertEqual(ws["B8"].value, "Total")
        self.assertEqual(ws["C8"].value, "=SUM(C4:C7)")
        self.assertEqual(ws["X8"].value, "=SUM(X4:X7)")
        self.assertEqual(ws["Y8"].value, '=IF(COUNT(V4:V7)=0,"",S8-V8)')

    def test_empty_data_sheets_still_produce_valid_formulas(self):
        wb = Workbook()
        write_summary_sheet(wb, ["Goa"], {"B2B": ["State Code", "State Name", "Month & Year"]})
        ws = wb["Summary"]
        self.assertEqual(ws["C4"].value, "=0")
        self.assertEqual(ws["O4"].value, "=0")
        self.assertEqual(ws["S4"].value, "=C4+F4-I4-L4+O4+P4")


if __name__ == "__main__":
    unittest.main()
