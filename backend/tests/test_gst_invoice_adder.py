import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.services.gst_invoice_adder.processor import (
    COL_GST_NEW, COL_INV_DATE, COL_INV_NO,
    _bulk_query, _extract_pairs_vectorized, _fetch_batch_pooled,
    build_gst_report_pure_python,
)


class ExtractPairsVectorizedTests(unittest.TestCase):
    """A genuinely blank/NaN invoice-number cell must never become a pair
    sent toward Oracle - regression coverage for a real bug found live: on
    pandas 3.x, Series.astype(str) no longer stringifies NaN to "nan" the
    way it used to, so a plain ``.astype(str)...isin(_INVALID_INV)`` filter
    silently lets a real NaN cell through as a non-string value."""

    def test_a_real_nan_invoice_cell_is_excluded(self):
        df = pd.DataFrame(
            [
                ("1001", datetime(2026, 1, 1)),
                (float("nan"), datetime(2026, 1, 2)),
                ("", datetime(2026, 1, 3)),
                ("None", datetime(2026, 1, 4)),
            ],
            columns=[COL_INV_NO, COL_INV_DATE],
            dtype=object,
        )

        pairs = _extract_pairs_vectorized(df, COL_INV_NO, COL_INV_DATE)

        self.assertEqual(pairs, [("1001", "01-Jan-26")])


class BulkQueryTests(unittest.TestCase):
    def test_generates_one_numbered_bind_pair_per_input_pair(self):
        pairs = [("1001", "01-Jan-26"), ("1002", "02-Jan-26"), ("1003", "03-Jan-26")]
        query, binds = _bulk_query(pairs)
        self.assertEqual(
            binds,
            {
                "inv_no_0": "1001", "trx_date_0": "01-Jan-26",
                "inv_no_1": "1002", "trx_date_1": "02-Jan-26",
                "inv_no_2": "1003", "trx_date_2": "03-Jan-26",
            },
        )
        self.assertEqual(query.count("UNION ALL"), 2)
        # Every bound value only ever appears as a bind value, never
        # interpolated directly into the SQL text (no injection surface).
        for value in binds.values():
            self.assertNotIn(value, query)

    def test_single_pair_has_no_union(self):
        query, binds = _bulk_query([("1001", "01-Jan-26")])
        self.assertNotIn("UNION ALL", query)
        self.assertEqual(binds, {"inv_no_0": "1001", "trx_date_0": "01-Jan-26"})


class _FakeLogQueue:
    def __init__(self):
        self.messages = []

    def put(self, item):
        self.messages.append(item)


class BuildGstReportPureWriterTests(unittest.TestCase):
    """Regression coverage for the write-only rewrite: it must produce the
    exact same output shape (column order, values, formatting hooks,
    total/found/blank counts) as the previous cell-by-cell writer, while
    also correctly excluding the same invalid-invoice rows that
    _extract_pairs_vectorized already excludes from the Oracle fetch -
    otherwise a row that was never queried could still be counted/written
    as a legitimate "blank" result instead of being skipped entirely."""

    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)
        self.output_path = str(Path(self.tmpdir.name) / "out.xlsx")
        self.log_q = _FakeLogQueue()

    def _build_df(self, rows):
        return pd.DataFrame(rows, columns=["Customer Name", COL_INV_NO, COL_INV_DATE, "Amount"])

    def test_matched_and_unmatched_rows_with_invalid_rows_excluded(self):
        df = self._build_df([
            ("Acme Traders", "1001", datetime(2026, 1, 1), 100.0),
            ("Acme Traders", "1002", datetime(2026, 1, 2), 200.0),
            ("Acme Traders", float("nan"), datetime(2026, 1, 3), 300.0),  # invalid inv_no -> excluded
            ("Acme Traders", "None", datetime(2026, 1, 4), 400.0),  # invalid inv_no -> excluded
            ("Acme Traders", "NaN", datetime(2026, 1, 5), 500.0),  # invalid (case-insensitive) -> excluded
        ])
        gst_map = {("1001", "01-Jan-26"): "GST1001", ("1002", "02-Jan-26"): ""}

        total, found, blank = build_gst_report_pure_python(
            df, COL_INV_NO, COL_INV_DATE, gst_map, self.output_path, self.log_q,
        )

        self.assertEqual((total, found, blank), (2, 1, 1))

        wb = load_workbook(self.output_path)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        # GST column is inserted immediately before the invoice-number column.
        self.assertEqual(
            header,
            ["Customer Name", COL_GST_NEW, COL_INV_NO, COL_INV_DATE, "Amount"],
        )
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 2)
        self.assertEqual(data_rows[0], ("Acme Traders", "GST1001", "1001", datetime(2026, 1, 1), 100.0))
        self.assertEqual(data_rows[1], ("Acme Traders", None, "1002", datetime(2026, 1, 2), 200.0))
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, f"A1:E{len(data_rows) + 1}")

    def test_progress_callback_is_invoked_on_the_final_row(self):
        df = self._build_df([("Acme Traders", "1001", datetime(2026, 1, 1), 100.0)])
        calls = []

        build_gst_report_pure_python(
            df, COL_INV_NO, COL_INV_DATE, {}, self.output_path, self.log_q,
            progress_cb=lambda pct, msg: calls.append((pct, msg)),
        )

        self.assertTrue(calls)
        last_pct, last_msg = calls[-1]
        self.assertIn("1 / 1", last_msg)
        self.assertGreater(last_pct, 0.92)

    def test_all_rows_invalid_produces_an_empty_but_valid_workbook(self):
        df = self._build_df([
            ("Acme Traders", float("nan"), datetime(2026, 1, 1), 100.0),
            ("Acme Traders", "", datetime(2026, 1, 2), 200.0),
        ])

        total, found, blank = build_gst_report_pure_python(
            df, COL_INV_NO, COL_INV_DATE, {}, self.output_path, self.log_q,
        )

        self.assertEqual((total, found, blank), (0, 0, 0))
        wb = load_workbook(self.output_path)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, ["Customer Name", COL_GST_NEW, COL_INV_NO, COL_INV_DATE, "Amount"])
        self.assertEqual(list(ws.iter_rows(min_row=2)), [])
        # No auto-filter range is set over zero data rows.
        self.assertIsNone(ws.auto_filter.ref)

    def test_date_column_gets_a_date_number_format(self):
        df = self._build_df([("Acme Traders", "1001", datetime(2026, 1, 1), 100.0)])

        build_gst_report_pure_python(
            df, COL_INV_NO, COL_INV_DATE, {}, self.output_path, self.log_q,
        )

        wb = load_workbook(self.output_path)
        ws = wb.active
        date_cell = ws.cell(row=2, column=4)  # Invoice/ Receipt Date column
        self.assertEqual(date_cell.number_format, "DD-MMM-YYYY")


class _FakeCursor:
    """Distinguishes the bulk call (positional binds dict) from the
    fallback per-pair call (inv_no=/trx_date= kwargs, the real GST_QUERY's
    own calling convention) exactly as _fetch_batch_pooled issues them."""

    def __init__(self, bulk_should_fail, gst_by_pair):
        self.bulk_should_fail = bulk_should_fail
        self.gst_by_pair = gst_by_pair
        self.prefetchrows = None
        self._bulk_binds = None
        self._last_pair = None

    def execute(self, sql, binds=None, **kwargs):
        if kwargs:
            self._last_pair = (kwargs["inv_no"], kwargs["trx_date"])
            return
        if self.bulk_should_fail:
            raise RuntimeError("simulated bulk failure")
        self._bulk_binds = binds

    def fetchall(self):
        n = len(self._bulk_binds) // 2
        rows = []
        for i in range(n):
            pair = (self._bulk_binds[f"inv_no_{i}"], self._bulk_binds[f"trx_date_{i}"])
            rows.append((pair[0], pair[1], self.gst_by_pair.get(pair) or None))
        return rows

    def fetchone(self):
        gst = self.gst_by_pair.get(self._last_pair)
        return (gst,) if gst else None

    def close(self):
        pass


class _FakeAcquireContext:
    def __init__(self, conn):
        self._conn = conn

    def __enter__(self):
        return self._conn

    def __exit__(self, *exc_info):
        return False


class _FakeConnection:
    def __init__(self, cursor):
        self._cursor = cursor

    def cursor(self):
        return self._cursor


class _FakePool:
    def __init__(self, cursor):
        self._cursor = cursor

    def acquire(self):
        return _FakeAcquireContext(_FakeConnection(self._cursor))


class FetchBatchPooledFallbackTests(unittest.TestCase):
    """The bulk query is new; one malformed pair or a transient Oracle
    hiccup on the UNION ALL statement must never blank an entire batch -
    it should fall back to the proven one-pair-per-query path instead."""

    def test_successful_bulk_query_returns_every_pair(self):
        cursor = _FakeCursor(bulk_should_fail=False, gst_by_pair={("1001", "01-Jan-26"): "GST1"})
        log_q = _FakeLogQueue()

        result = _fetch_batch_pooled(
            [("1001", "01-Jan-26"), ("1002", "02-Jan-26")], _FakePool(cursor), log_q,
        )

        self.assertEqual(result, {("1001", "01-Jan-26"): "GST1", ("1002", "02-Jan-26"): ""})

    def test_bulk_failure_falls_back_to_per_pair_queries(self):
        cursor = _FakeCursor(bulk_should_fail=True, gst_by_pair={("1001", "01-Jan-26"): "GST1"})
        log_q = _FakeLogQueue()

        result = _fetch_batch_pooled(
            [("1001", "01-Jan-26"), ("1002", "02-Jan-26")], _FakePool(cursor), log_q,
        )

        self.assertEqual(result, {("1001", "01-Jan-26"): "GST1", ("1002", "02-Jan-26"): ""})
        self.assertTrue(
            any("retrying them individually" in str(message) for _level, message in log_q.messages)
        )


if __name__ == "__main__":
    unittest.main()
