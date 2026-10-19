"""The compact copy of a workbook must hold exactly the same data and look."""

import datetime as dt
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services import xlsx_compact


def _workbook(path, rows=500, gaps=True):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["Name", "Amount", "Date", "Note", "Formula"])
    for i in range(1, rows + 1):
        ws.append([f"Customer {i % 37}", i * 1.25, dt.datetime(2026, 1, 1) + dt.timedelta(days=i), (None if gaps and i % 5 == 0 else "-"), f"=B{i + 1}*2"])
    other = wb.create_sheet("Second")
    other.append(["a", "b"])
    other.append([1, 2])
    wb.save(path)


def _cells(path):
    wb = openpyxl.load_workbook(path)
    return {
        ws.title: [[(c.value, c.number_format) for c in row] for row in ws.iter_rows()]
        for ws in wb.worksheets
    }


class CompactTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.src = Path(self.tmp.name) / "a.xlsx"
        self.dst = Path(self.tmp.name) / "b.xlsx"

    def test_same_values_and_formats_in_a_smaller_file(self):
        _workbook(self.src, gaps=False)
        xlsx_compact.compact_xlsx(self.src, self.dst)
        self.assertEqual(_cells(self.src), _cells(self.dst))
        self.assertLess(self.dst.stat().st_size, self.src.stat().st_size)

    def test_cell_addresses_are_gone_but_rows_keep_theirs(self):
        _workbook(self.src, rows=20, gaps=False)
        xlsx_compact.compact_xlsx(self.src, self.dst)
        with zipfile.ZipFile(self.dst) as z:
            sheet = z.read("xl/worksheets/sheet1.xml")
        self.assertNotIn(b'<c r="', sheet)
        self.assertIn(b'<row r="2"', sheet)

    def test_a_row_with_an_empty_cell_in_the_middle_keeps_its_addresses(self):
        # the empty cell is simply absent from the file; without addresses the
        # cells after it would slide one column to the left
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("G")
        ws.append(["a", None, "c", "d"])
        ws.append(["e", "f", "g", "h"])
        ws.append([None, "j", "k", "l"])
        wb.save(self.src)
        xlsx_compact.compact_xlsx(self.src, self.dst)
        out = openpyxl.load_workbook(self.dst)["G"]
        self.assertEqual([[c.value for c in row] for row in out.iter_rows()],
                         [["a", None, "c", "d"], ["e", "f", "g", "h"], [None, "j", "k", "l"]])
        with zipfile.ZipFile(self.dst) as z:
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn('<c r="C1"', sheet)         # the gappy rows were left alone
        self.assertIn('<c r="B3"', sheet)
        self.assertNotIn('<c r="A2"', sheet)      # the full row was compacted

    def test_a_chunk_boundary_never_splits_a_row(self):
        for gaps in (False, True):
            _workbook(self.src, rows=400, gaps=gaps)
            for chunk in (1, 7, 50, 333, 5000):
                with self.subTest(gaps=gaps, chunk=chunk), patch.object(xlsx_compact, "_CHUNK_BYTES", chunk):
                    xlsx_compact.compact_xlsx(self.src, self.dst)
                    self.assertEqual(_cells(self.src), _cells(self.dst))
                    if not gaps:
                        with zipfile.ZipFile(self.dst) as z:
                            self.assertNotIn(b'<c r="', z.read("xl/worksheets/sheet1.xml"))

    def test_styles_survive(self):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("S")
        from openpyxl.cell import WriteOnlyCell
        cell = WriteOnlyCell(ws, value=3)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FF0000")
        cell.number_format = '#,##0;-#,##0;"-"'
        ws.append([cell, 0])
        wb.save(self.src)
        xlsx_compact.compact_xlsx(self.src, self.dst)
        out = openpyxl.load_workbook(self.dst)["S"]["A1"]
        self.assertTrue(out.font.b)
        self.assertEqual(out.fill.fgColor.rgb, "00FF0000")
        self.assertEqual(out.number_format, '#,##0;-#,##0;"-"')

    def test_in_place_replaces_the_file_and_reports_the_saving(self):
        _workbook(self.src)
        before = self.src.stat().st_size
        saved = xlsx_compact.compact_xlsx_in_place(self.src)
        self.assertGreater(saved, 0)
        self.assertEqual(before - saved, self.src.stat().st_size)
        self.assertEqual(openpyxl.load_workbook(self.src)["Second"]["B2"].value, 2)
        self.assertEqual([p.name for p in Path(self.tmp.name).iterdir()], ["a.xlsx"])   # no temp file left


if __name__ == "__main__":
    unittest.main()
