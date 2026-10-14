import csv
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.jobs import JobUserError
from app.services.ultrafine_cashflow import generate, ingest, mapping_store as store, resolver, seed, summary
from app.services.ultrafine_cashflow.keys import description_key
from app.services.ultrafine_cashflow.models import (
    CfLocation,
    DescriptionAlias,
    Location,
    MonthInput,
    Party,
    PlanItem,
    PlanSection,
    Transaction,
    TransactionSplit,
)
from app.services.ultrafine_cashflow.statement_reader import Statement, StatementLine, read_statement


def _make_session():
    engine = create_engine("sqlite://")
    tables = [t for t in Base.metadata.sorted_tables if t.name.startswith("ufcf_")]
    Base.metadata.create_all(engine, tables=tables)
    return sessionmaker(bind=engine)()


def _line(day, hour, minute, direction, amount, description, ref="R1", month=9):
    return StatementLine(datetime(2026, month, day, hour, minute), direction, Decimal(str(amount)),
                         description, ref)


def _statement(lines, account="111222333"):
    return Statement(account_no=account, lines=lines)


def _basic_masters(db):
    """Two CF locations, a few locations, parties and an RDC sharing group."""
    store.save_cf_location(db, None, {"name": "Main Operation", "sort_order": 1, "payable_pct": 80,
                                      "reserve_pct": 20, "authorized_person": "A. Singh"})
    store.save_cf_location(db, None, {"name": "Flyash Business", "sort_order": 2, "payable_pct": 100})
    store.save_location(db, None, "Gujarat", "Main Operation")
    store.save_location(db, None, "Wada", "Main Operation")
    store.save_location(db, None, "Flyash", "Flyash Business")
    store.save_sharing_group(db, None, "RDC")
    store.save_party(db, None, "Vendor One - Crs", "Gujarat", None, None)
    store.save_party(db, None, "Rdc Concrete", None, None, "RDC")
    store.save_alias(db, None, "NEFT - VENDOR ONE", "Vendor One - Crs")
    store.save_alias(db, None, "CHQ DEP - RDC CONCRETE", "Rdc Concrete")
    db.commit()


def _rdc_ratio(db, effective="2026-09-01T00:00", main=0.6, fly=0.4):
    store.save_ratio_set(db, None, "RDC", effective,
                         [{"location": "Gujarat", "ratio": main}, {"location": "Flyash", "ratio": fly}])
    db.commit()


class DescriptionKeyTests(unittest.TestCase):
    def test_noise_is_removed_so_one_alias_covers_every_payment(self):
        a = description_key("NEFT  - FDRL0002301 - Rushika Enterprise - 23010200000646 - Rushika Enterprise")
        b = description_key("NEFT - FDRL0002301 - Rushika Enterprise - 99999999999999 - Rushika  Enterprise")
        self.assertEqual(a, b)
        self.assertEqual(a, "NEFT - RUSHIKA ENTERPRISE")

    def test_dates_references_and_own_company_name_are_ignored(self):
        self.assertEqual(description_key("FT-  2026-10-03- P2100P2100 - 0073041655 "), "FT - P2100P2100")
        self.assertEqual(
            description_key("NEFT Cr-ICIC0SF0002-SRK BUILDWELL PRIVATE LIMITED-ULTRAFINE MINERAL  ADMIXTURES-IN12627457688980"),
            "NEFT CR - SRK BUILDWELL PRIVATE LIMITED",
        )


class StatementReaderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def test_csv_with_extra_remark_column_and_header_totals(self):
        path = self.dir / "may.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            w = csv.writer(handle)
            for row in (
                ["Statement of Account"], ["Account No          : 575000"], ["From Date           : 01/05/2026"],
                ["To Date             : 31/05/2026"], ["Total Debit Count   : 1"], ["Total Credit Count  : 1"],
                ["Total Debit Amount  : 100.00"], ["Total Credit Amount : 250.50"], [],
                ["Transaction Date", "Description", "Amount", "My note", "C.D.Falg", "Reference No", "Value Date"],
                ["01-05-2026 04:51", "Vendor", "100", "", "D", "6.05E+11", "01-05-2026 00:00"],
                ["02-05-2026 10:00", "Customer", "250.50", "Already Done", "C", "X1", "02-05-2026 00:00"],
            ):
                w.writerow(row)
        statement = read_statement(path)
        self.assertEqual(statement.account_no, "575000")
        self.assertEqual([l.direction for l in statement.lines], ["D", "C"])
        self.assertEqual(statement.debit_total, Decimal("100.00"))
        self.assertEqual(statement.check_against_header(), [])

    def test_header_total_mismatch_is_reported(self):
        path = self.dir / "short.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            w = csv.writer(handle)
            for row in (
                ["Account No : 575000"], ["Total Credit Count  : 2"], ["Total Credit Amount : 300"],
                ["Transaction Date", "Description", "Amount", "C.D.Falg", "Reference No"],
                ["01-05-2026 04:51", "Customer", "100", "C", "R"],
            ):
                w.writerow(row)
        problems = read_statement(path).check_against_header()
        self.assertEqual(len(problems), 2)
        self.assertIn("300.00", problems[0])

    def test_xlsx_layout_with_debit_credit_column(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Account Number", "575000"])
        ws.append(["From Date", "01/09/2026"])
        ws.append([])
        ws.append(["Transaction Date", "Transaction Description", "Transaction Amount", "Debit / Credit",
                   "Reference No."])
        ws.append(["01/09/2026 01:30:54", "RTGS Cr-X", 2717867, "C", "REF1"])
        path = self.dir / "sep.xlsx"
        wb.save(path)
        statement = read_statement(path)
        self.assertEqual(statement.lines[0].amount, Decimal("2717867.00"))
        self.assertEqual(statement.lines[0].txn_dt, datetime(2026, 9, 1, 1, 30, 54))

    def test_non_statement_is_rejected(self):
        path = self.dir / "junk.csv"
        path.write_text("a,b\n1,2\n", encoding="utf-8")
        with self.assertRaises(JobUserError):
            read_statement(path)


class IngestTests(unittest.TestCase):
    def setUp(self):
        self.db = _make_session()

    def test_reupload_adds_nothing_and_overlap_adds_only_new_lines(self):
        first = _statement([_line(1, 10, 0, "D", 100, "A"), _line(2, 11, 0, "C", 50, "B")])
        out = ingest.ingest_statement(self.db, first, "a.csv", None)
        self.assertEqual((out["new_count"], out["duplicate_count"]), (2, 0))
        again = ingest.ingest_statement(self.db, first, "a.csv", None)
        self.assertEqual((again["new_count"], again["duplicate_count"]), (0, 2))
        overlap = _statement([_line(2, 11, 0, "C", 50, "B"), _line(3, 9, 0, "D", 7, "C")])
        out = ingest.ingest_statement(self.db, overlap, "b.csv", None)
        self.assertEqual((out["new_count"], out["duplicate_count"]), (1, 1))

    def test_csv_and_xls_exports_of_the_same_line_dedupe_despite_seconds_and_reference(self):
        csv_line = StatementLine(datetime(2026, 5, 1, 18, 32), "D", Decimal("12676"), "Shahabaz Khan", "6.05E+11")
        xls_line = StatementLine(datetime(2026, 5, 1, 18, 32, 41), "D", Decimal("12676"), "Shahabaz  Khan",
                                 "HDFCH00968195663")
        ingest.ingest_statement(self.db, _statement([csv_line]), "csv", None)
        out = ingest.ingest_statement(self.db, _statement([xls_line]), "xls", None)
        self.assertEqual((out["new_count"], out["duplicate_count"]), (0, 1))

    def test_identical_lines_in_one_file_are_kept_separately(self):
        twin = _line(1, 10, 0, "D", 100, "Same")
        out = ingest.ingest_statement(self.db, _statement([twin, twin]), "a", None)
        self.assertEqual(out["new_count"], 2)
        again = ingest.ingest_statement(self.db, _statement([twin, twin]), "a", None)
        self.assertEqual(again["new_count"], 0)
        three = ingest.ingest_statement(self.db, _statement([twin, twin, twin]), "a", None)
        self.assertEqual(three["new_count"], 1)

    def test_other_account_is_independent(self):
        line = _line(1, 10, 0, "D", 100, "A")
        ingest.ingest_statement(self.db, _statement([line], "111"), "a", None)
        out = ingest.ingest_statement(self.db, _statement([line], "222"), "b", None)
        self.assertEqual(out["new_count"], 1)


class ResolverTests(unittest.TestCase):
    def setUp(self):
        self.db = _make_session()
        _basic_masters(self.db)

    def _month(self, lines):
        ingest.ingest_statement(self.db, _statement(lines), "f", None)
        self.db.commit()
        return resolver.allocations_for_month(self.db, "2026-09")

    def test_plain_party_resolves_name_location_and_derived_cf_location(self):
        rows = self._month([_line(1, 10, 0, "D", 500, "NEFT - FDRL0002301 - Vendor One - 1234567890 - Vendor One")])
        self.assertEqual((rows[0].name, rows[0].location, rows[0].cf_location, rows[0].issues),
                         ("Vendor One - Crs", "Gujarat", "Main Operation", []))

    def test_location_named_like_a_cf_location_falls_back_to_it(self):
        store.save_location(self.db, None, "Capex", None)       # no default CF location set
        store.save_cf_location(self.db, None, {"name": "Capex"})
        store.save_party(self.db, None, "Plant Builder - Crs", "Capex", None, None)
        store.save_alias(self.db, None, "NEFT - Plant Builder", "Plant Builder - Crs")
        self.db.commit()
        rows = self._month([_line(1, 10, 0, "D", 500, "NEFT - Plant Builder")])
        self.assertEqual((rows[0].location, rows[0].cf_location, rows[0].issues), ("Capex", "Capex", []))

    def test_unknown_description_is_reported_not_guessed(self):
        rows = self._month([_line(1, 10, 0, "D", 500, "Someone New")])
        self.assertEqual(rows[0].issues, [resolver.ISSUE_NO_PARTY])
        report = resolver.unmapped_summary(rows, resolver.load_maps(self.db))
        self.assertEqual(report["no_party"][0]["desc_key"], "SOMEONE NEW")

    def test_sharing_split_adds_up_to_the_bank_amount_to_the_paisa(self):
        _rdc_ratio(self.db, main=0.3333333333, fly=0.6666666667)
        rows = self._month([_line(1, 10, 0, "C", 100.01, "CHQ DEP - RDC CONCRETE 8841123")])
        self.assertEqual(len(rows), 2)
        self.assertEqual(sum(r.amount for r in rows), Decimal("100.01"))
        self.assertEqual({r.cf_location for r in rows}, {"Main Operation", "Flyash Business"})

    def test_ratio_set_in_force_on_the_transaction_time_is_used(self):
        _rdc_ratio(self.db, "2026-09-01T00:00", 0.6, 0.4)
        _rdc_ratio(self.db, "2026-09-15T12:00", 0.1, 0.9)
        rows = self._month([
            _line(10, 9, 0, "C", 1000, "CHQ DEP - RDC CONCRETE"),
            _line(15, 11, 59, "C", 1000, "CHQ DEP - RDC CONCRETE 7000001"),
            _line(15, 12, 0, "C", 1000, "CHQ DEP - RDC CONCRETE 7000002"),
        ])
        by_day = {}
        for row in rows:
            by_day.setdefault(row.description, {})[row.location] = row.amount
        self.assertEqual(by_day["CHQ DEP - RDC CONCRETE"]["Gujarat"], Decimal("600.00"))
        self.assertEqual(by_day["CHQ DEP - RDC CONCRETE 7000001"]["Gujarat"], Decimal("600.00"))
        self.assertEqual(by_day["CHQ DEP - RDC CONCRETE 7000002"]["Gujarat"], Decimal("100.00"))

    def test_no_ratio_before_the_first_set_is_an_issue(self):
        _rdc_ratio(self.db, "2026-09-10T00:00")
        rows = self._month([_line(1, 10, 0, "C", 1000, "CHQ DEP - RDC CONCRETE")])
        self.assertEqual(rows[0].issues, [resolver.ISSUE_NO_RATIO])

    def test_manual_split_replaces_the_mapping_and_must_match_the_amount(self):
        rows = self._month([_line(30, 18, 48, "D", 1000, "Multiple Payees ULT44NES3009.001 -")])
        txn = self.db.query(Transaction).one()
        gujarat = store.find_named(self.db, Location, "Gujarat")
        flyash = store.find_named(self.db, Location, "Flyash")
        self.db.add_all([
            TransactionSplit(txn_id=txn.id, location_id=gujarat.id, amount=Decimal("700"), sort_order=0),
            TransactionSplit(txn_id=txn.id, location_id=flyash.id, amount=Decimal("300"), sort_order=1),
        ])
        self.db.commit()
        rows = resolver.allocations_for_month(self.db, "2026-09")
        self.assertEqual([(r.cf_location, r.amount, r.issues) for r in rows],
                         [("Main Operation", Decimal("700"), []), ("Flyash Business", Decimal("300"), [])])
        self.db.query(TransactionSplit).filter(TransactionSplit.location_id == flyash.id).one().amount = Decimal("250")
        self.db.commit()
        rows = resolver.allocations_for_month(self.db, "2026-09")
        self.assertTrue(all(resolver.ISSUE_SPLIT_MISMATCH in r.issues for r in rows))

    def test_party_override_on_one_transaction_wins_over_the_description(self):
        self._month([_line(1, 10, 0, "D", 500, "NEFT - Vendor One")])
        txn = self.db.query(Transaction).one()
        other = store.save_party(self.db, None, "Wada Vendor", "Wada", None, None)
        txn.party_id = other["id"]
        self.db.commit()
        rows = resolver.allocations_for_month(self.db, "2026-09")
        self.assertEqual((rows[0].name, rows[0].location), ("Wada Vendor", "Wada"))

    def test_renaming_a_party_changes_every_transaction_with_no_other_edit(self):
        self._month([_line(1, 10, 0, "D", 500, "NEFT - Vendor One"), _line(2, 10, 0, "D", 600, "NEFT - Vendor One")])
        party = store.find_named(self.db, Party, "Vendor One - Crs")
        store.save_party(self.db, party.id, "Vendor One Pvt Ltd", "Wada", None, None)
        self.db.commit()
        rows = resolver.allocations_for_month(self.db, "2026-09")
        self.assertEqual({(r.name, r.location) for r in rows}, {("Vendor One Pvt Ltd", "Wada")})

    def test_party_cf_override_wins_over_the_location_default(self):
        party = store.find_named(self.db, Party, "Vendor One - Crs")
        store.save_party(self.db, party.id, party.name, "Gujarat", "Flyash Business", None)
        self.db.commit()
        rows = self._month([_line(1, 10, 0, "D", 500, "NEFT - Vendor One")])
        self.assertEqual(rows[0].cf_location, "Flyash Business")


class MappingStoreTests(unittest.TestCase):
    def setUp(self):
        self.db = _make_session()
        _basic_masters(self.db)

    def test_names_are_unique_ignoring_case_and_spacing(self):
        with self.assertRaises(store.MappingError):
            store.save_party(self.db, None, "  vendor ONE - crs ", "Gujarat", None, None)
        self.assertEqual(self.db.query(Party).filter(Party.name.like("%endor%")).count(), 1)

    def test_archived_names_come_back_on_create_and_in_use_rows_cannot_be_archived(self):
        location = store.find_named(self.db, Location, "Wada")
        store.archive_location(self.db, location.id)
        self.assertTrue(location.is_deleted)
        store.save_location(self.db, None, "WADA", "Main Operation")
        self.assertFalse(location.is_deleted)
        with self.assertRaises(store.MappingError):
            store.archive_location(self.db, store.find_named(self.db, Location, "Gujarat").id)
        with self.assertRaises(store.MappingError):
            store.archive_party(self.db, store.find_named(self.db, Party, "Vendor One - Crs").id)

    def test_fix_panel_saves_the_cf_location_for_a_new_name_but_never_overwrites(self):
        created = store.get_or_create_party(self.db, "Fresh Vendor", "Wada", "Flyash Business")
        self.assertEqual((created.location_id is not None, created.cf_location_id is not None), (True, True))
        existing = store.find_named(self.db, Party, "Vendor One - Crs")
        kept_cf = existing.cf_location_id
        store.get_or_create_party(self.db, "Vendor One - Crs", "Wada", "Flyash Business")
        self.assertEqual(existing.cf_location_id, kept_cf)
        with self.assertRaises(store.MappingError):
            store.get_or_create_party(self.db, "Another", "Wada", "No Such CF")

    def test_fix_panel_rules_new_names_need_both_known_names_need_neither(self):
        with self.assertRaises(store.MappingError):
            store.get_or_create_party(self.db, "Brand New Co", "Wada", None, require_placement=True)
        with self.assertRaises(store.MappingError):
            store.get_or_create_party(self.db, "Brand New Co", None, "Flyash Business", require_placement=True)
        made = store.get_or_create_party(self.db, "Brand New Co", "Wada", "Flyash Business", require_placement=True)
        self.assertIsNotNone(made.cf_location_id)
        # a known, complete name needs nothing and is left exactly as it is
        known = store.get_or_create_party(self.db, "Vendor One - Crs", None, None, require_placement=True)
        self.assertEqual(known.id, store.find_named(self.db, Party, "Vendor One - Crs").id)

    def test_ratio_sets_must_total_one_hundred_percent(self):
        with self.assertRaises(store.MappingError) as ctx:
            store.save_ratio_set(self.db, None, "RDC", "2026-09-01T00:00",
                                 [{"location": "Gujarat", "ratio": 0.5}, {"location": "Flyash", "ratio": 0.4}])
        self.assertIn("100%", str(ctx.exception))

    def test_editing_a_ratio_set_updates_in_place_without_piling_up_rows(self):
        first = store.save_ratio_set(self.db, None, "RDC", "2026-09-01T00:00",
                                     [{"location": "Gujarat", "ratio": 0.6}, {"location": "Flyash", "ratio": 0.4}])
        store.save_ratio_set(self.db, first["id"], "RDC", "2026-09-01T00:00",
                             [{"location": "Gujarat", "ratio": 0.7}, {"location": "Flyash", "ratio": 0.3}])
        self.db.commit()
        sets = store.list_ratio_sets(self.db)
        self.assertEqual(len(sets), 1)
        self.assertEqual([l["ratio"] for l in sets[0]["lines"]], [0.7, 0.3])

    def test_a_sharing_party_ignores_location_fields(self):
        saved = store.save_party(self.db, None, "Shared Co", "Gujarat", "Main Operation", "RDC")
        party = self.db.get(Party, saved["id"])
        self.assertIsNone(party.location_id)
        self.assertIsNone(party.cf_location_id)

    def test_resolve_key_reuses_the_party_and_never_overwrites_its_location(self):
        party = store.get_or_create_party(self.db, "vendor one - CRS", "Wada")
        self.assertEqual(party.location_id, store.find_named(self.db, Location, "Gujarat").id)
        fresh = store.get_or_create_party(self.db, "Brand New", "Wada")
        self.assertEqual(fresh.location_id, store.find_named(self.db, Location, "Wada").id)


class SummaryAndWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.db = _make_session()
        _basic_masters(self.db)
        _rdc_ratio(self.db, "2026-01-01T00:00")
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        patcher = patch.object(generate, "STORAGE_DIR", Path(tmp.name))
        patcher.start()
        self.addCleanup(patcher.stop)

    def _upload(self, lines):
        ingest.ingest_statement(self.db, _statement(lines), "f", None)
        self.db.commit()

    def test_closing_balance_carries_into_the_next_month_unless_overridden(self):
        self._upload([
            _line(2, 10, 0, "C", 1_000_000, "CHQ DEP - RDC CONCRETE", month=8),
            _line(3, 10, 0, "D", 200_000, "NEFT - Vendor One", month=8),
            _line(2, 10, 0, "C", 500_000, "CHQ DEP - RDC CONCRETE", month=9),
        ])
        august = {r.name: r for r in summary.month_figures(self.db, "2026-08").rows}
        # RDC 60/40: Main Operation collects 6.00 lakh, pays 2.00 lakh; 80% payable.
        self.assertAlmostEqual(august["Main Operation"].collection, 6.0)
        self.assertAlmostEqual(august["Main Operation"].closing, 6.0 * 0.8 - 2.0)
        september = {r.name: r for r in summary.month_figures(self.db, "2026-09").rows}
        self.assertAlmostEqual(september["Main Operation"].opening, august["Main Operation"].closing)
        main = store.find_named(self.db, CfLocation, "Main Operation")
        self.db.add(MonthInput(month="2026-09", cf_location_id=main.id, opening_balance=Decimal("10")))
        self.db.commit()
        september = {r.name: r for r in summary.month_figures(self.db, "2026-09").rows}
        self.assertEqual(september["Main Operation"].opening, 10.0)
        self.assertTrue(september["Main Operation"].opening_is_manual)

    def test_planning_amounts_accept_simple_sums_only(self):
        self.assertEqual(summary.evaluate_amount("20+3+1.5"), 24.5)
        self.assertEqual(summary.evaluate_amount("=10*2"), 20.0)
        self.assertIsNone(summary.evaluate_amount("__import__('os')"))
        self.assertIsNone(summary.evaluate_amount("abc"))

    def test_generation_is_blocked_while_a_line_is_unmapped(self):
        self._upload([_line(1, 10, 0, "D", 500, "Mystery payee")])
        result = generate.generate_month(self.db, "2026-09", None)
        self.assertEqual(result["status"], "needs_mapping_fix")
        self.assertEqual(result["unmapped"]["no_party"][0]["desc_key"], "MYSTERY PAYEE")
        self.assertIsNone(generate.STORAGE_DIR and next(iter(generate.STORAGE_DIR.glob("**/*.xlsx")), None))

    def test_unmapped_lines_can_be_written_anyway_and_are_flagged(self):
        self._upload([
            _line(2, 10, 0, "C", 1_000_000, "CHQ DEP - RDC CONCRETE"),
            _line(3, 10, 0, "D", 200_000, "NEFT - Vendor One"),
            _line(4, 10, 0, "D", 500_000, "Mystery payee"),
        ])
        blocked = generate.generate_month(self.db, "2026-09", None)
        self.assertEqual(blocked["status"], "needs_mapping_fix")

        result = generate.generate_month(self.db, "2026-09", None, allow_unmapped=True)
        self.db.commit()
        self.assertEqual((result["status"], result["unmapped_count"]), ("generated", 1))
        path = generate.STORAGE_DIR / "2026-09" / "UF_Cashflow_Sep_26.xlsx"
        values = load_workbook(path, data_only=True)
        payments, summary_sheet = values["Payments"], values["Summary"]
        headers = [c.value for c in payments[3]]
        self.assertEqual(headers[-1], "Remark")
        flagged = [r for r in payments.iter_rows(min_row=4) if r[-1].value]
        self.assertEqual(len(flagged), 1)
        self.assertFalse(flagged[0][3].value)   # no name is left blank, never invented
        rows = {summary_sheet.cell(r, 3).value: r for r in range(3, 14)}
        # the unmapped 5.00 lakh gets its own row, so the Summary total still ties to the sheets
        self.assertAlmostEqual(summary_sheet.cell(rows["Not mapped"], 9).value, 5.0)
        self.assertAlmostEqual(summary_sheet.cell(rows["Total"], 9).value, 7.0)   # 2.00 + 5.00 paid
        self.assertEqual(payments["E2"].value, 700_000)

    def test_workbook_has_live_formulas_cached_values_and_every_line(self):
        section = PlanSection(month="2026-09", block_no=1, title="Salaries", sort_order=0)
        self.db.add(section)
        self.db.flush()
        self.db.add_all([PlanItem(section_id=section.id, label="Onroll", amount_expr="45", sort_order=0),
                         PlanItem(section_id=section.id, label="Labour", amount_expr="20+3+1.5", sort_order=1)])
        self._upload([
            _line(2, 10, 0, "C", 1_000_000, "CHQ DEP - RDC CONCRETE"),
            _line(3, 10, 0, "D", 200_000, "NEFT - Vendor One"),
        ])
        first = generate.generate_month(self.db, "2026-09", 7)
        self.db.commit()
        self.assertEqual((first["status"], first["version"]), ("generated", 1))
        path = generate.STORAGE_DIR / "2026-09" / "UF_Cashflow_Sep_26.xlsx"

        formulas = load_workbook(path)
        values = load_workbook(path, data_only=True)
        self.assertEqual(formulas.sheetnames, ["Summary", "Payments", "Receipts", "Working"])
        summary_sheet, cached = formulas["Summary"], values["Summary"]
        rows = {summary_sheet.cell(r, 3).value: r for r in range(3, 12)}
        main_row = rows["Main Operation"]
        self.assertIn("SUMIF(Receipts!", summary_sheet.cell(main_row, 6).value)
        self.assertAlmostEqual(cached.cell(main_row, 6).value, 6.0)       # collection, lakhs
        self.assertAlmostEqual(cached.cell(main_row, 9).value, 2.0)       # paid
        self.assertAlmostEqual(cached.cell(main_row, 11).value, 6.0 * 0.8 - 2.0)
        self.assertEqual(values["Payments"]["E2"].value, 200000)
        self.assertEqual(values["Receipts"]["E2"].value, 1_000_000)
        plan_totals = [c.value for row in cached.iter_rows(min_col=5, max_col=5) for c in row
                       if isinstance(c.value, (int, float)) and abs(c.value - 69.5) < 1e-9]
        self.assertTrue(plan_totals, "planning block total (45 + 24.5) should be cached")
        # Working sheet lists the dated ratio set
        working = [c.value for row in values["Working"].iter_rows() for c in row if c.value is not None]
        self.assertIn("RDC", working)

        # Regenerating replaces the stored file and bumps the version.
        second = generate.generate_month(self.db, "2026-09", 7)
        self.assertEqual(second["version"], 2)
        self.assertEqual(len(list(generate.STORAGE_DIR.glob("**/*.xlsx"))), 1)


class SeedTests(unittest.TestCase):
    def test_seed_is_additive_and_never_revives_archived_rows(self):
        db = _make_session()
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        storage = patch.object(generate, "STORAGE_DIR", Path(tmp.name))
        storage.start()
        self.addCleanup(storage.stop)
        added = seed.seed_missing_from_json(db)
        db.commit()
        self.assertGreater(added["parties"], 100)
        self.assertEqual(db.query(CfLocation).filter(CfLocation.show_in_summary == True).count(), 7)  # noqa: E712

        ratio_groups = {s["group"] for s in store.list_ratio_sets(db)}
        self.assertEqual(ratio_groups, {"RDC", "IOCL", "ETC Pool"})
        for ratio_set in store.list_ratio_sets(db):
            self.assertAlmostEqual(ratio_set["total"], 1.0, places=6)

        edited = store.find_named(db, Location, "Gujarat")
        edited.cf_location_id = None
        archived = db.query(DescriptionAlias).first()
        archived.is_deleted = True
        db.commit()
        again = seed.seed_missing_from_json(db)
        db.commit()
        self.assertEqual(again, {})
        self.assertIsNone(store.find_named(db, Location, "Gujarat").cf_location_id)
        self.assertTrue(db.get(DescriptionAlias, archived.id).is_deleted)


class ApiTests(unittest.TestCase):
    """The route handlers themselves, called with a throwaway SQLite
    session (no HTTP client, login or MySQL needed)."""

    def setUp(self):
        from sqlalchemy.pool import StaticPool

        from app.routers import ultrafine_cashflow as router

        engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(engine, tables=[t for t in Base.metadata.sorted_tables
                                                 if t.name.startswith("ufcf_")])
        self.session_factory = sessionmaker(bind=engine)
        self.router = router
        self.db = self.session_factory()
        self.addCleanup(self.db.close)
        _basic_masters(self.db)
        _rdc_ratio(self.db, "2026-01-01T00:00")

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        for target in (patch.object(generate, "STORAGE_DIR", self.tmp / "store"),
                       patch.object(router, "SessionLocal", self.session_factory)):
            target.start()
            self.addCleanup(target.stop)

    def _write_statement(self, name="stmt.csv"):
        path = self.tmp / name
        with path.open("w", newline="", encoding="utf-8") as handle:
            w = csv.writer(handle)
            for row in (
                ["Account No : 575000"],
                ["Transaction Date", "Description", "Amount", "C.D.Falg", "Reference No"],
                ["02-09-2026 10:00", "CHQ DEP - RDC CONCRETE 8841123", "1000000", "C", "R1"],
                ["03-09-2026 10:00", "NEFT - FDRL0002301 - Vendor One - 12345678901 - Vendor One", "200000", "D", "R2"],
                ["04-09-2026 10:00", "Mystery payee", "5000", "D", "R3"],
            ):
                w.writerow(row)
        return path

    def _http_error(self, fn, *args, **kwargs):
        from fastapi import HTTPException

        with self.assertRaises(HTTPException) as ctx:
            fn(*args, **kwargs)
        return ctx.exception

    def test_upload_job_stores_lines_and_blocks_generation_until_mapped(self):
        r, db = self.router, self.db
        result = r._job_ingest([str(self._write_statement())], ["stmt.csv"], 1)
        self.assertTrue(result["files"][0]["ok"])
        self.assertEqual(result["files"][0]["new_count"], 3)
        month = result["months"][0]
        self.assertEqual((month["month"], month["status"], month["issue_count"]), ("2026-09", "needs_mapping_fix", 1))

        again = r._job_ingest([str(self._write_statement("again.csv"))], ["again.csv"], 1)
        self.assertEqual(again["files"][0]["new_count"], 0)
        self.assertEqual(again["months"][0]["status"], "unchanged")

        listed = r.list_months(db=db, user=None)
        self.assertEqual((listed[0]["month"], listed[0]["transactions"], listed[0]["issue_count"], listed[0]["file"]),
                         ("2026-09", 3, 1, None))
        unmapped = r.month_unmapped("2026-09", db=db, user=None)
        self.assertEqual(unmapped["no_party"][0]["desc_key"], "MYSTERY PAYEE")

        # Fix it through the fix-up endpoint, regenerate, and fetch the stored workbook.
        r.resolve_key(r.ResolveKeyIn(desc_key="Mystery payee", party="Mystery Co", location="Wada",
                                  cf_location="Main Operation"),
                      db=db, user=None)
        outcome = r._job_generate("2026-09", 1)
        self.assertEqual((outcome["status"], outcome["version"]), ("generated", 1))
        response = r.download_month("2026-09", db=self.session_factory(), user=None)
        self.assertTrue(Path(response.path).read_bytes().startswith(b"PK"))
        self.assertEqual(self._http_error(r.download_month, "2026-08", db=db, user=None).status_code, 404)
        self.assertEqual(self._http_error(r.download_month, "bad", db=db, user=None).status_code, 400)

    def test_a_bad_file_is_reported_per_file_and_does_not_stop_the_others(self):
        bad = self.tmp / "bad.csv"
        bad.write_text("a,b\n1,2\n", encoding="utf-8")
        result = self.router._job_ingest([str(bad), str(self._write_statement())], ["bad.csv", "good.csv"], 1)
        self.assertFalse(result["files"][0]["ok"])
        self.assertIn("bank statement", result["files"][0]["error"])
        self.assertEqual(result["files"][1]["new_count"], 3)

    def test_transaction_override_and_split_flow_through_the_listing(self):
        r, db = self.router, self.db
        r._job_ingest([str(self._write_statement())], ["s.csv"], 1)
        page = r.month_transactions("2026-09", status="issues", db=db, user=None)
        self.assertEqual(page["total"], 1)
        txn_id = page["items"][0]["id"]

        bad = self._http_error(r.set_override, txn_id,
                               r.OverrideIn(splits=[r.SplitIn(location="Gujarat", amount=1)]), db=db, user=None)
        self.assertEqual(bad.status_code, 400)
        self.assertIn("adds up to 1.00", bad.detail)

        r.set_override(txn_id, r.OverrideIn(splits=[r.SplitIn(location="Gujarat", amount=3000),
                                                    r.SplitIn(location="Flyash", amount=2000)]), db=db, user=None)
        row = r.month_transactions("2026-09", status="manual", db=db, user=None)["items"][0]
        self.assertEqual([(a["cf_location"], a["amount"]) for a in row["allocations"]],
                         [("Main Operation", 3000.0), ("Flyash Business", 2000.0)])
        self.assertEqual(r.month_transactions("2026-09", status="issues", db=db, user=None)["total"], 0)

        r.set_override(txn_id, r.OverrideIn(splits=[]), db=db, user=None)
        self.assertEqual(r.month_transactions("2026-09", status="issues", db=db, user=None)["total"], 1)
        self.assertEqual(r.month_transactions("2026-09", q="vendor", db=db, user=None)["total"], 1)
        self.assertEqual(r.month_transactions("2026-09", direction="C", db=db, user=None)["total"], 1)
        self.assertEqual(self._http_error(r.set_override, 999999, r.OverrideIn(), db=db, user=None).status_code, 404)

    def test_a_date_or_time_cannot_be_saved_as_a_description_key(self):
        for junk in ("17:22:01", "08/10/2026 18:08:14"):
            error = self._http_error(self.router.add_alias,
                                     self.router.AliasIn(desc_key=junk, party="Vendor One"), db=self.db, user=None)
            self.assertEqual(error.status_code, 400)
            self.assertIn("date or time", error.detail)

    def test_ratio_lines_are_just_a_particular_and_a_percentage(self):
        r, db = self.router, self.db
        r.add_ratio_set(r.RatioSetIn(group="IOCL", effective_from="2026-10-01T00:00", lines=[
            r.RatioLineIn(location="Main", ratio=0.4, cf_location="Main Operation"),
            r.RatioLineIn(location="Flyash Business", ratio=0.6)]), db=db, user=None)   # a CF name needs no CF
        lines = {l["location"]: l["cf_location"] for l in
                 next(s for s in r.get_ratio_sets(db=db, user=None) if s["group"] == "IOCL")["lines"]}
        self.assertEqual(lines, {"Main": "Main Operation", "Flyash Business": "Flyash Business"})
        # A name can be given a brand-new location by typing it; the CF location must already exist.
        r.add_party(r.PartyIn(name="New Vendor", location="Brand New Site", cf_location="Main Operation"),
                    db=db, user=None)
        self.assertIn("Brand New Site", [l["name"] for l in r.get_locations(db=db, user=None)])
        bad = self._http_error(r.add_party, r.PartyIn(name="Other", location="X", cf_location="Nope"), db=db, user=None)
        self.assertEqual(bad.status_code, 400)

    def test_the_seeded_october_workbook_is_listed_and_downloadable(self):
        r, db = self.router, self.db
        seed.seed_missing_from_json(db)
        db.commit()
        month = next(m for m in r.list_months(db=db, user=None) if m["month"] == "2026-10")
        self.assertEqual((month["transactions"], month["file"]["source"], month["file"]["available"]),
                         (0, "uploaded", True))
        self.assertTrue(Path(r.download_month("2026-10", db=self.session_factory(), user=None).path)
                        .read_bytes().startswith(b"PK"))

    def test_mapping_crud_reports_mistakes_as_400_and_archives_instead_of_deleting(self):
        r, db = self.router, self.db
        created = r.add_location(r.LocationIn(name="Nagpur", cf_location="Main Operation"), db=db, user=None)
        self.assertEqual(self._http_error(r.add_location, r.LocationIn(name="nagpur"), db=db, user=None).status_code, 400)
        self.assertEqual(self._http_error(r.add_location, r.LocationIn(name="X", cf_location="Nope"),
                                          db=db, user=None).status_code, 400)
        r.remove_location(created["id"], db=db, user=None)
        self.assertNotIn("Nagpur", [l["name"] for l in r.get_locations(db=db, user=None)])
        self.assertTrue(self.session_factory().get(Location, created["id"]).is_deleted)   # archived, never removed
        too_small = r.RatioSetIn(group="RDC", effective_from="2026-12-01T00:00",
                                 lines=[r.RatioLineIn(location="Gujarat", ratio=0.5)])
        self.assertEqual(self._http_error(r.add_ratio_set, too_small, db=db, user=None).status_code, 400)
        self.assertEqual(r.mapping_options(db=db, user=None)["sharing_groups"], ["RDC"])

    def test_month_setup_round_trip_validates_planning_amounts(self):
        r, db = self.router, self.db
        r._job_ingest([str(self._write_statement())], ["s.csv"], 1)
        r.save_setup("2026-09", r.SetupIn(
            inputs=[r.InputIn(cf_location="Main Operation", opening_balance=12.5, statutory_amount=1)],
            plan=[r.PlanSectionIn(title="Salaries", items=[
                r.PlanItemIn(label="Onroll", amount_expr="45"), r.PlanItemIn(label="Labour", amount_expr="20+3+1.5")])]),
            db=db, user=None)
        data = r.get_setup("2026-09", db=db, user=None)
        main = next(row for row in data["inputs"] if row["cf_location"] == "Main Operation")
        self.assertEqual((main["opening_balance"], main["statutory_amount"], main["computed_opening"]), (12.5, 1.0, 12.5))
        section = data["plan"][0]["sections"][0]
        self.assertEqual([i["amount"] for i in section["items"]], [45.0, 24.5])

        r.save_setup("2026-09", r.SetupIn(plan=[r.PlanSectionIn(id=section["id"], title="Salaries", items=[
            r.PlanItemIn(id=section["items"][0]["id"], label="Onroll", amount_expr="50")])]), db=db, user=None)
        check = self.session_factory()
        self.assertEqual(check.query(PlanSection).filter(PlanSection.month == "2026-09").count(), 1)  # edited in place
        invalid = r.SetupIn(plan=[r.PlanSectionIn(title="x", items=[
            r.PlanItemIn(label="bad", amount_expr="__import__('os')")])])
        self.assertEqual(self._http_error(r.save_setup, "2026-09", invalid, db=db, user=None).status_code, 400)
        self.assertEqual(self._http_error(r.save_setup, "bad", r.SetupIn(), db=db, user=None).status_code, 400)


if __name__ == "__main__":
    unittest.main()
