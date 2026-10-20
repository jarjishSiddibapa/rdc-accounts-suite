"""Automatic report mail: settings validation, the weekday/time schedule, the
send-and-record flow, and the Untagged / GST mail builders."""

import datetime as dt
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.jobs import JobUserError
from app.models import (
    Application, ApplicationEmailRecipient, ScheduledReportMail, ScheduledReportMailLog, SystemEmailSettings,
)
from app.regional import IST
from app.services import scheduled_report_mail as mail
from app.services.gst_invoice_adder import auto_mail as gst_mail
from app.services.mailer_shared import sheet_to_html
from app.services.untagged_invoices import auto_mail as untagged_mail
from app.services.untagged_invoices import processor
from app.services.untagged_invoices.report_job import BuiltReport

APP = "untagged-invoices-report"
GOOD = {
    "enabled": True, "days": [0, 2, 5], "time": "08:30",
    "to": ["a@rdc.in"], "cc": ["b@rdc.in"],
    "subject_template": "Report as on {date}", "intro_template": "Hello {date}",
    "signature": "", "send_when_unmapped": True,
}


class _DbCase(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite://")
        Base.metadata.create_all(
            engine, tables=[Application.__table__, ApplicationEmailRecipient.__table__, ScheduledReportMail.__table__,
                    SystemEmailSettings.__table__, ScheduledReportMailLog.__table__],
        )
        self.factory = sessionmaker(bind=engine, expire_on_commit=False)
        self.db = self.factory()
        for target in ("SessionLocal",):
            patcher = patch.object(mail, target, self.factory)
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = patch.object(mail.audit_middleware, "log_event", lambda *a, **k: None)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.db.close)

    def row(self, app_key=APP):
        self.db.expire_all()
        return mail.get_or_create_settings(self.db, app_key)


class SaveSettingsTests(_DbCase):
    def test_saves_days_time_wording_and_recipients(self):
        view = mail.save_settings(self.db, APP, GOOD)
        self.assertEqual(view["days"], [0, 2, 5])
        self.assertEqual(view["time"], "08:30")
        self.assertEqual(view["to"], ["a@rdc.in"])
        self.assertEqual(view["cc"], ["b@rdc.in"])
        self.assertTrue(view["enabled"])
        self.assertIsNotNone(view["next_run"])

    def test_defaults_before_anything_is_saved(self):
        view = mail.settings_view(self.db, APP)
        self.assertFalse(view["enabled"])
        self.assertEqual(view["time"], "08:30")
        self.assertEqual(view["days"], [0, 1, 2, 3, 4, 5])
        self.assertIsNone(view["next_run"])

    def test_refuses_bad_input_with_a_plain_message(self):
        cases = [
            ({"days": []}, "at least one day"),
            ({"time": "8:30"}, "HH:MM"),
            ({"time": "25:00"}, "HH:MM"),
            ({"to": []}, "To recipient"),
            ({"to": ["not-an-email"]}, "valid email"),
            ({"subject_template": "  "}, "subject cannot be empty"),
            ({"subject_template": "a\nb"}, "single line"),
            ({"intro_template": "Hi {nope}"}, "not available"),
            ({"intro_template": "Hi {date"}, "unmatched"),
        ]
        for change, expected in cases:
            with self.subTest(change=change):
                with self.assertRaises(ValueError) as caught:
                    mail.save_settings(self.db, APP, {**GOOD, **change})
                self.assertIn(expected, str(caught.exception))

    def test_switched_off_needs_no_days_or_recipients(self):
        view = mail.save_settings(self.db, APP, {**GOOD, "enabled": False, "days": [], "to": []})
        self.assertFalse(view["enabled"])
        self.assertIsNone(view["next_run"])

    def test_saving_after_the_time_has_passed_does_not_fire_a_mail_today(self):
        with patch.object(mail, "now_ist", return_value=dt.datetime(2026, 10, 19, 10, 0, tzinfo=IST)):
            mail.save_settings(self.db, APP, GOOD)          # Monday 10:00, send time 08:30
            row = self.row()
            self.assertFalse(mail.due_now(row, dt.datetime(2026, 10, 19, 10, 1, tzinfo=IST)))
            self.assertEqual(mail.next_run(row, dt.datetime(2026, 10, 19, 10, 1, tzinfo=IST)),
                             dt.datetime(2026, 10, 21, 8, 30, tzinfo=IST))

    def test_saving_before_the_time_still_sends_today(self):
        with patch.object(mail, "now_ist", return_value=dt.datetime(2026, 10, 19, 7, 0, tzinfo=IST)):
            mail.save_settings(self.db, APP, GOOD)
            self.assertTrue(mail.due_now(self.row(), dt.datetime(2026, 10, 19, 8, 31, tzinfo=IST)))

    def test_duplicate_recipient_is_refused(self):
        with self.assertRaises(ValueError):
            mail.save_settings(self.db, APP, {**GOOD, "cc": ["a@rdc.in"]})

    def test_gst_report_accepts_its_own_placeholders(self):
        mail.save_settings(self.db, "gst-invoice-adder", {
            **GOOD, "intro_template": "{total} rows, {found} found, {blank} blank on {date}",
        })


class ScheduleTests(_DbCase):
    def _enabled(self):
        # saved on a Sunday morning, so no day's slot has been skipped yet
        with patch.object(mail, "now_ist", return_value=dt.datetime(2026, 10, 18, 7, 0, tzinfo=IST)):
            mail.save_settings(self.db, APP, GOOD)   # Mon, Wed, Sat at 08:30
        return self.row()

    def at(self, y, m, d, hh, mm):
        return dt.datetime(y, m, d, hh, mm, tzinfo=IST)

    def test_due_only_on_chosen_weekdays_after_the_time(self):
        row = self._enabled()
        # 2026-10-19 is a Monday
        self.assertFalse(mail.due_now(row, self.at(2026, 10, 19, 8, 29)))
        self.assertTrue(mail.due_now(row, self.at(2026, 10, 19, 8, 30)))
        self.assertTrue(mail.due_now(row, self.at(2026, 10, 19, 11, 0)))    # picked up late
        self.assertFalse(mail.due_now(row, self.at(2026, 10, 20, 9, 0)))    # Tuesday
        self.assertTrue(mail.due_now(row, self.at(2026, 10, 21, 9, 0)))     # Wednesday
        self.assertFalse(mail.due_now(row, self.at(2026, 10, 25, 9, 0)))    # Sunday

    def test_not_due_when_switched_off_or_already_attempted_today(self):
        row = self._enabled()
        row.last_attempt_date = dt.date(2026, 10, 19)
        self.assertFalse(mail.due_now(row, self.at(2026, 10, 19, 9, 0)))
        row.last_attempt_date = None
        row.enabled = False
        self.assertFalse(mail.due_now(row, self.at(2026, 10, 19, 9, 0)))

    def test_next_run_skips_to_the_next_chosen_day(self):
        row = self._enabled()
        self.assertEqual(mail.next_run(row, self.at(2026, 10, 19, 7, 0)), self.at(2026, 10, 19, 8, 30))
        self.assertEqual(mail.next_run(row, self.at(2026, 10, 19, 9, 0)), self.at(2026, 10, 19, 9, 0))  # due now
        row.last_attempt_date = dt.date(2026, 10, 19)
        self.assertEqual(mail.next_run(row, self.at(2026, 10, 19, 9, 0)), self.at(2026, 10, 21, 8, 30))
        self.assertEqual(mail.next_run(row, self.at(2026, 10, 22, 9, 0)), self.at(2026, 10, 24, 8, 30))

    def test_enqueue_starts_each_due_report_once_per_day(self):
        self._enabled()
        with patch("app.services.scheduled_report_mail.now_ist", return_value=self.at(2026, 10, 19, 8, 31)), \
                patch("app.jobs.submit_job", return_value="job-1") as submit:
            self.assertEqual(mail.enqueue_due_checks(), [APP])
            self.assertEqual(mail.enqueue_due_checks(), [])    # same day: not again
        self.assertEqual(submit.call_count, 1)
        self.assertEqual(submit.call_args.args[1], "scheduled")
        self.assertEqual(submit.call_args.kwargs["owner_id"], 0)
        self.assertEqual(self.row().last_attempt_date, dt.date(2026, 10, 19))


class SendLogTests(_DbCase):
    """Every run leaves a log entry: when, who asked, to whom, what was sent."""

    def setUp(self):
        super().setUp()
        mail.save_settings(self.db, APP, GOOD)
        self.sent = MagicMock()
        for target, value in (
            ("send_mail", self.sent),
            ("get_system_email_settings", lambda db: {"configured": True, "sender_email": "sys@rdc.in", "app_password": "x"}),
        ):
            patcher = patch.object(mail, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _file(self, name, size=10):
        path = Path(self.tmp.name) / name
        path.write_bytes(b"x" * size)
        return str(path)

    def log(self, **kw):
        self.db.expire_all()
        return mail.log_view(self.db, APP, **kw)

    def test_a_sent_run_records_time_recipients_and_each_message(self):
        a, b = self._file("p1.xlsx", 1234), self._file("p2.xlsx", 99)
        mail.run_mail_job(APP, "manual", lambda *x: mail.MailPayload(
            values={"date": "d"}, attachment_groups=[[a], [b]], detail={"rows": 7},
        ), requested_by="admin@rdc.in")
        log = self.log()
        self.assertEqual(log["total"], 1)
        entry = log["items"][0]
        self.assertEqual((entry["status"], entry["trigger"], entry["requested_by"]), ("sent", "manual", "admin@rdc.in"))
        self.assertEqual((entry["to"], entry["cc"]), (["a@rdc.in"], ["b@rdc.in"]))
        self.assertIsNotNone(entry["started_at"])
        self.assertIsNotNone(entry["finished_at"])
        self.assertEqual([m["subject"] for m in entry["messages"]],
                         ["Report as on d (Part 1 of 2)", "Report as on d (Part 2 of 2)"])
        self.assertEqual(entry["messages"][0]["attachments"], [{"name": "p1.xlsx", "bytes": 1234}])
        self.assertTrue(all(m["sent_at"] for m in entry["messages"]))
        self.assertEqual(entry["detail"]["rows"], 7)
        self.assertIsNone(entry["error"])

    def test_a_scheduled_run_has_no_requester(self):
        mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(values={"date": "d"}))
        entry = self.log()["items"][0]
        self.assertEqual((entry["trigger"], entry["requested_by"]), ("scheduled", None))

    def test_a_failure_part_way_keeps_what_was_already_sent(self):
        a, b = self._file("p1.xlsx"), self._file("p2.xlsx")
        self.sent.side_effect = [None, RuntimeError("SMTP dropped")]
        with self.assertRaises(JobUserError):
            mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(
                values={"date": "d"}, attachment_groups=[[a], [b]],
            ))
        entry = self.log()["items"][0]
        self.assertEqual(entry["status"], "failed")
        self.assertEqual(len(entry["messages"]), 1)          # part 1 did go out
        self.assertIn("SMTP dropped", entry["error"])

    def test_a_held_run_is_logged_with_the_reason_and_intended_recipients(self):
        mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(values={}, hold_reason="2 unmapped"))
        entry = self.log()["items"][0]
        self.assertEqual((entry["status"], entry["error"], entry["messages"]), ("held", "2 unmapped", []))
        self.assertEqual(entry["to"], ["a@rdc.in"])

    def test_a_refused_overlapping_run_is_logged_as_skipped(self):
        row = self.row()
        row.running_since = mail._utcnow()
        self.db.commit()
        with self.assertRaises(JobUserError):
            mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(values={}))
        entry = self.log()["items"][0]
        self.assertEqual(entry["status"], "skipped")
        self.assertIn("still in progress", entry["error"])

    def test_a_build_failure_is_logged_failed(self):
        def build(*_):
            raise RuntimeError("ERP down")
        with self.assertRaises(JobUserError):
            mail.run_mail_job(APP, "scheduled", build)
        entry = self.log()["items"][0]
        self.assertEqual(entry["status"], "failed")
        self.assertIn("ERP down", entry["error"])
        self.assertIsNotNone(entry["finished_at"])

    def test_newest_first_and_paged(self):
        for _ in range(5):
            mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(values={"date": "d"}))
        page = self.log(limit=2, offset=0)
        self.assertEqual(page["total"], 5)
        self.assertEqual(len(page["items"]), 2)
        ids = [item["id"] for item in self.log(limit=10)["items"]]
        self.assertEqual(ids, sorted(ids, reverse=True))
        self.assertEqual(len(self.log(limit=2, offset=4)["items"]), 1)

    def test_a_run_whose_worker_died_shows_as_interrupted(self):
        old = mail._utcnow() - dt.timedelta(hours=3)
        self.db.add(ScheduledReportMailLog(app_key=APP, trigger="scheduled", status="running", started_at=old))
        self.db.commit()
        self.assertEqual(self.log()["items"][0]["status"], "interrupted")

    def test_each_application_sees_only_its_own_log(self):
        mail.run_mail_job(APP, "scheduled", lambda *x: mail.MailPayload(values={"date": "d"}))
        self.assertEqual(mail.log_view(self.db, "gst-invoice-adder")["total"], 0)


class RunMailJobTests(_DbCase):
    def setUp(self):
        super().setUp()
        mail.save_settings(self.db, APP, GOOD)
        self.sent = MagicMock()
        for target, value in (
            ("send_mail", self.sent),
            ("get_system_email_settings", lambda db: {"configured": True, "sender_email": "sys@rdc.in", "app_password": "x"}),
        ):
            patcher = patch.object(mail, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _file(self, name="r.xlsx"):
        path = Path(self.tmp.name) / name
        path.write_bytes(b"xlsx")
        return str(path)

    def test_sends_to_the_default_recipients_and_records_it(self):
        attachment = self._file()

        def build(snapshot, progress_cb, cancel_event):
            return mail.MailPayload(
                values={"date": "19-Oct-2026"}, sections=[("Untagged Summary", "<table></table>")],
                attachment_groups=[[attachment]], detail={"rows": 5},
            )

        result = mail.run_mail_job(APP, "manual", build)
        self.assertEqual(result["status"], "sent")
        kwargs = self.sent.call_args.kwargs
        self.assertEqual(kwargs["to_addresses"], ["a@rdc.in"])
        self.assertEqual(kwargs["cc_addresses"], ["b@rdc.in"])
        self.assertEqual(kwargs["subject"], "Report as on 19-Oct-2026")
        self.assertEqual(kwargs["attachments"], [attachment])
        self.assertIn("Hello 19-Oct-2026", kwargs["html_body"])
        self.assertIn("Untagged Summary", kwargs["html_body"])
        row = self.row()
        self.assertEqual((row.last_status, row.last_trigger), ("sent", "manual"))
        self.assertIsNone(row.running_since)

    def test_held_payload_sends_nothing_and_says_why(self):
        result = mail.run_mail_job(APP, "scheduled", lambda *a: mail.MailPayload(values={}, hold_reason="3 unmapped"))
        self.assertEqual(result["status"], "held")
        self.sent.assert_not_called()
        self.assertEqual((self.row().last_status, self.row().last_error), ("held", "3 unmapped"))

    def test_a_failing_build_is_recorded_and_releases_the_run(self):
        def build(*_):
            raise RuntimeError("ERP down")

        with self.assertRaises(JobUserError):
            mail.run_mail_job(APP, "scheduled", build)
        row = self.row()
        self.assertEqual(row.last_status, "failed")
        self.assertIn("ERP down", row.last_error)
        self.assertIsNone(row.running_since)

    def test_a_second_run_is_refused_while_one_is_in_progress(self):
        row = self.row()
        row.running_since = mail._utcnow()
        self.db.commit()
        with self.assertRaises(JobUserError):
            mail.run_mail_job(APP, "manual", lambda *a: mail.MailPayload(values={}))

    def test_a_stale_run_does_not_block_forever(self):
        row = self.row()
        row.running_since = mail._utcnow() - dt.timedelta(hours=3)
        self.db.commit()
        result = mail.run_mail_job(APP, "manual", lambda *a: mail.MailPayload(values={"date": "x"}))
        self.assertEqual(result["status"], "sent")

    def test_split_report_goes_out_as_numbered_mails(self):
        first, second = self._file("p1.xlsx"), self._file("p2.xlsx")
        mail.run_mail_job(APP, "manual", lambda *a: mail.MailPayload(
            values={"date": "d"}, sections=[("S", "<p>t</p>")], attachment_groups=[[first], [second]],
        ))
        subjects = [call.kwargs["subject"] for call in self.sent.call_args_list]
        self.assertEqual(subjects, ["Report as on d (Part 1 of 2)", "Report as on d (Part 2 of 2)"])
        bodies = [call.kwargs["html_body"] for call in self.sent.call_args_list]
        self.assertIn("<p>t</p>", bodies[0])
        self.assertNotIn("<p>t</p>", bodies[1])

    def test_no_system_email_fails_clearly(self):
        with patch.object(mail, "get_system_email_settings", lambda db: {"configured": False}):
            with self.assertRaises(JobUserError) as caught:
                mail.run_mail_job(APP, "manual", lambda *a: mail.MailPayload(values={"date": "d"}))
        self.assertIn("system email", str(caught.exception))
        self.sent.assert_not_called()

    def test_scratch_files_are_removed_afterwards(self):
        work = Path(self.tmp.name) / "work"
        work.mkdir()
        (work / "a.xlsx").write_bytes(b"x")
        mail.run_mail_job(APP, "manual", lambda *a: mail.MailPayload(values={"date": "d"}, cleanup=[str(work)]))
        self.assertFalse(work.exists())


def _frames():
    """Small stand-ins for the report tables."""
    base = {
        "Location": ["MUMBAI"], "Account Incharges": ["Haresh"], "Count of Invoices": [2],
        "0-15 Days": [1.0], "16-30 Days": [0.0], "31-60 Days": [2.0], "61-90 Days": [0.0],
        "91-120 Days": [0.0], "121-150 Days": [0.0], "151-180 Days": [0.0], "181-360 Days": [0.0],
        "360+ Days": [0.0],
    }
    untagged = pd.DataFrame({**base, "Accounted Outstanding": [3.0], "Above 30 Days": [2.0]})
    below = pd.DataFrame({**base, "Total O/S": [3.0], "Above 30 days": [2.0]})
    return untagged, below


class SummaryTableTests(unittest.TestCase):
    def test_mail_tables_show_the_computed_above_30_column(self):
        untagged, below = _frames()
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "s.xlsx")
            processor.write_summary_workbook(untagged, below, path)
            html = sheet_to_html(path, "Untagged Summary")
            self.assertIn("MUMBAI", html)
            self.assertIn("Haresh", html)
            # Above 30 Days = 3 - 1 - 0 = 2 (a ROUND formula in the workbook)
            self.assertIn(">2<", html)
            self.assertNotIn("=ROUND", html)
            self.assertIn("Below 1k Summary", openpyxl.load_workbook(path).sheetnames)

    def test_mail_tables_carry_no_unit_footnote_row(self):
        untagged, below = _frames()
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "s.xlsx")
            processor.write_summary_workbook(untagged, below, path)
            self.assertNotIn("Lakhs", sheet_to_html(path, "Untagged Summary"))
            sheet = openpyxl.load_workbook(path)["Untagged Summary"]
            self.assertEqual(sheet.max_row, 3)                    # header, one location, Grand Total


class CompactTableTests(unittest.TestCase):
    def test_compact_markup_is_far_smaller_and_keeps_every_value(self):
        untagged, below = _frames()
        many = pd.concat([untagged] * 40, ignore_index=True)
        many["Location"] = [f"LOC {i}" for i in range(len(many))]
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "s.xlsx")
            processor.write_summary_workbook(many, below, path)
            inline = sheet_to_html(path, "Untagged Summary")
            sink: list[str] = []
            compact = sheet_to_html(path, "Untagged Summary", sink, "u")
        self.assertLess(len(compact) + sum(map(len, sink)), len(inline) * 0.4)
        self.assertGreater(len(sink), 0)
        self.assertTrue(all(rule.startswith(".u") for rule in sink))
        for i in (0, 17, 39):
            self.assertIn(f"LOC {i}<", compact)
        self.assertNotIn("style=\"background", compact)     # styles live in the head instead
        self.assertIn('class="u0"', compact)


def _big_workbook(path, rows=4000):
    """A real workbook with no empty cells (like the report's), big enough that
    compacting it makes a clear difference."""
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Ageing")
    ws.append([f"Col {i}" for i in range(12)])
    for r in range(rows):
        ws.append([f"Customer {r % 53}", f"INV{r:07d}", r * 1.37, r % 90, 0, 0, r * 0.5, 0, 0, 0, "-", "Transactions"])
    wb.save(path)


class UntaggedMailBuilderTests(unittest.TestCase):
    SNAPSHOT = {"send_when_unmapped": True}

    def _built(self, path, unmapped=False):
        untagged, below = _frames()
        df = pd.DataFrame({"x": [1]})

        def fake_build(output_path, allow_unmapped=False, progress_cb=None, cancel_event=None):
            if unmapped and not allow_unmapped:
                return BuiltReport(result={
                    "needs_mapping_fix": True, "total_rows": 1,
                    "validation_warnings": [{"category": "Location Not Mapped", "items": ["A", "B"]}],
                })
            _big_workbook(output_path)
            return BuiltReport(
                result={"needs_mapping_fix": False, "total_rows": 1, "validation_warnings": []},
                as_on_date=dt.date(2026, 10, 19), frames=(df, untagged, df, below, df),
            )
        return fake_build

    def test_body_has_untagged_then_below_1k_and_the_workbook_is_attached(self):
        with patch.object(untagged_mail.report_job, "build_report", self._built(None)):
            payload = untagged_mail.build(self.SNAPSHOT)
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])
        self.assertEqual([title for title, _ in payload.sections], ["Untagged Summary", "Below 1k Summary"])
        self.assertEqual(len(payload.attachment_groups), 1)
        self.assertTrue(payload.attachment_groups[0][0].endswith(f"Untagged_Invoices_Report_As_On_{dt.date.today().isoformat()}.xlsx"))
        self.assertEqual(payload.values["date"], "19-Oct-2026")
        self.assertEqual(payload.notes, [])
        html = mail._wrap_html("Hi", payload.sections, payload.notes, "", payload.css, payload.intro_note)
        self.assertIn("<style>.u0{", html)
        self.assertIn('<meta charset="utf-8">', html)
        self.assertIn(".b0{", html)
        self.assertLess(html.index("</style>"), html.index("<body"))

    def test_the_lakhs_unit_is_one_bold_line_in_the_mail_text_not_a_table_row(self):
        with patch.object(untagged_mail.report_job, "build_report", self._built(None)):
            payload = untagged_mail.build(self.SNAPSHOT)
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])
        html = mail._wrap_html("Dear Team,", payload.sections, payload.notes, "", payload.css, payload.intro_note)
        self.assertEqual(html.count("Lakhs"), 1)                  # stated once, nowhere in the tables
        self.assertIn("<b>Note: All amounts in Lakhs (÷ 1,00,000)</b>", html)
        self.assertLess(html.index("Dear Team,"), html.index("Lakhs"))
        self.assertLess(html.index("Lakhs"), html.index("Untagged Summary"))

    def _sizes(self):
        with tempfile.TemporaryDirectory() as tmp:
            full = Path(tmp) / "f.xlsx"
            _big_workbook(full)
            compact = Path(tmp) / "c.xlsx"
            from app.services import xlsx_compact
            xlsx_compact.compact_xlsx(full, compact)
            return full.stat().st_size, compact.stat().st_size

    def _build(self, limit):
        with patch.object(untagged_mail.report_job, "build_report", self._built(None)), \
                patch.object(mail, "MAX_ATTACHMENT_BYTES", limit):
            payload = untagged_mail.build(self.SNAPSHOT)
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])
        return payload

    def test_an_oversized_workbook_is_attached_in_full_after_compacting(self):
        full, compact = self._sizes()
        self.assertLess(compact, full * 0.7)
        payload = self._build((full + compact) // 2)
        attached = payload.attachment_groups[0][0]
        self.assertTrue(attached.endswith(f"Untagged_Invoices_Report_As_On_{dt.date.today().isoformat()}.xlsx"))
        self.assertNotIn("Without_Full_Ageing", attached)
        self.assertEqual(payload.notes, [])                       # nothing was left out, so no apology
        self.assertIn("compacted", payload.detail["attachment"])
        self.assertLessEqual(Path(attached).stat().st_size, (full + compact) // 2)
        sheet = openpyxl.load_workbook(attached)["Ageing"]
        self.assertEqual(sheet.max_row, 4001)                     # every row is still there
        self.assertEqual(sheet["B4001"].value, "INV0003999")
        self.assertEqual(payload.downloads, [("Untagged Invoices Report", attached)])

    def test_if_even_the_compact_file_is_too_big_it_is_split_and_no_row_is_lost(self):
        full, compact = self._sizes()
        ageing = pd.DataFrame({"Customer": [f"C{i}" for i in range(4000)], "Amount": [i * 1.37 for i in range(4000)]})
        untagged, below = _frames()
        small = pd.DataFrame({"x": [1]})

        def fake_build(output_path, allow_unmapped=False, progress_cb=None, cancel_event=None):
            _big_workbook(output_path)
            return BuiltReport(
                result={"needs_mapping_fix": False, "total_rows": 4000, "validation_warnings": []},
                as_on_date=dt.date(2026, 10, 19), frames=(ageing, untagged, small, below, small),
            )

        limit = compact * 6 // 10                                 # forces the split
        with patch.object(untagged_mail.report_job, "build_report", fake_build), \
                patch.object(mail, "MAX_ATTACHMENT_BYTES", limit):
            payload = untagged_mail.build(self.SNAPSHOT)
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])

        files = [group[0] for group in payload.attachment_groups]
        self.assertGreaterEqual(len(files), 2)
        self.assertTrue(all(Path(f).stat().st_size <= limit for f in files))
        self.assertEqual(
            openpyxl.load_workbook(files[0]).sheetnames,
            ["Untagged Summary", "Untagged Ageing", "Below 1k Summary", "Below 1k Ageing"],
        )
        seen = []
        for path in files[1:]:
            sheet = openpyxl.load_workbook(path)["Ageing"]
            seen += [row[0].value for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(seen, [f"C{i}" for i in range(4000)])    # every Ageing row, once, in order
        self.assertTrue(any("nothing has been left out" in note.lower() for note in payload.notes))
        self.assertEqual(payload.downloads[0][0], "Untagged Invoices Report")   # the complete file stays downloadable

    def test_missing_mappings_hold_the_mail_when_that_is_the_setting(self):
        with patch.object(untagged_mail.report_job, "build_report", self._built(None, unmapped=True)):
            payload = untagged_mail.build({"send_when_unmapped": False})
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])
        self.assertIn("Location Not Mapped (2): A, B", payload.hold_reason)
        self.assertEqual(payload.attachment_groups, [])


class GstMailBuilderTests(unittest.TestCase):
    def _fake_build(self, rows):
        df = pd.DataFrame({
            "Invoice No/ Receipt No": [str(i) for i in range(rows)],
            "Invoice/ Receipt Date": [pd.Timestamp("2026-01-01")] * rows,
            "Accounted Outstanding": [float(i) for i in range(rows)],
        })
        gst_map = {}

        def build_report(output_path, download_name, progress_cb=None, cancel_event=None):
            report_job = gst_mail.report_job
            total, found, blank, _ = report_job.write_workbook(output_path, df, gst_map)
            return report_job.BuiltGstReport(
                result={"total": total, "found": found, "blank": blank, "output_path": output_path,
                        "download_name": download_name, "log": []},
                df=df, gst_map=gst_map,
            )
        return build_report

    def _run(self, rows, limit):
        direct = lambda fn, *a, **k: fn(*a, **k)
        with patch.object(gst_mail.report_job, "build_report", self._fake_build(rows)), \
                patch.object(gst_mail, "run_cpu_phase", direct), \
                patch.object(mail, "MAX_ATTACHMENT_BYTES", limit), \
                patch.object(gst_mail, "_TARGET_PART_BYTES", int(limit * 0.85)):
            payload = gst_mail.build({"send_when_unmapped": True})
        self.addCleanup(lambda: [shutil.rmtree(p, ignore_errors=True) for p in payload.cleanup])
        return payload

    def test_a_report_that_fits_is_one_mail_with_the_counts(self):
        payload = self._run(20, 10 ** 9)
        self.assertEqual(len(payload.attachment_groups), 1)
        self.assertEqual(payload.values["total"], "20")
        self.assertIn("20", payload.sections[0][1])
        self.assertTrue(payload.attachment_groups[0][0].endswith(".xlsx"))

    def _whole_sizes(self):
        direct = lambda fn, *a, **k: fn(*a, **k)
        from app.services import xlsx_compact
        with patch.object(gst_mail.report_job, "build_report", self._fake_build(3000)), \
                patch.object(gst_mail, "run_cpu_phase", direct):
            payload = self._run(3000, 10 ** 9)
        standard = Path(payload.attachment_groups[0][0]).stat().st_size
        with tempfile.TemporaryDirectory() as tmp:
            compact = Path(tmp) / "c.xlsx"
            xlsx_compact.compact_xlsx(payload.attachment_groups[0][0], compact)
            return standard, compact.stat().st_size

    def test_an_oversized_report_is_compacted_into_one_mail_when_that_fits(self):
        standard, compact = self._whole_sizes()
        self.assertLess(compact, standard * 0.7)
        payload = self._run(3000, (standard + compact) // 2)
        self.assertEqual(len(payload.attachment_groups), 1)
        self.assertIn("compacted", payload.detail["attachment"])
        self.assertEqual(openpyxl.load_workbook(payload.attachment_groups[0][0]).active.max_row, 3001)

    def test_only_if_even_the_compact_file_is_too_big_is_it_split_into_parts(self):
        standard, compact = self._whole_sizes()
        limit = int(compact * 0.6)
        payload = self._run(3000, limit)
        self.assertGreater(len(payload.attachment_groups), 1)
        total_rows = 0
        for (path,) in payload.attachment_groups:
            self.assertLessEqual(Path(path).stat().st_size, limit)
            total_rows += openpyxl.load_workbook(path).active.max_row - 1
        self.assertEqual(total_rows, 3000)    # nothing lost or repeated
        self.assertEqual(len(payload.downloads), len(payload.attachment_groups))


class AutoMailRouteTests(_DbCase):
    def setUp(self):
        super().setUp()
        from app import auth
        from app.routers import auto_mail_routes
        self.routes_module = auto_mail_routes
        self.router = auto_mail_routes.make_auto_mail_router(APP)
        self.endpoints = {(r.path, tuple(sorted(r.methods))): r.endpoint for r in self.router.routes}
        self.auth = auth
        patcher = patch.object(mail, "get_system_email_settings", lambda db: {"configured": True})
        patcher.start()
        self.addCleanup(patcher.stop)

    def endpoint(self, path, method):
        return self.endpoints[(path, (method,))]

    def test_every_route_is_administrators_only(self):
        self.assertTrue(any(dep.dependency is self.auth.require_admin for dep in self.router.dependencies))

    def test_there_is_no_way_to_send_without_a_preview_from_here(self):
        self.assertNotIn(("/auto-mail/send-now", ("POST",)), self.endpoints)

    def test_save_then_read_back(self):
        body = self.routes_module.AutoMailBody(**GOOD)
        saved = self.endpoint("/auto-mail", "PUT")(body, self.db)
        read = self.endpoint("/auto-mail", "GET")(self.db)
        self.assertEqual(saved["days"], read["days"])
        self.assertEqual(read["to"], ["a@rdc.in"])

    def test_a_bad_save_is_a_400_with_the_reason(self):
        from fastapi import HTTPException
        body = self.routes_module.AutoMailBody(**{**GOOD, "days": []})
        with self.assertRaises(HTTPException) as caught:
            self.endpoint("/auto-mail", "PUT")(body, self.db)
        self.assertEqual(caught.exception.status_code, 400)
        self.assertIn("at least one day", caught.exception.detail)

    def test_the_log_route_returns_the_history(self):
        self.db.add(ScheduledReportMailLog(app_key=APP, trigger="manual", status="sent", requested_by="a@rdc.in",
                                           started_at=dt.datetime(2026, 10, 19, 3, 0)))
        self.db.commit()
        result = self.endpoint("/auto-mail/log", "GET")(20, 0, self.db)
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["items"][0]["requested_by"], "a@rdc.in")
        self.assertTrue(result["items"][0]["started_at"].startswith("2026-10-19T03:00"))


class RegistrationTests(unittest.TestCase):
    def test_the_mail_jobs_are_allowed_to_run_from_the_queue(self):
        from app import jobs
        for task in (
            "app.services.untagged_invoices.auto_mail:run_auto_mail_job",
            "app.services.gst_invoice_adder.auto_mail:run_auto_mail_job",
            "app.services.untagged_invoices.auto_mail:run_preview_job",
            "app.services.gst_invoice_adder.auto_mail:run_preview_job",
        ):
            self.assertIn(task, jobs._ALLOWED_TASKS)

    def test_both_applications_appear_in_the_admin_recipient_list(self):
        from app.permissions import REPORT_RECIPIENT_APP_KEYS
        self.assertLessEqual({"untagged-invoices-report", "gst-invoice-adder"}, REPORT_RECIPIENT_APP_KEYS)

    def test_both_applications_expose_the_preview_and_the_admin_routes(self):
        from app.main import app
        paths = app.openapi()["paths"]
        for key in ("untagged-invoices-report", "gst-invoice-adder"):
            for suffix in ("/mail/preview", "/mail/confirm-send", "/mail/defaults", "/auto-mail", "/auto-mail/log"):
                self.assertIn(f"/api/tools/{key}{suffix}", paths)
            self.assertNotIn(f"/api/tools/{key}/auto-mail/send-now", paths)


if __name__ == "__main__":
    unittest.main()
