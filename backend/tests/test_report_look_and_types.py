"""The Debtors Reports Auto Mailer's workbooks: the Unaccounted / MRN / PO look,
real Excel dates and numbers (not text), and every total a live Excel formula."""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.services.gst_invoice_adder import processor as gst_processor
from app.services.mailer_shared import sheet_to_html
from app.services.unapplied_receipts import processor as unapplied
from app.services.untagged_invoices import processor as untagged
from app.services.untagged_invoices.processor import LOCATION_COL

PINK, BLUE = "F4B8C8", "C2D8ED"


def _rgb(cell):
    """The fill colour as RRGGBB (the alpha byte differs between xlsxwriter and openpyxl files)."""
    return cell.fill.fgColor.rgb[-6:]


class _Log:
    def __init__(self):
        self.messages = []

    def put(self, item):
        self.messages.append(item)


def _ageing():
    rows = []
    for i, (outstanding, kind) in enumerate([(-1000.0, "Transactions"), (500.0, "Transactions"),
                                              (750.0, "Transactions"), (-2000.0, "Receipts")]):
        rows.append({
            "Customer Name": f"Customer {i}", "Customer Account": str(24455 + i),
            "Location Name": "LOC A", "Sales Person Name": "Sales", "Credit Limit": 100000.0,
            "Payment Terms": "30", "Invoice No/ Receipt No": f"INV/{i}",
            "Invoice/ Receipt Date": dt.datetime(2026, 1, 1 + i),
            "Type (Transaction/Receipt)": kind, "Invoice Type": "Invoice",
            "Accounted Outstanding": outstanding, "Days O/S": 0,
        })
    return pd.DataFrame(rows)


class UntaggedAccountsLookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name) / "untagged.xlsx"
        as_on = dt.date(2026, 1, 31)
        df = untagged.add_location_column(_ageing(), {"LOC A": "Mumbai"})
        df = untagged.recompute_ageing_buckets(df, as_on, _Log())
        below_ageing, untagged_ageing = untagged.build_below_1k_ageing(df), untagged.build_untagged_ageing(df)
        incharges = {"Mumbai": "Haresh"}
        untagged.write_report(
            df, untagged.build_untagged_summary(untagged_ageing, incharges), untagged_ageing,
            untagged.build_below_1k_summary(below_ageing, incharges), below_ageing,
            self.path, as_on, look="accounts",
        )
        self.wb = load_workbook(self.path)

    def test_summary_tables_have_the_pink_header_and_live_formula_totals(self):
        ws = self.wb["Below 1k Summary"]
        header = ws.cell(row=1, column=1)
        self.assertEqual(_rgb(header), PINK)
        self.assertEqual((header.font.name, header.font.size, header.font.bold), ("Calibri", 11, True))
        self.assertEqual(header.border.left.color.rgb[-6:], "000000")
        self.assertEqual(ws.cell(row=2, column=3).alignment.horizontal, "center")
        last = ws.max_row
        while ws.cell(row=last, column=1).value != "Grand Total":
            last -= 1
        self.assertEqual(_rgb(ws.cell(row=last, column=1)), PINK)
        for column, letter in ((3, "C"), (4, "D"), (5, "E")):
            self.assertEqual(ws.cell(row=last, column=column).value, f"=SUBTOTAL(9,{letter}2:{letter}{last - 1})")
        # "Above 30 Days" is a formula in every row, not a typed-in value
        above = [c for c in ws[1] if c.value == "Above 30 Days"][0].column
        self.assertTrue(str(ws.cell(row=2, column=above).value).startswith("=ROUND("))

    def test_the_lists_have_the_light_blue_header_and_plain_white_rows(self):
        ws = self.wb["Ageing"]
        self.assertEqual(_rgb(ws.cell(row=1, column=1)), BLUE)
        self.assertEqual(_rgb(ws.cell(row=2, column=1)), "FFFFFF")
        self.assertEqual(_rgb(ws.cell(row=3, column=1)), "FFFFFF")   # no banding
        self.assertEqual(ws.cell(row=2, column=1).font.name, "Calibri")

    def test_dates_numbers_and_text_are_real_types_with_proper_formats(self):
        ws = self.wb["Ageing"]
        headers = [c.value for c in ws[1]]
        cell = lambda name: ws.cell(row=2, column=headers.index(name) + 1)   # noqa: E731

        date = cell("Invoice/ Receipt Date")
        self.assertIsInstance(date.value, dt.datetime)
        self.assertEqual(date.number_format, "DD-MMM-YYYY")
        account = cell("Customer Account")
        self.assertEqual((account.value, account.number_format), (24455, "0"))
        self.assertEqual(cell("Payment Terms").value, 30)
        self.assertIsInstance(cell("Accounted Outstanding").value, (int, float))
        self.assertIn("#,##0.00", cell("Accounted Outstanding").number_format)
        self.assertEqual(cell("Invoice No/ Receipt No").value, "INV/0")           # identifiers stay text
        self.assertEqual(cell(LOCATION_COL).value, "Mumbai")

    def test_the_default_look_is_unchanged_for_the_standalone_report(self):
        as_on = dt.date(2026, 1, 31)
        df = untagged.recompute_ageing_buckets(untagged.add_location_column(_ageing(), {"LOC A": "Mumbai"}), as_on, _Log())
        ageing = untagged.build_below_1k_ageing(df)
        path = self.path.with_name("default.xlsx")
        untagged.write_report(df, untagged.build_untagged_summary(untagged.build_untagged_ageing(df), {}),
                              untagged.build_untagged_ageing(df), untagged.build_below_1k_summary(ageing, {}),
                              ageing, path, as_on)
        self.assertEqual(_rgb(load_workbook(path)["Ageing"].cell(row=1, column=1)), "1F3864")

    def test_the_mail_table_carries_the_pink_header(self):
        css = []
        sheet_to_html(str(self.path), "Below 1k Summary", css, "b")
        self.assertIn("f4b8c8", "".join(css).lower())


class UnappliedAccountsLookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = str(Path(tmp.name) / "unapplied.xlsx")
        main = pd.DataFrame({
            "Customer Name": ["A Ltd", "B Ltd"], "Customer Number": ["139554", "150289"],
            "Payment Method": ["HDFC", "HDFC"], "Payment Number": ["ON/1", "ON/2"],
            "Payment Date": ["09-OCT-26", "12-JAN-26"], "Location": ["Mumbai", "Mumbai"],
            "Sales Person": ["S", "S"],
            "On Account/ Prepayment Amount": [0.0, 5970.4], "Unapplied Amount": [40000.0, 0.0],
            "Due Days": [19, 291], "Ageing Bucket": ["16 - 30", "181 - 360"],
        })
        unidentified = main.iloc[:1].copy()
        unidentified["Customer Name"] = "***** Unidentified"
        unapplied.write_formatted_excel(
            main, main.iloc[1:], self.path, dt.date(2026, 10, 28), df_unidentified=unidentified,
            incharge_map={"Mumbai": "Haresh"}, look="accounts",
        )
        self.wb = load_workbook(self.path)

    def test_payment_dates_are_real_dates_on_every_list(self):
        for sheet in ("Unapplied Receipts", "Advance of Customers", "Unidentified Customers"):
            ws = self.wb[sheet]
            headers = [c.value for c in ws[2]]
            cell = ws.cell(row=3, column=headers.index("Payment Date") + 1)
            self.assertIsInstance(cell.value, dt.datetime, sheet)
            self.assertEqual(cell.number_format, "DD-MMM-YYYY", sheet)
        ws = self.wb["Unapplied Receipts"]
        headers = [c.value for c in ws[2]]
        self.assertEqual(ws.cell(row=3, column=headers.index("Payment Date") + 1).value, dt.datetime(2026, 10, 9))

    def test_amounts_and_days_are_numbers_and_every_total_is_a_formula(self):
        ws = self.wb["Unapplied Receipts"]
        headers = [c.value for c in ws[2]]
        unapplied_col = headers.index("Unapplied Amount") + 1
        self.assertEqual(ws.cell(row=3, column=unapplied_col).value, 40000.0)
        self.assertIn("#,##0.00", ws.cell(row=3, column=unapplied_col).number_format)
        self.assertEqual(ws.cell(row=3, column=headers.index("Customer Number") + 1).value, 139554)
        last = ws.max_row
        self.assertEqual(ws.cell(row=last, column=1).value, "Grand Total")
        self.assertEqual(ws.cell(row=last, column=unapplied_col).value, "=SUBTOTAL(9,{0}3:{0}{1})".format(
            ws.cell(row=2, column=unapplied_col).column_letter, last - 1))
        summary = self.wb["Summary"]
        grand_row = summary.max_row
        self.assertTrue(str(summary.cell(row=4, column=3).value).startswith("=SUM("))
        self.assertTrue(str(summary.cell(row=grand_row, column=3).value).startswith("=SUBTOTAL(9,"))

    def test_the_look_is_the_unaccounted_one(self):
        ws = self.wb["Unapplied Receipts"]
        self.assertEqual(_rgb(ws.cell(row=2, column=1)), BLUE)
        self.assertEqual(_rgb(ws.cell(row=3, column=1)), "FFFFFF")
        self.assertEqual(_rgb(ws.cell(row=4, column=1)), "FFFFFF")
        self.assertEqual(_rgb(ws.cell(row=ws.max_row, column=1)), PINK)
        summary = self.wb["Summary"]
        self.assertEqual(_rgb(summary.cell(row=3, column=1)), PINK)
        self.assertEqual(summary.cell(row=4, column=3).alignment.horizontal, "center")
        self.assertEqual(summary.cell(row=4, column=4).font.name, "Calibri")


class GstWorkbookTypesTests(unittest.TestCase):
    def test_the_accounts_look_with_real_types(self):
        df = pd.DataFrame({
            "Customer Account": ["24455", "24456"], "Payment Terms": ["30", ""],
            "Invoice No/ Receipt No": ["INV/1", "INV/2"],
            "Invoice/ Receipt Date": [dt.datetime(2026, 1, 1), dt.datetime(2026, 1, 2)],
            "Accounted Outstanding": [100.5, 200.0],
        })
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "gst.xlsx")
            gst_processor.build_gst_report_pure_python(
                df, "Invoice No/ Receipt No", "Invoice/ Receipt Date",
                {("INV/1", gst_processor._to_oracle_date(dt.datetime(2026, 1, 1))): "GST-1"},
                path, _Log(), look="accounts",
            )
            ws = load_workbook(path).active
        headers = [c.value for c in ws[1]]
        cell = lambda name: ws.cell(row=2, column=headers.index(name) + 1)   # noqa: E731
        self.assertEqual(_rgb(ws.cell(row=1, column=1)), BLUE)
        self.assertEqual((cell("Customer Account").value, cell("Customer Account").number_format), (24455, "0"))
        self.assertEqual(cell("Payment Terms").value, 30)
        self.assertIsInstance(cell("Invoice/ Receipt Date").value, dt.datetime)
        self.assertEqual(cell("Invoice/ Receipt Date").number_format, "DD-MMM-YYYY")
        self.assertEqual(cell("Invoice No/ Receipt No").value, "INV/1")
        self.assertEqual(cell(gst_processor.COL_GST_NEW).value, "GST-1")
        self.assertIsInstance(cell("Accounted Outstanding").value, (int, float))


if __name__ == "__main__":
    unittest.main()
