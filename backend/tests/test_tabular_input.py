"""Any-format spreadsheet input and plain-language errors.

The .xls / .xlsb fixtures in tests/fixtures were produced by Excel itself from
aging_sample.xlsx, so they are genuine BIFF / binary-OOXML files.
"""
import datetime as dt
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

from app.jobs import JobUserError
from app.services import tabular_input as ti
from app.services.untagged_invoices.processor import read_ageing_file

from test_untagged_invoices import _HEADERS, _LogQueue, _base_row, _make_ageing_workbook

FIXTURES = Path(__file__).parent / "fixtures"


def _aging_rows():
    return [
        _base_row(**{"Invoice/ Receipt Date": dt.datetime(2026, 8, 1), "Accounted Outstanding": -1500.0}),
        _base_row(**{"Customer Name": "Beta Ltd", "Invoice/ Receipt Date": dt.datetime(2026, 7, 15), "Accounted Outstanding": 800.5}),
    ]


class DetectFormatTests(unittest.TestCase):
    def test_real_excel_files_are_detected_by_content(self):
        self.assertEqual(ti.detect_format(FIXTURES / "aging_sample.xlsx"), "xlsx")
        self.assertEqual(ti.detect_format(FIXTURES / "aging_sample.xls"), "xls")
        self.assertEqual(ti.detect_format(FIXTURES / "aging_sample.xlsb"), "xlsb")

    def test_extension_is_ignored(self):
        with tempfile.TemporaryDirectory() as tmp:
            disguised = Path(tmp) / "report.xls"          # really an .xlsx
            shutil.copyfile(FIXTURES / "aging_sample.xlsx", disguised)
            self.assertEqual(ti.detect_format(disguised), "xlsx")
            html = Path(tmp) / "report.xlsx"              # really HTML
            html.write_text("<html><table><tr><td>a</td></tr></table></html>")
            self.assertEqual(ti.detect_format(html), "html")
            csv_file = Path(tmp) / "report.xls"           # really CSV
            csv_file.write_text("a,b\n1,2\n")
            self.assertEqual(ti.detect_format(csv_file), "csv")

    def test_unreadable_things_get_a_plain_message(self):
        with tempfile.TemporaryDirectory() as tmp:
            empty = Path(tmp) / "empty.xlsx"
            empty.write_bytes(b"")
            with self.assertRaisesRegex(JobUserError, "is empty"):
                ti.detect_format(empty)
            pdf = Path(tmp) / "x.xlsx"
            pdf.write_bytes(b"%PDF-1.7 stuff")
            with self.assertRaisesRegex(JobUserError, "PDF"):
                ti.detect_format(pdf)
            broken = Path(tmp) / "broken.xlsx"
            broken.write_bytes(b"PK\x03\x04 not really a zip at all")
            with self.assertRaisesRegex(JobUserError, "damaged"):
                ti.detect_format(broken)
            plain_zip = Path(tmp) / "a.xlsx"
            with zipfile.ZipFile(plain_zip, "w") as zf:
                zf.writestr("hello.txt", "hi")
            with self.assertRaisesRegex(JobUserError, "not an Excel workbook"):
                ti.detect_format(plain_zip)
            binary = Path(tmp) / "b.xls"
            binary.write_bytes(b"\x01\x02\x00\x03\x04" * 50)
            with self.assertRaisesRegex(JobUserError, "not a spreadsheet"):
                ti.detect_format(binary)

    def test_upload_prefix_is_not_shown_to_users(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "0b9f2f6e-5f1c-4c52-9b36-3a4f0d6d8a11_my report.xlsx"
            path.write_bytes(b"")
            with self.assertRaisesRegex(JobUserError, r"'my report\.xlsx'"):
                ti.detect_format(path)


class ReadGridTests(unittest.TestCase):
    def test_csv_with_semicolons_ragged_rows_and_blank_lines(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bank.csv"
            path.write_text("Date;Narration;Amount\n01/08/2026;Rent;1,500.00\n\n02/08/2026;Fuel;300;extra\n", encoding="cp1252")
            grid = ti.read_grid(path)
            self.assertEqual(grid.shape, (3, 4))
            self.assertEqual(grid.iloc[0, 1], "Narration")

    def test_sheet_not_found_lists_the_sheets(self):
        with self.assertRaisesRegex(JobUserError, "no sheet named 'Nope'.*'Aging'"):
            ti.read_grid(FIXTURES / "aging_sample.xlsx", sheet="Nope")

    def test_empty_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "blank.xlsx"
            openpyxl.Workbook().save(path)
            with self.assertRaisesRegex(JobUserError, "empty"):
                ti.read_grid(path)

    def test_to_xlsx_converts_other_formats_and_keeps_xlsx_as_is(self):
        with tempfile.TemporaryDirectory() as tmp:
            same = ti.to_xlsx(FIXTURES / "aging_sample.xlsx", tmp)
            self.assertEqual(same, FIXTURES / "aging_sample.xlsx")
            for name in ("aging_sample.xls", "aging_sample.xlsb"):
                converted = ti.to_xlsx(FIXTURES / name, tmp)
                self.assertEqual(converted.suffix, ".xlsx")
                ws = openpyxl.load_workbook(converted).active
                self.assertEqual(ws["A14"].value, "Customer Name")
                self.assertEqual(ws["A15"].value, "Acme Co")


class CleaningTests(unittest.TestCase):
    def test_clean_numeric(self):
        out, bad = ti.clean_numeric(pd.Series(["1,234.50", "(500)", "700-", 12, None, "", "abc", "Rs. 1,000", "-"], dtype=object))
        self.assertEqual(out.tolist()[:4], [1234.5, -500.0, -700.0, 12.0])
        self.assertTrue(pd.isna(out[4]) and pd.isna(out[5]) and pd.isna(out[8]))
        self.assertEqual(out[7], 1000.0)
        self.assertEqual(bad, ["abc"])

    def test_parse_dates_handles_text_serials_and_real_dates(self):
        out, bad = ti.parse_dates(pd.Series(["31/08/2026", "01-Aug-2026", 46000, dt.datetime(2026, 1, 2), "not a date", None], dtype=object))
        self.assertEqual(out[0], pd.Timestamp(2026, 8, 31))
        self.assertEqual(out[1], pd.Timestamp(2026, 8, 1))
        self.assertEqual(out[2], pd.Timestamp("1899-12-30") + pd.Timedelta(days=46000))
        self.assertEqual(out[3], pd.Timestamp(2026, 1, 2))
        self.assertTrue(pd.isna(out[4]) and pd.isna(out[5]))
        self.assertEqual(bad, ["not a date"])


class AgeingReaderAnyFormatTests(unittest.TestCase):
    def test_xlsx_xls_xlsb_give_identical_results(self):
        frames = {ext: read_ageing_file(str(FIXTURES / f"aging_sample.{ext}"), _LogQueue()) for ext in ("xlsx", "xls", "xlsb")}
        for ext in ("xls", "xlsb"):
            self.assertEqual(frames[ext]["Accounted Outstanding"].tolist(), frames["xlsx"]["Accounted Outstanding"].tolist())
            self.assertEqual(frames[ext]["Invoice/ Receipt Date"].tolist(), frames["xlsx"]["Invoice/ Receipt Date"].tolist())
            self.assertEqual(frames[ext]["Customer Name"].tolist(), frames["xlsx"]["Customer Name"].tolist())

    def test_csv_and_html_versions_work_too(self):
        with tempfile.TemporaryDirectory() as tmp:
            rows = [[f"meta {i}"] for i in range(1, 14)] + [_HEADERS] + [
                [("" if h not in r else r[h]) for h in _HEADERS] for r in _aging_rows()
            ]
            csv_path = Path(tmp) / "aging.csv"
            pd.DataFrame(rows).to_csv(csv_path, header=False, index=False)
            df = read_ageing_file(str(csv_path), _LogQueue())
            self.assertEqual(len(df), 2)
            self.assertEqual(df["Accounted Outstanding"].tolist(), [-1500.0, 800.5])

            html_path = Path(tmp) / "aging.xls"
            cells = "".join("<tr>" + "".join(f"<td>{'' if v is None else v}</td>" for v in row) + "</tr>" for row in rows)
            html_path.write_text(f"<html><body><table>{cells}</table></body></html>")
            df = read_ageing_file(str(html_path), _LogQueue())
            self.assertEqual(len(df), 2)
            self.assertEqual(df["Invoice/ Receipt Date"].dt.month.tolist(), [8, 7])

    def test_header_row_can_be_anywhere(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "aging.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["title"])
            ws.append(_HEADERS)
            ws.append([_base_row().get(h) for h in _HEADERS])
            wb.save(path)
            self.assertEqual(len(read_ageing_file(str(path), _LogQueue())), 1)


class AgeingReaderErrorTests(unittest.TestCase):
    def _build(self, tmp, mutate):
        path = Path(tmp) / "aging.xlsx"
        _make_ageing_workbook(path, _aging_rows())
        wb = openpyxl.load_workbook(path)
        mutate(wb.active)
        wb.save(path)
        return str(path)

    def test_missing_columns_are_named(self):
        def drop(ws):
            col = _HEADERS.index("Accounted Outstanding") + 1
            ws.cell(row=14, column=col, value="Balance")
            col = _HEADERS.index("Location Name") + 1
            ws.cell(row=14, column=col, value="Branch")
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(JobUserError) as ctx:
                read_ageing_file(self._build(tmp, drop), _LogQueue())
            msg = str(ctx.exception)
            self.assertIn("Location Name", msg)
            self.assertIn("Accounted Outstanding", msg)
            self.assertIn("Columns found", msg)

    def test_wrong_file_has_no_header(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "other.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Item", "Qty"])
            wb.active.append(["Pen", 3])
            wb.save(path)
            with self.assertRaisesRegex(JobUserError, "doesn't look like the Aging export"):
                read_ageing_file(str(path), _LogQueue())

    def test_all_amounts_unreadable(self):
        def junk(ws):
            col = _HEADERS.index("Accounted Outstanding") + 1
            for r in (15, 16):
                ws.cell(row=r, column=col, value="n/a")
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(JobUserError, "not a number"):
                read_ageing_file(self._build(tmp, junk), _LogQueue())

    def test_all_dates_unreadable(self):
        def junk(ws):
            col = _HEADERS.index("Invoice/ Receipt Date") + 1
            for r in (15, 16):
                ws.cell(row=r, column=col, value="last week")
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(JobUserError, "could be read as a date"):
                read_ageing_file(self._build(tmp, junk), _LogQueue())

    def test_some_bad_values_only_warn(self):
        def junk(ws):
            ws.cell(row=15, column=_HEADERS.index("Accounted Outstanding") + 1, value="oops")
            ws.cell(row=16, column=_HEADERS.index("Invoice/ Receipt Date") + 1, value="never")
        with tempfile.TemporaryDirectory() as tmp:
            log = _LogQueue()
            df = read_ageing_file(self._build(tmp, junk), log)
            self.assertEqual(len(df), 2)
            warnings = " ".join(m for tag, m in log.messages if tag == "warn")
            self.assertIn("oops", warnings)
            self.assertIn("never", warnings)

    def test_no_locations_at_all(self):
        def blank(ws):
            col = _HEADERS.index("Location Name") + 1
            for r in (15, 16):
                ws.cell(row=r, column=col).value = None
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(JobUserError, "empty for every row"):
                read_ageing_file(self._build(tmp, blank), _LogQueue())

    def test_header_only_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "aging.xlsx"
            _make_ageing_workbook(path, [], include_grand_total=False)
            with self.assertRaisesRegex(JobUserError, "no data rows"):
                read_ageing_file(str(path), _LogQueue())

    def test_blank_type_and_location_cells_do_not_break_the_report_steps(self):
        from app.services.untagged_invoices.processor import (
            add_location_column, build_below_1k_ageing, build_untagged_ageing, missing_location_mappings)

        def blank_one(ws):
            ws.cell(row=15, column=_HEADERS.index("Type (Transaction/Receipt)") + 1).value = None
            ws.cell(row=16, column=_HEADERS.index("Location Name") + 1).value = None
        with tempfile.TemporaryDirectory() as tmp:
            df = read_ageing_file(self._build(tmp, blank_one), _LogQueue())
            self.assertEqual(missing_location_mappings(df, {"COIMBATORE - SARAVANAMPATTI": "X"}), [])
            df = add_location_column(df, {"COIMBATORE - SARAVANAMPATTI": "X"})
            self.assertEqual(len(build_untagged_ageing(df)), 0)
            self.assertEqual(len(build_below_1k_ageing(df)), 1)   # Beta Ltd 800.50, location blank


class ConvertedWorkbookFeedsOpenpyxlAppsTests(unittest.TestCase):
    """Apps that only know openpyxl (Tally reports) get an equivalent .xlsx."""

    def test_xls_keeps_the_dr_cr_number_formats_tally_relies_on(self):
        with tempfile.TemporaryDirectory() as tmp:
            converted = ti.to_xlsx(FIXTURES / "tally_export.xls", tmp)
            ws = openpyxl.load_workbook(converted)["Sheet2"]
            self.assertIn("Cr", ws["F3"].number_format)
            self.assertIn("Dr", ws["F5"].number_format)
            self.assertEqual(ws["A3"].value.date(), dt.date(2026, 8, 1))

    def test_lossy_formats_warn_about_number_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            log = _LogQueue()
            ti.to_xlsx(FIXTURES / "tally_export.xlsb", tmp, log)
            self.assertTrue(any(tag == "warn" and "Dr/Cr" in m for tag, m in log.messages))

    def test_creditors_ageing_gives_the_same_answer_for_xlsx_and_xls(self):
        from app.services.creditors_ageing import processor
        mapping = {
            "VENDOR A": {"name": "Vendor A", "loc": "Goa", "vt": "Expenses", "vst": "Testing", "intercompany": False},
            "VENDOR B": {"name": "Vendor B", "loc": "Wada", "vt": "Advance", "vst": "Regular", "intercompany": False},
        }
        results = {}
        for ext in ("xlsx", "xls", "xlsb"):
            with tempfile.TemporaryDirectory() as tmp:
                source = ti.to_xlsx(FIXTURES / f"tally_export.{ext}", tmp)
                results[ext] = processor.generate_report(str(source), str(Path(tmp) / "r.xlsx"), mapping, as_on_date="2026-08-28")["counts"]
        self.assertEqual(results["xls"], results["xlsx"])
        self.assertEqual(results["xlsb"], results["xlsx"])


if __name__ == "__main__":
    unittest.main()
