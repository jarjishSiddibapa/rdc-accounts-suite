"""The openpyxl style-assignment speed-up must not change any written file."""

import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]

SCRIPT = r'''
import sys
sys.path.insert(0, {backend!r})
if {fast!r}:
    from app import openpyxl_speedup
    openpyxl_speedup.install()
    from openpyxl.styles import styleable
    assert styleable.StyleDescriptor.__set__.__name__ == "fast_set"
from copy import copy
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="1F3864")
side = Side(style="thin", color="D0D8E4")
border = Border(left=side, right=side, top=side, bottom=side)
align = Alignment(horizontal="left", vertical="center", indent=1)
for _ in range(2):                                   # two workbooks: caches must not leak between them
    wb = Workbook()
    ws = wb.active
    for r in range(1, 60):
        for c in range(1, 6):
            cell = ws.cell(r, c, "x")
            cell.font, cell.fill, cell.border, cell.alignment = font, fill, border, align
            cell.number_format = "0.00"
    ws["A1"].font = font
    other = Font(name="Consolas", size=9)
    ws["A3"].font = other
    ws["B3"].font = copy(ws["A3"].font)              # a style copied from another cell
    ws["C3"].font = other
    if {mutate!r}:
        # An object changed after it was used: openpyxl keeps it by reference and then lines it up
        # by hash, which is not stable from one run to the next, so only check this does not break.
        font.bold = False
        ws["A2"].font = font
        other.italic = True
        ws["D3"].font = other
        font.bold = True
wb.save({out!r})
'''


class OpenpyxlSpeedupTests(unittest.TestCase):
    def _build(self, fast: bool, directory: str, mutate: bool = False) -> Path:
        out = Path(directory) / ("fast.xlsx" if fast else "plain.xlsx")
        subprocess.run(
            [sys.executable, "-I", "-c", SCRIPT.format(backend=str(BACKEND), fast=fast, out=str(out), mutate=mutate)],
            check=True, capture_output=True, timeout=120,
        )
        return out

    def test_written_styles_and_cells_are_identical_with_and_without_it(self):
        with tempfile.TemporaryDirectory() as directory:
            fast, plain = self._build(True, directory), self._build(False, directory)
            with zipfile.ZipFile(fast) as a, zipfile.ZipFile(plain) as b:
                for part in ("xl/styles.xml", "xl/worksheets/sheet1.xml"):
                    self.assertEqual(a.read(part), b.read(part), part)

    def test_an_object_changed_after_use_still_writes_a_loadable_workbook(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            ws = load_workbook(self._build(True, directory, mutate=True)).active
            self.assertEqual(ws["A3"].font.name, "Consolas")
            self.assertEqual(ws["B3"].font.name, "Consolas")


if __name__ == "__main__":
    unittest.main()
