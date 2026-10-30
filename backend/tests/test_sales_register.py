"""RDC Sales Register Report: the reporting period, the workbook written exactly
as the ERP returned the rows, always one file, admin-only access, and the
richer schedule and sender choice it comes with."""

import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import permissions, security
from app.database import Base
from app.jobs import JobUserError
from app.models import EmailSettings, ScheduledReportMail, SystemEmailSettings, User
from app.regional import IST
from app.services import scheduled_report_mail as mail
from app.services.sales_register import auto_mail, report
from app.services.sales_register.sql import SALES_REGISTER_SQL


class PeriodTests(unittest.TestCase):
    def test_a_normal_day_reports_the_month_so_far_through_yesterday(self):
        self.assertEqual(report.get_report_periods(date(2026, 10, 30)),
                         [("Current Month MTD", date(2026, 10, 1), date(2026, 10, 29))])
        self.assertEqual(report.get_report_periods(date(2026, 10, 2)),
                         [("Current Month MTD", date(2026, 10, 1), date(2026, 10, 1))])

    def test_the_first_of_a_month_reports_the_whole_previous_month_and_nothing_else(self):
        self.assertEqual(report.get_report_periods(date(2026, 10, 1)),
                         [("Previous Month", date(2026, 9, 1), date(2026, 9, 30))])
        self.assertEqual(report.get_report_periods(date(2027, 1, 1)),
                         [("Previous Month", date(2026, 12, 1), date(2026, 12, 31))])
        self.assertEqual(report.get_report_periods(date(2028, 3, 1)),
                         [("Previous Month", date(2028, 2, 1), date(2028, 2, 29))])

    def test_there_is_always_at_most_one_period_so_one_file(self):
        day = date(2026, 1, 1)
        for _ in range(400):
            self.assertEqual(len(report.get_report_periods(day)), 1)
            day += timedelta(days=1)

    def test_the_query_is_the_supplied_one_with_its_binds(self):
        for bind in (":p_from_date", ":p_to_date + 1", ":p_org_id", ":p_trkinv", ":p_location",
                     ":p_cust_name", ":p_cust_number"):
            self.assertIn(bind, SALES_REGISTER_SQL)
        self.assertIn("rct.complete_flag = 'Y'", SALES_REGISTER_SQL)
        self.assertEqual(report.P_ORG_ID, 102)


class _FakeCursor:
    def __init__(self, columns, rows):
        self.description = [(c,) for c in columns]
        self.rows = list(rows)
        self.executed = None
        self.arraysize = 0
        self.prefetchrows = 0

    def execute(self, sql, **binds):
        self.executed = (sql, binds)

    def fetchmany(self, size):
        batch, self.rows = self.rows[:size], self.rows[size:]
        return batch

    def close(self):
        pass


class _FakeConnection:
    def __init__(self, cursor):
        self._cursor = cursor
        self.outputtypehandler = None

    def cursor(self):
        return self._cursor

    def close(self):
        pass


def _run_build(columns, rows, path, **patches):
    cursor = _FakeCursor(columns, rows)
    with patch.object(report, "_init_oracle_client"), \
            patch.object(report, "_connect_oracle", lambda cfg: _FakeConnection(cursor)), \
            patch.object(report, "_FETCH_ROWS", 2):
        summary = report.build_workbook(SimpleNamespace(instant_client_dir=""), date(2026, 10, 1),
                                        date(2026, 10, 29), str(path))
    return summary, cursor


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def test_every_value_is_written_exactly_as_the_erp_returned_it(self):
        columns = ["INVOICE_NUMBER", "INVOICE_DATE", "LINE_AMOUNT", "QTY", "SALES_ORDER_DATE", "NOTE", "GSTIN"]
        stamp = datetime(2026, 10, 5, 13, 45, 7)
        rows = [
            ("0012345", "05-OCT-2026", 1234.5, 3, stamp, "=SUM(A1:A2)", None),       # text that looks like a number / formula
            ("INV/9", "06-OCT-2026", -0.1, 0, None, "http://example.com", "27AAAAA0000A1Z5"),
            ("X" * 300, "07-OCT-2026", 10, 1, stamp, "multi\nline", ""),
        ]
        path = self.dir / "r.xlsx"
        summary, cursor = _run_build(columns, rows, path)
        self.assertEqual(summary, {"rows": 3, "columns": 7, "sheets": 1})

        ws = openpyxl.load_workbook(path)[report.SHEET_NAME]
        self.assertEqual([c.value for c in ws[1]], columns)
        first, second, third = ([c.value for c in ws[i]] for i in (2, 3, 4))
        self.assertEqual(first, ["0012345", "05-OCT-2026", 1234.5, 3, stamp, "=SUM(A1:A2)", None])
        self.assertIsInstance(first[0], str)               # leading zeros kept: text stays text
        self.assertEqual(first[5], "=SUM(A1:A2)")           # never turned into a formula
        self.assertEqual(ws["F2"].data_type, "s")
        self.assertEqual(second, ["INV/9", "06-OCT-2026", -0.1, 0, None, "http://example.com", "27AAAAA0000A1Z5"])
        self.assertIsNone(ws["F3"].hyperlink)               # nor into a link
        self.assertEqual(third[0], "X" * 300)
        self.assertEqual(third[5], "multi\nline")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:G4")

    def test_the_query_runs_with_the_fixed_org_and_the_dates_and_no_other_filter(self):
        path = self.dir / "r.xlsx"
        _, cursor = _run_build(["A"], [("x",)], path)
        sql, binds = cursor.executed
        self.assertIs(sql, SALES_REGISTER_SQL)
        self.assertEqual(binds, {
            "p_from_date": datetime(2026, 10, 1), "p_to_date": datetime(2026, 10, 29),
            "p_trkinv": None, "p_location": None, "p_cust_name": None, "p_cust_number": None, "p_org_id": 102,
        })

    def test_an_empty_period_still_gives_a_workbook_with_the_headings(self):
        path = self.dir / "r.xlsx"
        summary, _ = _run_build(["A", "B"], [], path)
        self.assertEqual(summary["rows"], 0)
        ws = openpyxl.load_workbook(path)[report.SHEET_NAME]
        self.assertEqual([c.value for c in ws[1]], ["A", "B"])
        self.assertEqual(ws.max_row, 1)

    def test_more_rows_than_a_sheet_hold_continue_on_another_sheet_of_the_same_file(self):
        rows = [(i,) for i in range(1, 8)]
        path = self.dir / "r.xlsx"
        with patch.object(report, "_MAX_SHEET_ROWS", 4):         # header + 3 rows a sheet
            summary, _ = _run_build(["N"], rows, path)
        self.assertEqual((summary["rows"], summary["sheets"]), (7, 3))
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, [report.SHEET_NAME, f"{report.SHEET_NAME} (2)", f"{report.SHEET_NAME} (3)"])
        values = [c.value for name in wb.sheetnames for (c,) in wb[name].iter_rows(min_row=2)]
        self.assertEqual(values, list(range(1, 8)))
        self.assertEqual([wb[n].cell(1, 1).value for n in wb.sheetnames], ["N", "N", "N"])

    def test_lobs_are_read_as_text_and_unwritable_characters_dropped(self):
        self.assertEqual(report._cell(b"caf\xc3\xa9"), "café")
        self.assertEqual(report._cell("a\ud800b￾"), "ab")
        self.assertEqual(report._cell("plain"), "plain")
        self.assertEqual(report._cell(12.5), 12.5)
        self.assertIsNone(report._cell(None))


class OneFileTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)

    def _file(self, size):
        path = self.dir / "r.xlsx"
        path.write_bytes(b"x" * size)
        return path

    def test_a_report_within_the_limit_goes_as_it_is(self):
        path = self._file(100)
        self.assertEqual(auto_mail.fit_in_one_file(path), (path, "complete workbook"))

    def test_too_big_but_fine_once_compacted(self):
        path = self._file(2_000_000)
        with patch.object(mail, "MAX_ATTACHMENT_BYTES", 1_000_000), \
                patch.object(auto_mail.xlsx_compact, "compact_xlsx_in_place",
                             lambda p: Path(p).write_bytes(b"y" * 500)):
            final, how = auto_mail.fit_in_one_file(path)
        self.assertEqual(final, path)
        self.assertIn("compacted", how)

    def test_still_too_big_becomes_xlsb(self):
        path = self._file(2_000_000)

        def fake_xlsb(source, destination):
            Path(destination).write_bytes(b"z" * 300)
            return True

        with patch.object(mail, "MAX_ATTACHMENT_BYTES", 1_000_000), \
                patch.object(auto_mail.xlsx_compact, "compact_xlsx_in_place"), \
                patch.object(auto_mail.xlsb_convert, "convert_to_xlsb", fake_xlsb):
            final, how = auto_mail.fit_in_one_file(path)
        self.assertEqual(final.suffix, ".xlsb")
        self.assertIn(".xlsb", how)

    def test_when_even_that_is_too_big_it_is_never_split_it_stops(self):
        path = self._file(2_000_000)

        def fake_xlsb(source, destination):
            Path(destination).write_bytes(b"z" * 1_500_000)
            return True

        with patch.object(mail, "MAX_ATTACHMENT_BYTES", 1_000_000), \
                patch.object(auto_mail.xlsx_compact, "compact_xlsx_in_place"), \
                patch.object(auto_mail.xlsb_convert, "convert_to_xlsb", fake_xlsb):
            with self.assertRaises(JobUserError) as caught:
                auto_mail.fit_in_one_file(path)
        self.assertIn("never split", str(caught.exception))

    def test_the_real_converter_keeps_values_types_and_dates(self):
        """A real workbook through the real compaction and .xlsb conversion."""
        import pyxlsb
        columns = ["INV", "AMT", "WHEN", "NOTE"]
        stamp = datetime(2026, 10, 5, 13, 45)
        rows = [("0012", 1234.5, stamp, "a & b <c>"), ("7", 10, None, "=1+1"), ("", 0, stamp, "multi\nline")]
        path = self.dir / "real.xlsx"
        _run_build(columns, rows, path)
        xlsx_values = [[c.value for c in row] for row in openpyxl.load_workbook(path).active.iter_rows()]

        from app.services import xlsb_writer
        binary = self.dir / "real.xlsb"
        xlsb_writer.convert(str(path), str(binary))
        with pyxlsb.open_workbook(str(binary)) as wb, wb.get_sheet(1) as sheet:
            got = [[c.v for c in row] for row in sheet.rows()]
        self.assertEqual([r[0] for r in got], [r[0] for r in xlsx_values])
        self.assertEqual([r[1] for r in got][1:], [1234.5, 10.0, 0.0])
        self.assertEqual([r[3] for r in got][1:], ["a & b <c>", "=1+1", "multi\nline"])


class AdminOnlyTests(unittest.TestCase):
    KEY = "rdc-sales-register-report"

    def test_it_is_a_registered_application_for_the_report_mail(self):
        self.assertIn(self.KEY, permissions.APP_KEYS)
        self.assertEqual(permissions.APP_LABELS[self.KEY], "RDC Sales Register Report")
        self.assertIn(self.KEY, permissions.REPORT_RECIPIENT_APP_KEYS)
        spec = mail.SPECS[self.KEY]
        self.assertTrue(spec.scheduled)
        mail.validate_template(spec.default_subject, spec.placeholders)
        mail.validate_template(spec.default_intro, spec.placeholders)
        self.assertIs(mail._job_for(self.KEY), auto_mail.run_auto_mail_job)

    def test_only_an_administrator_can_use_it_even_if_it_was_granted(self):
        regular = SimpleNamespace(role="user", allowed_apps='["%s", "gst-invoice-adder"]' % self.KEY)
        self.assertEqual(permissions.parse_allowed_apps(regular), ["gst-invoice-adder"])
        self.assertFalse(permissions.user_has_app_access(regular, self.KEY))
        self.assertTrue(permissions.user_has_app_access(SimpleNamespace(role="admin", allowed_apps=None), self.KEY))

    def test_its_jobs_are_allowed_and_share_the_oracle_allowance(self):
        from app import jobs
        for task in ("run_auto_mail_job", "run_preview_job"):
            name = f"app.services.sales_register.auto_mail:{task}"
            self.assertIn(name, jobs._ALLOWED_TASKS)
            self.assertEqual(jobs._RESOURCE_BY_TASK[name], "oracle-gst")


class _Db(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine, tables=[
            User.__table__, EmailSettings.__table__, ScheduledReportMail.__table__, SystemEmailSettings.__table__])
        self.db = sessionmaker(bind=engine)()
        self.addCleanup(self.db.close)

    def person(self, email, first, sender=None, password=None):
        user = User(email=email, first_name=first, last_name="X", password_hash="x", role="admin",
                    is_active=True, is_deleted=False)
        self.db.add(user)
        self.db.commit()
        self.db.add(EmailSettings(user_id=user.id, sender_email=sender, is_deleted=False,
                                  app_password_encrypted=security.encrypt(password) if password else None))
        self.db.commit()
        return user

    def system(self, email="no-reply@rdc.in", password="sys-pass"):
        self.db.add(SystemEmailSettings(id=1, sender_email=email, app_password_encrypted=security.encrypt(password),
                                        is_deleted=False))
        self.db.commit()


class SenderChoiceTests(_Db):
    def test_the_options_are_me_the_system_email_and_other_people(self):
        me = self.person("me@rdc.in", "Meera", "meera@gmail.com", "pw1")
        self.person("asha@rdc.in", "Asha", "asha@gmail.com", "pw2")
        self.person("nomail@rdc.in", "Ben")
        self.system()
        options = mail.manual_sender_options(self.db, me)
        self.assertEqual([(o["value"], o["email"], o["configured"]) for o in options],
                         [("me", "meera@gmail.com", True), ("system", "no-reply@rdc.in", True),
                          (str(options[2]["value"]), "asha@gmail.com", True)])

    def test_each_choice_resolves_to_its_own_account(self):
        me = self.person("me@rdc.in", "Meera", "meera@gmail.com", "pw1")
        asha = self.person("asha@rdc.in", "Asha", "asha@gmail.com", "pw2")
        self.system()
        got = {c: mail.resolve_manual_sender(self.db, c, me.id)
               for c in (None, "me", "system", str(asha.id), "garbage", "999")}
        self.assertEqual((got[None]["email"], got["me"]["email"]), ("meera@gmail.com", "meera@gmail.com"))
        self.assertEqual((got["system"]["email"], got["system"]["app_password"]), ("no-reply@rdc.in", "sys-pass"))
        self.assertEqual((got[str(asha.id)]["email"], got[str(asha.id)]["app_password"]), ("asha@gmail.com", "pw2"))
        self.assertFalse(got["garbage"]["configured"])
        self.assertFalse(got["999"]["configured"])

    def test_a_system_email_that_is_not_set_up_cannot_send(self):
        me = self.person("me@rdc.in", "Meera", "meera@gmail.com", "pw1")
        sender = mail.resolve_manual_sender(self.db, "system", me.id)
        self.assertFalse(sender["configured"])
        self.assertIn("not set up", sender["problem"])

    def test_the_preview_records_who_it_is_from_and_the_confirm_uses_that(self):
        from app.services import report_mail_preview
        me = self.person("me@rdc.in", "Meera", "meera@gmail.com", "pw1")
        self.system()

        def build(snapshot, progress_cb=None, cancel_event=None):
            return mail.MailPayload(values={"date": "d"})

        with patch.object(report_mail_preview, "SessionLocal", sessionmaker(bind=self.db.get_bind())):
            result = report_mail_preview.build_preview(
                "gst-invoice-adder", build, me.id, ["a@b.in"], [], "S", "Hi {date}", sender="system")
            own = report_mail_preview.build_preview(
                "gst-invoice-adder", build, me.id, ["a@b.in"], [], "S", "Hi {date}", sender="me")
        self.assertEqual((result["status"], result["from_email"], result["sender"]), ("preview", "no-reply@rdc.in", "system"))
        self.assertEqual((own["from_email"], own["sender"]), ("meera@gmail.com", "me"))


class ScheduleTests(unittest.TestCase):
    def row(self, **kw):
        base = dict(enabled=True, is_deleted=False, send_mode="weekly", send_days="0,1,2,3,4,5", send_time="08:30",
                    send_times=None, send_month_days=None, last_attempt_date=None, last_attempt_slot=None)
        base.update(kw)
        return SimpleNamespace(**base)

    def at(self, text):
        return datetime.strptime(text, "%Y-%m-%d %H:%M").replace(tzinfo=IST)

    def test_weekly_with_one_time_behaves_as_before(self):
        row = self.row()
        self.assertTrue(mail.due_now(row, self.at("2026-10-30 08:30")))           # Friday
        self.assertFalse(mail.due_now(row, self.at("2026-10-30 08:29")))
        self.assertFalse(mail.due_now(row, self.at("2026-11-01 09:00")))          # Sunday
        row.last_attempt_date = date(2026, 10, 30)
        self.assertFalse(mail.due_now(row, self.at("2026-10-30 12:00")))
        self.assertEqual(mail.next_run(row, self.at("2026-10-30 12:00")), self.at("2026-10-31 08:30"))

    def test_several_times_a_day_each_go_out_once(self):
        row = self.row(send_time="08:30", send_times="08:30,14:00,18:15")
        self.assertEqual(mail.due_slot(row, self.at("2026-10-30 08:40")), "2026-10-30 08:30")
        row.last_attempt_slot = "2026-10-30 08:30"
        self.assertIsNone(mail.due_slot(row, self.at("2026-10-30 13:59")))
        self.assertEqual(mail.next_run(row, self.at("2026-10-30 13:00")), self.at("2026-10-30 14:00"))
        self.assertEqual(mail.due_slot(row, self.at("2026-10-30 14:00")), "2026-10-30 14:00")
        row.last_attempt_slot = "2026-10-30 18:15"
        self.assertIsNone(mail.due_slot(row, self.at("2026-10-30 23:00")))
        self.assertEqual(mail.next_run(row, self.at("2026-10-30 23:00")), self.at("2026-10-31 08:30"))

    def test_after_downtime_the_mail_goes_out_once_not_once_per_missed_slot(self):
        row = self.row(send_times="08:30,14:00,18:15")
        self.assertEqual(mail.due_slot(row, self.at("2026-10-30 19:00")), "2026-10-30 18:15")
        row.last_attempt_slot = "2026-10-30 18:15"
        self.assertIsNone(mail.due_slot(row, self.at("2026-10-30 19:05")))

    def test_monthly_on_chosen_days_and_on_the_last_day(self):
        row = self.row(send_mode="monthly", send_month_days="1,15,0")
        due = [d for d in (date(2026, 11, 1) + timedelta(days=i) for i in range(30))
               if mail.runs_on(row, d)]
        self.assertEqual(due, [date(2026, 11, 1), date(2026, 11, 15), date(2026, 11, 30)])
        self.assertTrue(mail.runs_on(row, date(2028, 2, 29)))        # last day of a leap February
        self.assertFalse(mail.runs_on(row, date(2028, 2, 28)))
        self.assertEqual(mail.next_run(row, self.at("2026-11-02 09:00")), self.at("2026-11-15 08:30"))

    def test_day_31_only_in_the_months_that_have_it(self):
        row = self.row(send_mode="monthly", send_month_days="31")
        self.assertTrue(mail.runs_on(row, date(2026, 10, 31)))
        self.assertFalse(mail.runs_on(row, date(2026, 11, 30)))
        self.assertEqual(mail.next_run(row, self.at("2026-11-02 09:00")), self.at("2026-12-31 08:30"))

    def test_switched_off_or_no_days_is_never_due(self):
        self.assertIsNone(mail.next_run(self.row(enabled=False)))
        self.assertFalse(mail.due_now(self.row(enabled=False), self.at("2026-10-30 09:00")))
        self.assertIsNone(mail.next_run(self.row(send_days="")))


class SaveScheduleTests(_Db):
    DATA = {"enabled": True, "mode": "weekly", "days": [0, 4], "month_days": [], "time": "08:30",
            "times": ["14:00", "08:30"], "to": ["a@rdc.in"], "cc": [], "subject_template": "S {date}",
            "intro_template": "Hi {date}"}

    def save(self, **changes):
        with patch.object(mail, "set_report_recipient_defaults"), \
                patch.object(mail, "get_report_recipient_defaults", lambda db, key: {"to": ["a@rdc.in"], "cc": []}), \
                patch.object(mail.audit_middleware, "log_event"), \
                patch.object(mail, "now_ist", lambda: datetime(2026, 10, 30, 7, 0, tzinfo=IST)):
            return mail.save_settings(self.db, "gst-invoice-adder", {**self.DATA, **changes})

    def test_weekly_with_several_times_is_saved_and_read_back(self):
        view = self.save()
        self.assertEqual((view["mode"], view["days"], view["times"], view["time"]),
                         ("weekly", [0, 4], ["08:30", "14:00"], "08:30"))

    def test_monthly_days_including_the_last_day_are_saved(self):
        view = self.save(mode="monthly", month_days=[15, 1, 0])
        self.assertEqual((view["mode"], view["month_days"]), ("monthly", [0, 1, 15]))

    def test_a_schedule_with_nothing_to_run_on_is_refused_unless_switched_off(self):
        with self.assertRaises(ValueError):
            self.save(mode="monthly", month_days=[])
        with self.assertRaises(ValueError):
            self.save(days=[])
        self.assertFalse(self.save(days=[], enabled=False)["enabled"])

    def test_bad_times_and_too_many_times_are_refused(self):
        with self.assertRaises(ValueError):
            self.save(times=["25:00"])
        with self.assertRaises(ValueError):
            self.save(times=[f"0{i}:00" for i in range(7)])
        with self.assertRaises(ValueError):
            self.save(mode="yearly")

    def test_saving_after_a_slot_has_passed_does_not_fire_it_but_a_later_slot_still_will(self):
        with patch.object(mail, "set_report_recipient_defaults"), \
                patch.object(mail, "get_report_recipient_defaults", lambda db, key: {"to": ["a@rdc.in"], "cc": []}), \
                patch.object(mail.audit_middleware, "log_event"), \
                patch.object(mail, "now_ist", lambda: datetime(2026, 10, 30, 10, 0, tzinfo=IST)):
            mail.save_settings(self.db, "gst-invoice-adder", self.DATA)
        row = mail.get_or_create_settings(self.db, "gst-invoice-adder")
        self.assertEqual(row.last_attempt_slot, "2026-10-30 08:30")
        self.assertIsNone(mail.due_slot(row, datetime(2026, 10, 30, 11, 0, tzinfo=IST)))
        self.assertEqual(mail.due_slot(row, datetime(2026, 10, 30, 14, 1, tzinfo=IST)), "2026-10-30 14:00")


if __name__ == "__main__":
    unittest.main()
