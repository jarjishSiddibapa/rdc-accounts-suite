"""Parses the bank's own "Account Activity" statement export into rows this
suite can store and search.

The export is a metadata block (Account Number, From/To Date, ...) followed
by a table whose header row starts with "Transaction Date". Both binary
BIFF (.xls, produced by every sample export seen so far) and genuine .xlsx
are supported; the metadata and table are located by label/header text
rather than a fixed row/column position, since that's the only part of the
layout confirmed stable across exports - column order is not assumed.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date as date_type, datetime
from decimal import Decimal, InvalidOperation

from .errors import StatementParseError

REQUIRED_COLUMNS = ("Transaction Date", "Transaction Description", "Transaction Amount")
COLUMN_ALIASES = {
    "transaction date": "Transaction Date",
    "transaction description": "Transaction Description",
    "transaction amount": "Transaction Amount",
    "debit / credit": "Debit / Credit",
    "debit/credit": "Debit / Credit",
    "reference no.": "Reference No.",
    "reference no": "Reference No.",
    "value date": "Value Date",
    "transaction branch": "Transaction Branch",
    "running balance": "Running Balance",
}


@dataclass
class BankTransactionRow:
    transaction_date: datetime
    transaction_description: str
    transaction_amount: Decimal
    debit_credit: str | None
    reference_no: str | None
    value_date: date_type | None
    transaction_branch: str | None
    running_balance: Decimal | None


@dataclass
class ParsedStatement:
    account_number: str | None
    from_date: date_type | None
    to_date: date_type | None
    rows: list[BankTransactionRow] = field(default_factory=list)


def _clean(value) -> str:
    return str(value).strip() if value is not None else ""


def _load_grid(path: str) -> list[list]:
    """Return the sheet as a plain 2D grid of raw cell values, trying the
    legacy binary format first (every sample export is BIFF) and falling
    back to a genuine .xlsx if the bank ever changes its export format."""
    try:
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        return [[sheet.cell(r, c).value for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    except Exception:
        pass
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = wb.worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            wb.close()
    except Exception as exc:
        raise StatementParseError(
            "Could not read this file as a bank statement export (neither legacy .xls nor .xlsx parsing succeeded)."
        ) from exc


def _parse_date(raw) -> date_type | None:
    text = _clean(raw)
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_datetime(raw) -> datetime | None:
    text = _clean(raw)
    if not text:
        return None
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _parse_amount(raw) -> Decimal | None:
    text = _clean(raw).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(str(raw)) if isinstance(raw, (int, float)) else Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def parse_bank_statement(path: str) -> ParsedStatement:
    grid = _load_grid(path)

    account_number: str | None = None
    from_date: date_type | None = None
    to_date: date_type | None = None
    header_row_index: int | None = None
    header_index: dict[str, int] = {}

    for r, row in enumerate(grid):
        if not row:
            continue
        label = _clean(row[0]).casefold()
        if label == "account number" and len(row) > 1:
            account_number = _clean(row[1]) or None
        elif label == "from date" and len(row) > 1:
            from_date = _parse_date(row[1])
        elif label == "to date" and len(row) > 1:
            to_date = _parse_date(row[1])
        elif label == "transaction date":
            for c, cell in enumerate(row):
                normalized = COLUMN_ALIASES.get(_clean(cell).casefold())
                if normalized:
                    header_index[normalized] = c
            header_row_index = r
            break

    if header_row_index is None or not all(col in header_index for col in REQUIRED_COLUMNS):
        raise StatementParseError(
            "This file doesn't look like a bank statement export - no 'Transaction Date' "
            "header row with the expected columns was found."
        )

    rows: list[BankTransactionRow] = []
    for row in grid[header_row_index + 1:]:
        if not row or len(row) <= header_index["Transaction Date"]:
            continue
        raw_date = row[header_index["Transaction Date"]]
        transaction_date = _parse_datetime(raw_date)
        description = _clean(row[header_index["Transaction Description"]]) if header_index["Transaction Description"] < len(row) else ""
        if transaction_date is None or not description:
            continue
        amount_idx = header_index["Transaction Amount"]
        amount = _parse_amount(row[amount_idx]) if amount_idx < len(row) else None
        if amount is None:
            continue

        def _get(col: str):
            idx = header_index.get(col)
            return row[idx] if idx is not None and idx < len(row) else None

        rows.append(
            BankTransactionRow(
                transaction_date=transaction_date,
                transaction_description=description,
                transaction_amount=amount,
                debit_credit=(_clean(_get("Debit / Credit")) or None),
                reference_no=(_clean(_get("Reference No.")) or None),
                value_date=_parse_date(_get("Value Date")),
                transaction_branch=(_clean(_get("Transaction Branch")) or None),
                running_balance=_parse_amount(_get("Running Balance")),
            )
        )

    if not rows:
        raise StatementParseError("No transaction rows were found under the header row.")

    return ParsedStatement(account_number=account_number, from_date=from_date, to_date=to_date, rows=rows)
