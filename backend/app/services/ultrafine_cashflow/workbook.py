"""Write the monthly Ultrafine Cashflow workbook.

Sheets: Summary, Payments, Receipts, Working. The Summary figures are live
formulas over the Payments / Receipts sheets (SUMIF on the CF Location
column), exactly like the manual workbook, so anyone can still edit a line
in Excel and see the Summary follow. Because Excel's Protected View does not
recalculate, every formula also gets its value cached in the file (see
app.services.xlsx_formula_cache).
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.xlsx_formula_cache import cache_formula_values, inject_cached_values

from .resolver import ISSUE_TEXT, Allocation, Maps
from .summary import MonthFigures

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=11)
_BOLD = Font(bold=True)
_SUBTOTAL_FILL = PatternFill("solid", fgColor="E8EEF6")
_UNMAPPED_FILL = PatternFill("solid", fgColor="FFF1D6")
_THIN = Side(style="thin", color="BBBBBB")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DATETIME_FORMAT = "dd/mm/yyyy hh:mm:ss"
_MONEY = "#,##0.00"
_LAKH = "#,##0.00"
# Summary table: B Sr.No. ... I Cl. Balance. Columns that carry a sum on the total rows.
LAST_COL = 9
_TOTAL_COLUMNS = (("E", "opening"), ("F", "collection"), ("G", "payable"), ("H", "paid"), ("I", "closing"))


def _ordinal(day: int) -> str:
    if 10 <= day % 100 <= 20:
        return f"{day}th"
    return f"{day}{ {1: 'st', 2: 'nd', 3: 'rd'}.get(day % 10, 'th') }"


def _style_header(ws, row: int, first_col: int, last_col: int) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_ledger(ws, rows: list[Allocation], payments: bool) -> tuple[int, str]:
    """Payments / Receipts sheet. Returns (last data row, CF column letter)."""
    headers = (
        ["Date", "Description", "Name", "Amount", "UTR No", "Location", "CF Location"]
        if payments
        else ["Date", "Description", "Name", "Amount", "Tally Location", "CF Location"]
    )
    flagged = any(a.issues for a in rows)   # written "anyway": say what is missing, line by line
    if flagged:
        headers.append("Remark")
    for offset, header in enumerate(headers):
        ws.cell(row=3, column=2 + offset, value=header)
    _style_header(ws, 3, 2, 1 + len(headers))
    row_number = 3
    for alloc in rows:
        row_number += 1
        values = [alloc.txn_dt, alloc.description, alloc.name, float(alloc.amount)]
        values += [alloc.reference_no, alloc.location, alloc.cf_location] if payments \
            else [alloc.location, alloc.cf_location]
        if flagged:
            values.append("; ".join(ISSUE_TEXT.get(code, code) for code in alloc.issues))
        for offset, value in enumerate(values):
            cell = ws.cell(row=row_number, column=2 + offset, value=value)
            if alloc.issues:
                cell.fill = _UNMAPPED_FILL
        ws.cell(row=row_number, column=2).number_format = _DATETIME_FORMAT
        ws.cell(row=row_number, column=5).number_format = _MONEY
    last = max(row_number, 4)
    ws["C2"], ws["E2"] = "Total", f"=SUBTOTAL(9,E4:E{last})"
    ws["C2"].font = ws["E2"].font = _BOLD
    ws["E2"].number_format = _MONEY
    for column, width in zip("BCDEFGH", (20, 70, 42, 16, 22, 26, 28)):
        ws.column_dimensions[column].width = width
    if flagged:
        ws.column_dimensions[get_column_letter(1 + len(headers))].width = 36
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"B3:{get_column_letter(1 + len(headers))}{last}"
    return last, "H" if payments else "G"


def _ratio_tables(maps: Maps) -> list[dict]:
    tables = []
    for group_id, group in sorted(maps.groups.items(), key=lambda item: item[1].name.casefold()):
        sets = maps.ratio_sets.get(group_id, [])
        if not sets:
            continue
        locations: list[str] = []
        for _, lines in sets:
            for line in lines:
                name = maps.locations[line.location_id].name if line.location_id in maps.locations else "?"
                if name not in locations:
                    locations.append(name)
        tables.append({"group": group.name, "locations": locations, "sets": [
            (effective, {
                (maps.locations[l.location_id].name if l.location_id in maps.locations else "?"): l.ratio
                for l in lines
            }) for effective, lines in sets
        ]})
    return tables


def _write_working(ws, maps: Maps) -> None:
    row = 1
    for table in _ratio_tables(maps):
        ws.cell(row=row, column=2, value=table["group"]).font = _TITLE_FONT
        row += 1
        ws.cell(row=row, column=2, value="PARTICULARS")
        for index, (effective, _) in enumerate(table["sets"]):
            cell = ws.cell(row=row, column=3 + index, value=effective)
            if effective.time() != datetime.min.time():
                cell.number_format = "dd-mmm-yyyy hh:mm"
            else:
                # Ratios are set per month, so a month-start date reads as "Oct 2026".
                cell.number_format = "mmm yyyy" if effective.day == 1 else "dd-mmm-yyyy"
        ws.cell(row=row, column=2 + len(table["sets"]) + 1, value=None)
        _style_header(ws, row, 2, 2 + len(table["sets"]))
        first = row + 1
        for name in table["locations"]:
            row += 1
            ws.cell(row=row, column=2, value=name)
            for index, (_, ratios) in enumerate(table["sets"]):
                cell = ws.cell(row=row, column=3 + index, value=ratios.get(name))
                cell.number_format = "0.00%"
        row += 1
        ws.cell(row=row, column=2, value="Total").font = _BOLD
        for index in range(len(table["sets"])):
            letter = get_column_letter(3 + index)
            cell = ws.cell(row=row, column=3 + index, value=f"=SUM({letter}{first}:{letter}{row - 1})")
            cell.number_format, cell.font = "0.00%", _BOLD
        row += 3
    ws.column_dimensions["B"].width = 34
    for index in range(3, 20):
        ws.column_dimensions[get_column_letter(index)].width = 16
    if row == 1:
        ws["B1"] = "No sharing ratios are configured."


def _write_summary(ws, figures: MonthFigures, as_of: datetime,
                   pay_last: int, rec_last: int) -> dict[str, float]:
    """Summary sheet; returns {coordinate: value} for every formula cell."""
    values: dict[str, float] = {}
    month_label = datetime.strptime(figures.month, "%Y-%m").strftime("%B- %Y")
    # Row 1 is left empty; row 2 is the title band: the heading centred over B:H and
    # the unit right-aligned in the last column; the table header sits on row 3.
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=LAST_COL - 1)
    ws["B2"] = f"Cash Flow as of {_ordinal(as_of.day)} {as_of:%B, %Y}"
    ws["B2"].font = _TITLE_FONT
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=2, column=LAST_COL, value="Amt in Lacs")
    ws.cell(row=2, column=LAST_COL).font = _TITLE_FONT
    ws.cell(row=2, column=LAST_COL).alignment = Alignment(horizontal="right", vertical="center")
    # "Payable (80 % Of Collection)": the heading names the percentage most rows use;
    # the odd location paying out a different share still shows its own figure below it.
    pct_counts = Counter(r.payable_pct for r in figures.rows)
    payable_header = "Payable"
    if pct_counts:
        top = pct_counts.most_common(1)[0][0]   # ties go to the one listed first
        payable_header = f"Payable ({top:g} % Of Collection)"
    headers = ["Sr.No.", "Location", "Authorized Person", "Op. Balance", f"Collection {month_label}",
               payable_header, "Actual Paid", "Cl. Balance"]
    for offset, header in enumerate(headers):
        ws.cell(row=3, column=2 + offset, value=header)
    _style_header(ws, 3, 2, LAST_COL)
    ws.row_dimensions[3].height = 30

    receipts = f"Receipts!$G$4:$G${rec_last}", f"Receipts!$E$4:$E${rec_last}"
    payments = f"Payments!$H$4:$H${pay_last}", f"Payments!$E$4:$E${pay_last}"

    row = 3
    leaf_rows: list[int] = []        # ungrouped rows counted in the grand total
    group_members: list = []         # rows of the group being built
    subtotal_rows: list[int] = []
    sr = 0

    def flush_group(group_name: str) -> None:
        nonlocal row
        row += 1
        first, last = group_members[0][0], group_members[-1][0]
        ws.cell(row=row, column=4, value=f"Total {group_name}").font = _BOLD
        for col_letter, attr in _TOTAL_COLUMNS:
            members = [getattr(r, attr) for _, r in group_members]
            total = sum(m or 0.0 for m in members)
            cell = ws[f"{col_letter}{row}"]
            cell.value = f"=SUM({col_letter}{first}:{col_letter}{last})"
            cell.number_format, cell.font = _LAKH, _BOLD
            values[cell.coordinate] = total
        for col in range(2, LAST_COL + 1):
            ws.cell(row=row, column=col).fill = _SUBTOTAL_FILL
        subtotal_rows.append(row)
        group_members.clear()

    rows = figures.rows
    for index, item in enumerate(rows):
        row += 1
        sr += 1
        ws.cell(row=row, column=2, value=sr)
        ws.cell(row=row, column=3, value=item.name)
        ws.cell(row=row, column=4, value=item.authorized_person or None)
        ws.cell(row=row, column=5, value=round(item.opening, 6))
        ws[f"F{row}"] = f"=SUMIF({receipts[0]},C{row},{receipts[1]})/100000"
        ws[f"G{row}"] = f"=F{row}*{item.payable_pct:g}%"
        ws[f"H{row}"] = f"=SUMIF({payments[0]},C{row},{payments[1]})/100000"
        # Statutory expenses keyed in Month setup still reduce the closing balance;
        # they just no longer have a column of their own on the Summary.
        statutory = f"-{round(item.statutory, 6):g}" if item.statutory else ""
        ws[f"I{row}"] = f"=E{row}+G{row}-H{row}{statutory}"
        values[f"F{row}"], values[f"G{row}"] = item.collection, item.payable
        values[f"H{row}"], values[f"I{row}"] = item.paid, item.closing
        for col in range(5, LAST_COL + 1):
            ws.cell(row=row, column=col).number_format = _LAKH
        for col in range(2, LAST_COL + 1):
            ws.cell(row=row, column=col).border = _BOX
        if item.group:
            group_members.append((row, item))
            next_group = rows[index + 1].group if index + 1 < len(rows) else None
            if next_group != item.group:
                flush_group(item.group)
        else:
            leaf_rows.append(row)

    # Lines without a CF location belong to no row, so they are not on the Summary;
    # they stay on the Payments / Receipts sheets, flagged in the Remark column.
    row += 2
    ws.cell(row=row, column=3, value="Total").font = _BOLD
    for col_letter, attr in _TOTAL_COLUMNS:
        total = sum((getattr(r, attr) or 0.0) for r in rows)
        refs = [f"{col_letter}{r}" for r in sorted(leaf_rows + subtotal_rows)]
        cell = ws[f"{col_letter}{row}"]
        cell.value = f"=SUM({','.join(refs)})" if refs else 0
        cell.number_format, cell.font = _LAKH, _BOLD
        values[cell.coordinate] = total
    for col in range(2, LAST_COL + 1):
        ws.cell(row=row, column=col).fill = _SUBTOTAL_FILL

    for column, width in zip("ABCDEFGHI", (3, 8, 30, 34, 14, 18, 14, 14, 14)):
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False
    return values


def write_cashflow_workbook(path: Path, figures: MonthFigures, maps: Maps) -> dict:
    allocations = figures.allocations
    payments = [a for a in allocations if a.direction == "D"]
    receipts = [a for a in allocations if a.direction == "C"]
    as_of = max((a.txn_dt for a in allocations), default=datetime.strptime(figures.month, "%Y-%m"))

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    payments_ws = wb.create_sheet("Payments")
    receipts_ws = wb.create_sheet("Receipts")
    working_ws = wb.create_sheet("Working")

    pay_last, _ = _write_ledger(payments_ws, payments, payments=True)
    rec_last, _ = _write_ledger(receipts_ws, receipts, payments=False)
    _write_working(working_ws, maps)
    summary_values = _write_summary(summary_ws, figures, as_of, pay_last, rec_last)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    cached = cache_formula_values(wb)
    cached["Summary"] = summary_values   # SUMIF cells cannot be resolved generically
    inject_cached_values(str(path), cached)
    return {
        "payment_rows": len(payments),
        "receipt_rows": len(receipts),
        "payment_total": float(sum((a.amount for a in payments), Decimal(0))),
        "receipt_total": float(sum((a.amount for a in receipts), Decimal(0))),
    }
