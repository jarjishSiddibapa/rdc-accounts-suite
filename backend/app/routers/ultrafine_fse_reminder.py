"""Ultrafine FSE Bulk Reminder routes.

Reads the credit-control team's "Coll vs Target" tracker sheet (see
app/services/ultrafine_fse_reminder/processor.py for the column-detection
and mail-building logic) and sends: one reminder email PER FSE (Field Sales
Executive), plus one broadcast combining every FSE for management.

Mirrors this suite's established two-step "preview then confirm-send" mail
pattern used by the ultrafine-payment-reminder / ultrafine-balance-
confirmation siblings, with one deliberate difference: those tools always
send the exact content their own preview job built, while this tool lets the
user directly edit each row's subject/body/to/cc in the browser before
sending (per the user's explicit request), so /send and /send-broadcast take
the (possibly edited) content straight from the request body instead of
re-reading a stored job result.
"""

import io
import uuid
from pathlib import Path
from typing import Optional

import openpyxl
import pandas as pd
from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.config import SCRATCH_DIR
from app.database import get_db
from app.jobs import JobUserError, cancel_job, get_job, run_cpu_phase, submit_job
from app.permissions import require_app_access
from app.services import mailer_shared
from app.services.ultrafine_fse_reminder import mapping_store, processor
from app.uploads import save_upload

router = APIRouter(
    prefix="/api/tools/ultrafine-fse-reminder",
    tags=["ultrafine-fse-reminder"],
    dependencies=[Depends(require_app_access("ultrafine-fse-reminder"))],
)

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


_TITLE_FILL = PatternFill("solid", fgColor="A9D08E")
_HEADER_FILL = PatternFill("solid", fgColor="F4B183")
_TARGET_FILL = PatternFill("solid", fgColor="FFFF00")
_BODY_FILL = PatternFill("solid", fgColor="DCE6F1")
_BORDER = Border(*(Side(style="thin", color="4472C4"),) * 4)
_BOLD = Font(bold=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_row(ws, row: int, ncols: int, fill: PatternFill, bold: bool = False, align: Optional[Alignment] = None) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        cell.border = _BORDER
        if bold:
            cell.font = _BOLD
        if align:
            cell.alignment = align


@router.get("/template")
def download_template():
    """An example workbook shaped like the real tracker's 'Coll vs Target'
    sheet, generated on the fly: title row, single FSE/Party's Name/Target/
    Received/Short Fall header (no duplicate FSE column and no unrelated
    decoy Target columns - the app never needed either; they were just
    clutter carried over from the original tracker), two real example FSEs
    with their real customer names and figures, so the shape is immediately
    recognisable. Deliberately has NO per-FSE "Total" rows, no "Grand
    Total" row, and no instructional note explaining that either (it was a
    merged cell that just got in the way): read_coll_vs_target() only ever
    sums the detail rows itself (see build_send_plan) and never reads a
    Total row's own numbers, so there's nothing for a filler to hand-total
    or keep in sync - just one row per party under each FSE. The "as
    on"/month-end dates stay dynamic so the downloaded file never looks
    stale, while the illustrative rows use real historical figures for
    clarity."""
    today = pd.Timestamp.now()
    month_end = (today + pd.offsets.MonthEnd(0)).strftime("%d-%b-%y")
    as_on = today.strftime("%d-%b-%y")
    month_label = today.strftime("%b-%y")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = processor.SHEET_NAME
    ncols = 5

    ws.cell(1, 1).value = f"Collection VS Target Summary for {month_label}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    _style_row(ws, 1, ncols, _TITLE_FILL, bold=True, align=_CENTER)

    headers = [
        "FSE", "Party's Name",
        f"Collection Target considering dues upto  {month_end}",
        f"Coll Received as on {as_on}",
        "Short Fall",
    ]
    ws.append(headers)
    _style_row(ws, 2, ncols, _HEADER_FILL, bold=True, align=_CENTER)
    ws.cell(2, 3).fill = _TARGET_FILL

    # Real illustrative figures (kept from an actual month's tracker) so the
    # layout is self-explanatory: one FSE with two customers, one FSE with a
    # single customer - both patterns a filler will actually hit. No Total /
    # Grand Total rows - see the docstring above for why they're not needed.
    rows = [
        ("Abhishek Nayak", "Dineshchandra R. Agarwal Infracon Pvt Ltd - Drs", 3.0403255, 2.4426, 0.5977255),
        ("Abhishek Nayak", "New India Ceramic Engineers - Drs", 10.72791, 0, 10.72791),
        ("Balram Chakrawarti", "Ahluwalia Construction Group - Drs", 19.4248, 0, 19.4248),
    ]
    row_num = 2
    for fse, party, target, received, shortfall in rows:
        row_num += 1
        ws.append([fse, party, target, received, shortfall])
        _style_row(ws, row_num, ncols, _BODY_FILL)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 14
    ws.freeze_panes = "A3"

    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=_XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="FSE_Collection_Reminder_Input_Template.xlsx"'},
    )


def _cpu_phase_preview(
    path: str, mapping: dict, signature: str, as_on_date: Optional[str],
    advisory_html: Optional[str],
) -> dict:
    """100% CPU (openpyxl read + plan building), no I/O - runs on the CPU
    process pool, matching the sibling ultrafine mail tools."""
    parsed = processor.read_coll_vs_target(path)
    if as_on_date:
        processor.apply_as_on_override(parsed, as_on_date)
    return processor.build_send_plan(parsed, mapping, signature, advisory_html)


def _job_preview(
    path: str, mapping: dict, signature: str, as_on_date: Optional[str],
    advisory_html: Optional[str], progress_cb=None,
) -> dict:
    try:
        if progress_cb:
            progress_cb(0.05, "Reading tracker file...")
        plan = run_cpu_phase(_cpu_phase_preview, path, mapping, signature, as_on_date, advisory_html)
        if progress_cb:
            progress_cb(0.95, "Preview ready")
    finally:
        Path(path).unlink(missing_ok=True)
    return {"status": "preview", **plan}


@router.post("/preview")
async def preview(
    file: UploadFile = File(...),
    as_on_date: Optional[str] = Form(None),
    advisory_html: Optional[str] = Form(None),
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Step 1 of 2: parse the uploaded tracker workbook and build both the
    per-FSE and broadcast send plans as a background job. `as_on_date`
    (ISO "YYYY-MM-DD", from the frontend's date picker, defaulting to today
    but user-editable) overrides the date baked into the subject/body/table
    header - it's independent of the file's own "received as on" column,
    which is only ever used to locate the right column, never its value.
    `advisory_html` (from the frontend's shared, optional rich-text field)
    overrides the fixed advisory paragraph for EVERY individual reminder and
    the broadcast alike - see processor.build_send_plan's docstring. Never
    sends anything - the frontend shows this to the user, who can edit
    subject/body/to/cc directly before calling /send or /send-broadcast."""
    settings = mailer_shared.get_email_settings(user.id)
    if not settings.get("configured"):
        raise HTTPException(
            status_code=400,
            detail="You haven't set up your email sender yet — go to Settings.",
        )

    job_dir = SCRATCH_DIR / f"fse-reminder-{uuid.uuid4()}"
    job_dir.mkdir(parents=True, exist_ok=True)
    path = str(await save_upload(file, job_dir))

    mapping = mapping_store.load_all(db)
    job_id = submit_job(
        _job_preview,
        path,
        mapping,
        settings.get("signature", ""),
        as_on_date,
        advisory_html,
        owner_id=user.id,
    )
    return {"job_id": job_id}


@router.get("/jobs/{job_id}")
def job_status(job_id: str, user=Depends(get_current_user)):
    job = get_job(job_id, owner_id=user.id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")
    return job


@router.post("/jobs/{job_id}/cancel")
def cancel(job_id: str, user=Depends(get_current_user)):
    job = cancel_job(job_id, owner_id=user.id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")
    return job


# ── Send (takes the user's current, possibly-edited content directly) ──────

class SendRow(BaseModel):
    fse_key: str
    to: list[str] = []
    cc: list[str] = []
    subject: str
    body_html: str


class SendBody(BaseModel):
    rows: list[SendRow]


def _job_send_rows(user_id: int, rows: list[dict], progress_cb=None) -> dict:
    settings = mailer_shared.get_email_settings(user_id)
    if not settings.get("configured"):
        raise JobUserError("Your email sender is no longer configured. Update Settings and try again.")

    report = []
    total = len(rows) or 1
    for index, row in enumerate(rows, start=1):
        to = [addr for addr in row["to"] if addr.strip()]
        cc = [addr for addr in row["cc"] if addr.strip()]
        if not to:
            report.append({"fse_key": row["fse_key"], "status": "skipped", "detail": "No To recipient"})
        else:
            invalid = [addr for addr in to + cc if not processor.is_valid_email_syntax(addr)]
            if invalid:
                report.append({"fse_key": row["fse_key"], "status": "failed", "detail": f"Invalid email address: {invalid[0]}"})
            else:
                try:
                    mailer_shared.send_mail(
                        settings["email"], settings["app_password"], to, cc,
                        row["subject"], row["body_html"], attachments=[],
                    )
                    report.append({"fse_key": row["fse_key"], "status": "sent", "detail": "Sent successfully"})
                except Exception as exc:  # noqa: BLE001 - per-row send boundary
                    report.append({"fse_key": row["fse_key"], "status": "failed", "detail": str(exc)})
        if progress_cb:
            progress_cb(min(0.99, index / total), f"Sent {index} of {len(rows)}")

    return {
        "status": "sent",
        "report": report,
        "sent": sum(1 for r in report if r["status"] == "sent"),
        "failed": sum(1 for r in report if r["status"] == "failed"),
        "skipped": sum(1 for r in report if r["status"] == "skipped"),
    }


@router.post("/send")
def send(body: SendBody, user=Depends(get_current_user)):
    """Send one or more individual FSE reminders, using exactly the
    subject/body/to/cc the user currently has on screen (this tool lets
    those be edited after preview, unlike its siblings - see module
    docstring)."""
    if not body.rows:
        raise HTTPException(status_code=400, detail="No rows to send.")
    job_id = submit_job(
        _job_send_rows,
        user.id,
        [row.model_dump() for row in body.rows],
        owner_id=user.id,
    )
    return {"job_id": job_id}


class TableRowIn(BaseModel):
    fse: str
    party: str
    target: float
    received: float
    shortfall: float


class SendBroadcastBody(BaseModel):
    to: list[str] = []
    cc: list[str] = []
    subject: str
    body_html: str
    attach_excel: bool = True
    table_title: str = ""
    target_header: str = ""
    received_header: str = ""
    table_rows: list[TableRowIn] = []


def _job_send_broadcast(
    user_id: int, to: list[str], cc: list[str], subject: str, body_html: str,
    attach_excel: bool, table_title: str, target_header: str, received_header: str,
    table_rows: list[dict], progress_cb=None,
) -> dict:
    settings = mailer_shared.get_email_settings(user_id)
    if not settings.get("configured"):
        raise JobUserError("Your email sender is no longer configured. Update Settings and try again.")

    to = [addr for addr in to if addr.strip()]
    cc = [addr for addr in cc if addr.strip()]
    if not to:
        raise JobUserError("No To recipient for the broadcast mail.")
    invalid = [addr for addr in to + cc if not processor.is_valid_email_syntax(addr)]
    if invalid:
        raise JobUserError(f"Invalid email address: {invalid[0]}")

    attachments: list[str] = []
    tmp_path: Optional[Path] = None
    if attach_excel and table_rows:
        if progress_cb:
            progress_cb(0.1, "Building Excel attachment...")
        workbook_bytes = processor.build_broadcast_workbook(table_title, target_header, received_header, table_rows)
        tmp_path = SCRATCH_DIR / f"fse-broadcast-{uuid.uuid4()}.xlsx"
        tmp_path.write_bytes(workbook_bytes)
        attachments = [str(tmp_path)]

    if progress_cb:
        progress_cb(0.5, "Sending broadcast mail...")
    try:
        mailer_shared.send_mail(settings["email"], settings["app_password"], to, cc, subject, body_html, attachments=attachments)
    finally:
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
    if progress_cb:
        progress_cb(0.99, "Sent")
    return {"status": "sent", "to": to, "cc": cc}


@router.post("/send-broadcast")
def send_broadcast(body: SendBroadcastBody, user=Depends(get_current_user)):
    job_id = submit_job(
        _job_send_broadcast,
        user.id,
        body.to,
        body.cc,
        body.subject,
        body.body_html,
        body.attach_excel,
        body.table_title,
        body.target_header,
        body.received_header,
        [row.model_dump() for row in body.table_rows],
        owner_id=user.id,
    )
    return {"job_id": job_id}


# ── FSE -> Email mapping CRUD (independent of sending) ──────────────────────

class MappingBody(BaseModel):
    fse_name: str
    email: str = ""


@router.get("/mappings")
def list_mappings(db: Session = Depends(get_db)):
    mapping = mapping_store.load_all(db)
    return [{"fse_name": name, "email": email} for name, email in sorted(mapping.items())]


@router.post("/mappings")
def create_mapping(body: MappingBody, db: Session = Depends(get_db)):
    fse_name = body.fse_name.strip()
    if not fse_name:
        raise HTTPException(status_code=400, detail="fse_name is required")
    mapping_store.upsert_fse_mapping(db, fse_name, body.email)
    return {"ok": True}


@router.put("/mappings/{fse_name}")
def update_mapping(fse_name: str, body: MappingBody, db: Session = Depends(get_db)):
    new_name = body.fse_name.strip()
    if not new_name:
        raise HTTPException(status_code=400, detail="fse_name is required")
    if new_name != fse_name:
        if not mapping_store.delete_fse_mapping(db, fse_name):
            raise HTTPException(status_code=404, detail="Mapping not found")
    mapping_store.upsert_fse_mapping(db, new_name, body.email)
    return {"ok": True}


@router.delete("/mappings/{fse_name}")
def delete_mapping(fse_name: str, db: Session = Depends(get_db)):
    if not mapping_store.delete_fse_mapping(db, fse_name):
        raise HTTPException(status_code=404, detail="Mapping not found")
    return {"ok": True}
